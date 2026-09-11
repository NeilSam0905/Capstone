"""
scripts/model_benchmark_ml.py
------------------------------------------------------------------
Extends scripts/model_benchmark.py's ranking with three tree-based ML
baselines that are NOT part of the committed benchmark: XGBoost,
LightGBM, Random Forest (forecasting/ml_models.py). EXPERIMENTAL - see
requirements/requirements-ml-experimental.txt for why these stay out of
requirements.txt.

Scored on the SAME walk-forward folds, the SAME DB, and the SAME
service-metric formula as scripts/model_benchmark.py (folds are a
deterministic function of series length, so re-deriving them here from
the identical DB and defaults reproduces the committed run's layout
exactly - no fold is shared by import, only by construction). `naive` is
re-run alongside the ML methods so pct_skus_beating_naive is computed
from THIS run's folds, not spliced in from the committed CSV.

Run (from the repo root):
    python scripts/model_benchmark_ml.py [--max-folds N] [--limit N]

Writes:
    data/model_benchmark_ml_results.csv    (per SKU, per method, per fold)
    data/model_benchmark_ml_summary.csv    (ranked summary, ML methods only)
    Model_Comparison.xlsx                  (combined with the committed
                                             statistical methods, repo root)
------------------------------------------------------------------
"""
import argparse
import os
import sys
import time

import numpy as np
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import model_benchmark as mb
from forecasting.category import build_service_class_fn
from forecasting.evaluate import evaluate_methods, summarise
from forecasting.ml_models import (
    lightgbm_fit_predict, random_forest_fit_predict, xgboost_fit_predict,
)
from forecasting.baselines import naive_fit_predict

OUT_CSV = "data/model_benchmark_ml_results.csv"
SUMMARY_CSV = "data/model_benchmark_ml_summary.csv"
XLSX_OUT = "Model_Comparison.xlsx"
COMMITTED_SUMMARY_CSV = "data/model_benchmark_summary.csv"


def build_ml_methods():
    return {
        "naive": naive_fit_predict(),
        "xgboost": xgboost_fit_predict(n_estimators=50),
        "lightgbm": lightgbm_fit_predict(n_estimators=50),
        "random_forest": random_forest_fit_predict(n_estimators=50),
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--max-folds", type=int, default=mb.MAX_FOLDS)
    ap.add_argument("--limit", type=int, default=None,
                    help="benchmark only the first N SKUs (for a smoke run)")
    args = ap.parse_args()

    con = mb.sqlite3.connect(mb.DB_NAME)
    series, names, index = mb.load_daily_series(con, args.limit)
    con.close()

    print(f"Loaded {len(series)} moving SKUs over {len(index)} calendar days "
          f"({index[0].date()} .. {index[-1].date()})")

    methods = build_ml_methods()
    print(f"ML methods ({len(methods)}): {', '.join(methods)}\n"
          f"(naive is re-run here only so pct_skus_beating_naive is computed "
          f"on THIS run's folds)")

    t0 = time.time()
    results, insufficient = evaluate_methods(
        series, methods, horizon=mb.HORIZON, min_folds=mb.MIN_FOLDS,
        max_folds=args.max_folds, min_train=mb.MIN_TRAIN)
    elapsed = time.time() - t0

    if results.empty:
        print("No SKU had enough history to score. Nothing written.")
        return 1

    # Fold-scoped Fast/Slow for the safety-stock class, not the committed
    # full-history Dim_Product.fsn_class - see mb.service_metrics().
    results = mb.service_metrics(results, series, build_service_class_fn(series))
    results["item_name"] = results["sku"].map(names)
    results.to_csv(OUT_CSV, index=False, lineterminator="\n")

    n_scored = results["sku"].nunique()
    print(f"Scored {n_scored} SKUs x {len(methods)} methods "
          f"in {elapsed:.1f}s -> {OUT_CSV} ({len(results):,} rows)")
    print(f"Insufficient history (<{mb.MIN_FOLDS} folds): {len(insufficient)} SKUs\n")

    summary = summarise(results)
    beats = mb.beats_naive(results)
    summary["pct_skus_beating_naive"] = (
        summary["method"].map(beats).astype(float).round(1))

    priced = mb.skus_priced(series, methods)
    svc = (results.groupby("method")
                  .agg(units_served=("units_served", "sum"),
                       units_short=("units_short", "sum"),
                       units_held=("units_held", "sum"),
                       demand=("actual_30d", "sum"))
                  .reset_index())
    svc["fill_rate_at_target"] = (svc["units_served"] / svc["demand"]).round(4)
    summary = summary.merge(
        svc[["method", "fill_rate_at_target", "units_short", "units_held"]],
        on="method", how="left")
    summary["n_skus_priced"] = summary["method"].map(priced).astype(int)
    summary["n_skus_priced_per_fold"] = summary["method"].map(
        mb.skus_priced_per_fold(results)).astype(float)
    # drop the re-run naive row - the committed CSV already has it, and this
    # run's copy exists only to seed pct_skus_beating_naive
    ml_only = summary[summary["method"] != "naive"].reset_index(drop=True)
    ml_only.to_csv(SUMMARY_CSV, index=False, lineterminator="\n")
    print(f"Wrote {SUMMARY_CSV}\n")

    write_combined_xlsx(ml_only)
    return 0


def write_combined_xlsx(ml_summary: pd.DataFrame):
    """Committed statistical methods (data/model_benchmark_summary.csv) +
    the new ML rows, in the same two-table format model_benchmark.py
    prints, written to the repo root as an .xlsx so it opens immediately."""
    committed = pd.read_csv(COMMITTED_SUMMARY_CSV)
    is_new = pd.Series(
        [False] * len(committed) + [True] * len(ml_summary), name="new_model")
    combined = pd.concat([committed, ml_summary], ignore_index=True)
    combined["new_model"] = is_new

    err_cols = ["method", "mae", "rmse", "mase", "pct_skus_beating_naive", "new_model"]
    dec_cols = ["method", "n_skus_priced", "n_skus_priced_per_fold",
                "fill_rate_at_target",
                "units_short", "units_held", "new_model"]

    by_mase = combined.sort_values(["mase", "mae"], na_position="last",
                                   kind="stable")[err_cols].reset_index(drop=True)
    by_fill = combined.sort_values(["fill_rate_at_target", "n_skus_priced"],
                                   ascending=False,
                                   kind="stable")[dec_cols].reset_index(drop=True)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    highlight = PatternFill("solid", fgColor="FFF2CC")
    bold = Font(bold=True)

    wb = Workbook()

    def write_sheet(ws, df, title):
        ws.title = title
        ws.append(list(df.columns))
        for cell in ws[1]:
            cell.font = bold
        for _, row in df.iterrows():
            ws.append(list(row))
        new_col = df.columns.get_loc("new_model") + 1
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=new_col).value:
                for c in range(1, ws.max_column + 1):
                    ws.cell(row=r, column=c).fill = highlight
        ws.delete_cols(new_col)
        for c, col in enumerate(df.columns, start=1):
            if col == "new_model":
                continue
            width = max(len(str(col)), *(len(f"{v:.4f}" if isinstance(v, float) else str(v))
                                         for v in df[col])) + 2
            ws.column_dimensions[get_column_letter(c if c < new_col else c - 1)].width = width

    write_sheet(wb.active, by_mase, "Error metrics (by MASE)")
    write_sheet(wb.create_sheet(), by_fill, "Decision metrics (by fill rate)")

    wb.save(XLSX_OUT)
    print(f"Wrote {XLSX_OUT} ({len(combined)} methods: "
          f"{len(committed)} committed + {len(ml_summary)} new)")
    print("Highlighted rows (light yellow) are the new ML methods.")


if __name__ == "__main__":
    sys.exit(main())
