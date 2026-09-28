"""
scripts/build_testing_summary_workbook.py
------------------------------------------------------------------
Builds data/USTore_Forecast_Testing_Summary.xlsx: the forecasting tests run
in this project, in one workbook. Category level: MASE, MAPE, the Prophet
re-scoring correction, the TSB/intermittent-model test, the
best-model-per-category scan, the parameter-tuning sweep and the daily-shape
test. Item level: all 47 methods scored on the Fast items, their robustness
checks and the item daily-shape test.

Pulls only from files/tables other scripts already produced (does not
re-run any test). Run those first if this errors on a missing file:
  scripts/step4c_category_forecast.py
  scripts/compare_category_forecast_methods.py
  scripts/tune_category_forecast_params.py
  scripts/test_category_forecast_shape.py
  scripts/test_item_forecast_methods.py
  scripts/test_lead_time_gap.py
  scripts/test_calendar_adjustment.py
  scripts/test_transformations.py
data/category_benchmark_folds.csv and category_benchmark_summary.csv are
marco's original 27-method run (scripts/benchmark_category_level.py).

Run:  python scripts/build_testing_summary_workbook.py
------------------------------------------------------------------
"""
import os
import sqlite3
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
OUT = os.path.join(DATA, "USTore_Forecast_Testing_Summary.xlsx")

FONT = "Calibri"
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT, size=14, bold=True, color="1F4E78")
NOTE_FONT = Font(name=FONT, size=10, italic=True, color="595959")
BODY_FONT = Font(name=FONT, size=11)
BOLD_FONT = Font(name=FONT, size=11, bold=True)
SHIP_FILL = PatternFill("solid", fgColor="E2EFDA")     # highlight the shipped model's row
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CATS = ["Apparel Accessories", "Bags", "Drinkware", "Home & Novelty",
       "Keychains & Charms", "Lanyards & IDs", "Outerwear",
       "Plush & Souvenirs", "Shirts & Tops", "Stationery",
       "Umbrellas & Gear", "Uncategorised"]


def style_header(ws, row, ncols, start_col=1):
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_table(ws, start_row, headers, rows, widths, highlight_rows=None, number_formats=None):
    """headers: list[str]; rows: list[list]; number_formats: dict{col_idx0: fmt}."""
    highlight_rows = highlight_rows or set()
    for j, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=j, value=h)
    style_header(ws, start_row, len(headers))
    for i, row in enumerate(rows):
        r = start_row + 1 + i
        for j, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=j, value=val)
            cell.font = BOLD_FONT if i in highlight_rows else BODY_FONT
            cell.border = BORDER
            if i in highlight_rows:
                cell.fill = SHIP_FILL
            if number_formats and (j - 1) in number_formats:
                cell.number_format = number_formats[j - 1]
    autosize(ws, widths)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1).coordinate
    return start_row + 1 + len(rows)


def title_block(ws, title, subtitle_lines):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    r = 2
    for line in subtitle_lines:
        ws.cell(row=r, column=1, value=line).font = NOTE_FONT
        r += 1
    return r + 1


# ------------------------------------------------------------ data pulls
def load_shipped():
    con = sqlite3.connect("file:%s?mode=ro" % os.path.join(ROOT, "ustore.db"), uri=True)
    df = pd.read_sql_query("""
        SELECT forecast_category, n_obs, mae, rmse, mape, mase,
               naive_mae, naive_mase, beats_naive_mae, mean_actual_30d
        FROM Result_Category_Forecast_Metrics ORDER BY forecast_category
    """, con)
    con.close()
    return df


def per_category_mase(df, methods, actual_col="actual", scale_col="scale_cal"):
    out = {}
    for cat, g in df.groupby("forecast_category"):
        scale = g[scale_col].mean()
        out[cat] = {m: (g[actual_col] - g[m]).abs().mean() / scale for m in methods}
    return pd.DataFrame(out).T.reindex(CATS)


def per_category_mape(df, methods, actual_col="actual"):
    out = {}
    for cat, g in df.groupby("forecast_category"):
        row = {}
        for m in methods:
            nz = g[actual_col] != 0
            row[m] = (100 * (g[actual_col][nz] - g[m][nz]).abs() / g[actual_col][nz]).mean()
        out[cat] = row
    return pd.DataFrame(out).T.reindex(CATS)


# ---------------------------------------------------------------- sheets
def sheet_overview(wb):
    ws = wb.active
    ws.title = "Overview"
    r = title_block(ws, "USTore Forecast — Testing Summary",
                    ["The forecasting tests run for this project (category level and item level), in one workbook.",
                     "Generated from scripts/build_testing_summary_workbook.py — re-run it after any test changes."])
    lines = [
        ("Tab", "What it shows"),
        ("Shipped Model Results", "The model actually running in the app today (6-month average per category, lowered when the school calendar shows quieter days ahead): MASE, MAPE, MAE, RMSE per category, from Result_Category_Forecast_Metrics."),
        ("Method Comparison", "6-month average vs. 3-month, 12-month, Prophet (3 versions), and a Prophet+average blend — all scored on the same test months. This is what decided the shipped model."),
        ("Prophet Re-scoring", "Why the original “8 of 12 categories beat naive, MASE 0.99” result did not hold up, with the corrected numbers."),
        ("Intermittent Models (TSB)", "Whether TSB/SBA/Croston (built for slow, irregular sellers) beat the shipped model — tested per category."),
        ("Best Model per Category", "Scanning all 27 previously-tested methods for each category's best scorer. Kept for the record; see the caution note on that tab before using it."),
        ("Parameter Tuning", "Whether adjusting TSB's and Prophet's own settings (not swapping the model) could close the gap. It didn't."),
        ("Daily Shape Test", "Whether a date-aware model (Prophet etc.) can give a better day-by-day pattern than a flat line, keeping the 6-month total. Small but real gain, mostly in semester-break months."),
        ("Item Method Results", "ITEM level: all 47 methods (Prophet, averages, TSB, category-share, blends such as 50/50) scored on the 58 Fast items. This is what replaced Prophet as the item forecast."),
        ("Item Robustness", "ITEM level: is the winner luck? Old-vs-new period ranking, paired comparisons with a 95% range, a settings grid, and error per 30-day window."),
        ("Item Daily Shape", "ITEM level: which day-by-day pattern (the item's own or its category's; Prophet or weekday) beats a flat line."),
        ("Calendar Adjustment", "Lowering the 30-day total when the school calendar shows quieter days ahead (semester break, exams). Small but consistent gain at both levels; now used by the app."),
        ("Transformations", "Whether log, square-root or Yeo-Johnson transforms of the sales data make the forecasts more accurate, at category and item level. They do not; Yeo-Johnson fails on this data."),
        ("Lead-Time Validation Check", "Disclosure only, not the headline: re-tests both shipped models as if a forecast could only be acted on after the real reorder delay (14-28 days), not the next day."),
    ]
    for i, (a, b) in enumerate(lines):
        row = r + i
        ws.cell(row=row, column=1, value=a).font = BOLD_FONT if i == 0 else BODY_FONT
        ws.cell(row=row, column=2, value=b).font = BOLD_FONT if i == 0 else BODY_FONT
        if i == 0:
            style_header(ws, row, 2)
        for c in (1, 2):
            ws.cell(row=row, column=c).border = BORDER
            ws.cell(row=row, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    autosize(ws, [26, 95])
    ws.row_dimensions[r].height = 30
    for i in range(1, len(lines)):
        ws.row_dimensions[r + i].height = 42

    r2 = r + len(lines) + 2
    ws.cell(row=r2, column=1, value="Bottom line").font = BOLD_FONT
    ws.cell(row=r2 + 1, column=1,
           value=("MASE is used instead of MAPE because MAPE is skewed by small-volume "
                  "categories even after grouping by category. The 6-month average beat every "
                  "alternative tested — other trailing-average windows, three versions of "
                  "Prophet, TSB/SBA/Croston, and 22 tuned parameter settings — so it is what "
                  "the app uses, now lowered when the school calendar shows quieter days ahead "
                  "(macro MASE 1.14 to 1.09; see the Calendar Adjustment tab). It beats a strict "
                  "‘naive’ benchmark (MASE below 1) in 6 of 12 categories; the manuscript's ≤20% "
                  "MAPE target is not reachable with the current ~2 years of sales history.")).font = BODY_FONT
    ws.cell(row=r2 + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r2 + 1, start_column=1, end_row=r2 + 1, end_column=2)
    ws.row_dimensions[r2 + 1].height = 75

    ws.cell(row=r2 + 3, column=1, value="Item level").font = BOLD_FONT
    ws.cell(row=r2 + 4, column=1,
            value=("The previous item forecast (Prophet, mean MASE 2.52) lost to simple methods. The app now "
                   "uses a 50/50 blend of two things: the item's category level split by the item's recent "
                   "share, and TSB, a smoothed recent-sales rate. Mean MASE is 1.71, better on 46 of 57 items, and 1.67 after the calendar adjustment. "
                   "13 of the 58 “Fast” items sold nothing in the last year and Prophet kept forecasting them; "
                   "on the 45 items still selling the gain is smaller but real (2.61 to 2.16). Item forecasts "
                   "are still not accurate in absolute terms, and the blend under-forecasts busy months and "
                   "over-forecasts quiet ones because it has no seasonal term.")).font = BODY_FONT
    ws.cell(row=r2 + 4, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r2 + 4, start_column=1, end_row=r2 + 4, end_column=2)
    ws.row_dimensions[r2 + 4].height = 92


def sheet_shipped(wb, shipped):
    ws = wb.create_sheet("Shipped Model Results")
    r = title_block(ws, "Shipped Model — 6-Month Average per Category, Calendar-Adjusted",
                    ["Source: Result_Category_Forecast_Metrics (scripts/step4c_category_forecast.py). "
                     "Walk-forward test, up to 12 past 30-day periods per category.",
                     "MASE = the model's typical error ÷ a simple benchmark's typical error. Below 1 means it beat that benchmark.",
                     "MAPE = typical error as a % of actual sales — shown for completeness, but unreliable for small categories (see the Overview tab)."])
    headers = ["Category", "Periods tested", "MAE (units)", "RMSE (units)", "MAPE (%)",
              "MASE", "Naive MAE", "Naive MASE", "Beats naive?", "Avg. actual units/30d"]
    rows = []
    for rec in shipped.sort_values("mase").to_dict("records"):
        rows.append([rec["forecast_category"], rec["n_obs"], rec["mae"], rec["rmse"],
                    rec["mape"], rec["mase"], rec["naive_mae"], rec["naive_mase"],
                    "Yes" if rec["beats_naive_mae"] else "No", rec["mean_actual_30d"]])
    end = write_table(ws, r, headers, rows,
                      widths=[22, 13, 12, 12, 11, 9, 11, 11, 12, 18],
                      number_formats={2: "0.0", 3: "0.0", 4: "0.0", 5: "0.00", 6: "0.0", 7: "0.00", 9: "0.0"})
    n = len(rows)
    ws.cell(row=end, column=1, value="Macro average").font = BOLD_FONT
    for j, col in [(2, "B"), (3, "C"), (4, "D"), (5, "E"), (6, "F")]:
        ws.cell(row=end, column=j, value="=AVERAGE(%s%d:%s%d)" % (col, r + 1, col, end - 1))
        ws.cell(row=end, column=j).font = BOLD_FONT
        ws.cell(row=end, column=j).number_format = "0.00"
    ws.cell(row=end, column=9, value='=COUNTIF(I%d:I%d,"Yes")&" of "&%d' % (r + 1, end - 1, n))
    ws.cell(row=end, column=9).font = BOLD_FONT
    for c in range(1, 11):
        ws.cell(row=end, column=c).border = BORDER
    ws.cell(row=end + 2, column=1,
           value="“Beats naive” here compares MAE against repeating the previous month's total — "
                 "an easier bar than MASE < 1.").font = NOTE_FONT


def sheet_method_comparison(wb, folds):
    ws = wb.create_sheet("Method Comparison")
    r = title_block(ws, "Method Comparison — MASE by Category",
                    ["Source: data/category_forecast_method_folds.csv (scripts/compare_category_forecast_methods.py).",
                     "Every method scored on the SAME test months, using only data known before each forecast date (no hindsight).",
                     "MASE denominator: in-sample naive error on 30-calendar-day blocks (same unit for every method — see the Prophet Re-scoring tab for why this matters)."])
    methods = ["naive_last30", "mean_90", "mean_180", "mean_365",
              "prophet_marco", "prophet_zf_hind", "prophet_zf_expected", "blend_50"]
    labels = {"naive_last30": "Repeat last month", "mean_90": "3-month average",
             "mean_180": "6-month average (SHIPPED)", "mean_365": "12-month average",
             "prophet_marco": "Prophet (as originally run)",
             "prophet_zf_hind": "Prophet, zero-filled (hindsight)",
             "prophet_zf_expected": "Prophet, zero-filled (realistic)",
             "blend_50": "50/50 blend: Prophet + 6-month avg"}
    tbl = per_category_mase(folds, methods)
    headers = ["Category"] + [labels[m] for m in methods]
    rows = [[cat] + [tbl.loc[cat, m] for m in methods] for cat in CATS]
    shipped_col = methods.index("mean_180") + 1
    end = write_table(ws, r, headers, rows, widths=[20] + [16] * len(methods),
                      number_formats={i: "0.00" for i in range(1, len(headers))})
    ws.cell(row=end, column=1, value="Macro average").font = BOLD_FONT
    for j in range(2, len(headers) + 1):
        col = get_column_letter(j)
        ws.cell(row=end, column=j, value="=AVERAGE(%s%d:%s%d)" % (col, r + 1, col, end - 1))
        ws.cell(row=end, column=j).font = BOLD_FONT
        ws.cell(row=end, column=j).number_format = "0.00"
    for c in range(1, len(headers) + 1):
        ws.cell(row=end, column=c).border = BORDER
    # highlight the shipped column
    for row_i in range(r, end + 1):
        cell = ws.cell(row=row_i, column=shipped_col + 1)
        if row_i > r:
            cell.fill = SHIP_FILL
    ws.cell(row=r, column=shipped_col + 1).fill = HEAD_FILL
    ws.cell(row=end + 2, column=1,
           value="Lower MASE is better. Green column = the model shipped in the app.").font = NOTE_FONT


def sheet_prophet_correction(wb, folds):
    ws = wb.create_sheet("Prophet Re-scoring")
    r = title_block(ws, "Why “8 of 12 categories, MASE 0.99” Did Not Hold Up",
                    ["The original result (docs/REBUILD_PIPELINE.md §6) used the same Prophet predictions but three test-design",
                     "advantages that don't exist in production: scoring only over days the store WAS open that test month (hindsight),",
                     "training on sale-days only (learns an inflated rate), and a MASE denominator measured in the wrong unit (~1.4x too small).",
                     "The predictions are identical in both columns below — only the yardstick changed."])
    g = folds.groupby("forecast_category")
    marco_scale = g["scale_marco"].mean().reindex(CATS)
    cal_scale = g["scale_cal"].mean().reindex(CATS)
    mae = g.apply(lambda x: (x["actual"] - x["prophet_marco"]).abs().mean(), include_groups=False).reindex(CATS)
    headers = ["Category", "Model's typical error (MAE, units)",
              "As originally scored — denominator (units)", "As originally scored — MASE",
              "Correct-unit denominator (units)", "Correct-unit MASE"]
    rows = []
    for cat in CATS:
        rows.append([cat, mae[cat], marco_scale[cat], mae[cat] / marco_scale[cat],
                    cal_scale[cat], mae[cat] / cal_scale[cat]])
    end = write_table(ws, r, headers, rows, widths=[20, 24, 24, 12, 22, 12],
                      number_formats={i: "0.00" for i in (1, 2, 3, 4, 5)})
    ws.cell(row=end, column=1, value="Macro average").font = BOLD_FONT
    for j, col in [(2, "B"), (4, "D"), (6, "F")]:
        ws.cell(row=end, column=j, value="=AVERAGE(%s%d:%s%d)" % (col, r + 1, col, end - 1))
        ws.cell(row=end, column=j).font = BOLD_FONT
        ws.cell(row=end, column=j).number_format = "0.00"
    ws.cell(row=end, column=4, value="=AVERAGE(D%d:D%d)" % (r + 1, end - 1))
    ws.cell(row=end, column=4).number_format = "0.00"
    ws.cell(row=end, column=6, value="=AVERAGE(F%d:F%d)" % (r + 1, end - 1))
    ws.cell(row=end, column=6).number_format = "0.00"
    for c in range(1, 7):
        ws.cell(row=end, column=c).font = BOLD_FONT
        ws.cell(row=end, column=c).border = BORDER
    n_lt1_orig = "=SUMPRODUCT(--(D%d:D%d<1))" % (r + 1, end - 1)
    n_lt1_fix = "=SUMPRODUCT(--(F%d:F%d<1))" % (r + 1, end - 1)
    ws.cell(row=end + 2, column=1, value="Categories with MASE < 1, as originally scored:").font = BODY_FONT
    ws.cell(row=end + 2, column=4, value=n_lt1_orig).font = BOLD_FONT
    ws.cell(row=end + 3, column=1, value="Categories with MASE < 1, correct units:").font = BODY_FONT
    ws.cell(row=end + 3, column=4, value=n_lt1_fix).font = BOLD_FONT


def sheet_intermittent(wb, bench):
    ws = wb.create_sheet("Intermittent Models (TSB)")
    r = title_block(ws, "Do TSB / SBA / Croston Beat the Average for Slow Sellers?",
                    ["Source: data/category_benchmark_folds.csv (scripts/benchmark_category_level.py) — marco's original 27-method run,",
                     "zero-filled full-calendar series. Different data span than the shipped model's own numbers, so absolute values here",
                     "differ slightly from the Shipped Model Results tab; the RELATIVE comparison between methods on this tab is fair.",
                     "TSB/SBA/Croston are built for “intermittent demand” (long quiet stretches, occasional sales) — raised as a candidate",
                     "specifically for the weaker-selling categories."])
    methods = ["RM6_6month_180d", "tsb", "sba", "croston"]
    labels = {"RM6_6month_180d": "6-month average (SHIPPED)", "tsb": "TSB", "sba": "SBA", "croston": "Croston"}
    tbl = {}
    for m in methods:
        g = bench[bench.method == m]
        tbl[m] = g.groupby("sku").apply(lambda x: x.abs_error.mean() / x.naive_scale.mean(),
                                        include_groups=False).reindex(CATS)
    headers = ["Category"] + [labels[m] for m in methods] + ["Best of these four"]
    rows = []
    for cat in CATS:
        vals = [tbl[m][cat] for m in methods]
        best = labels[methods[int(np.argmin(vals))]]
        rows.append([cat] + vals + [best])
    end = write_table(ws, r, headers, rows, widths=[20, 22, 10, 10, 10, 24],
                      number_formats={1: "0.00", 2: "0.00", 3: "0.00", 4: "0.00"})
    ws.cell(row=end, column=1, value="Macro average").font = BOLD_FONT
    for j, col in enumerate(["B", "C", "D", "E"], start=2):
        ws.cell(row=end, column=j, value="=AVERAGE(%s%d:%s%d)" % (col, r + 1, col, end - 1))
        ws.cell(row=end, column=j).font = BOLD_FONT
        ws.cell(row=end, column=j).number_format = "0.00"
    for c in range(1, 7):
        ws.cell(row=end, column=c).border = BORDER
    for row_i in range(r + 1, end):
        ws.cell(row=row_i, column=2).fill = SHIP_FILL
    ws.cell(row=r, column=2).fill = HEAD_FILL
    ws.cell(row=end + 2, column=1,
           value=("Result: TSB/SBA/Croston beat the average in only 2 of 12 categories (Plush & Souvenirs, "
                  "Umbrellas & Gear), and lose badly overall (macro MASE ~1.4 for the average vs. ~1.6–1.9 for "
                  "the others). “Low-selling” alone does not predict which model wins — Home & Novelty is "
                  "also low-selling and TSB loses there.")).font = NOTE_FONT
    ws.cell(row=end + 2, column=1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=end + 2, start_column=1, end_row=end + 2, end_column=6)
    ws.row_dimensions[end + 2].height = 45


# Raw method name -> (family, display label). Order here fixes column order:
# grouped by family so related methods sit together instead of appearing random.
METHOD_INFO = [
    ("naive",              "Naive",        "Repeat last value"),
    ("seasonal_naive_7",   "Naive",        "Seasonal naive (7d)"),
    ("rolling_mean_14",    "Trailing avg", "14-day average"),
    ("rolling_mean_30",    "Trailing avg", "30-day average"),
    ("rolling_mean_60",    "Trailing avg", "60-day average"),
    ("RM3_3month_90d",     "Trailing avg", "3-month average"),
    ("RM6_6month_180d",    "Trailing avg", "6-month average (SHIPPED)"),
    ("rolling_median_30",  "Trailing avg", "30-day median"),
    ("rolling_q75_30",     "Trailing avg", "30-day 75th percentile"),
    ("ewma_a0.1",          "Smoothing",    "EWMA (alpha=0.1)"),
    ("ewma_a0.3",          "Smoothing",    "EWMA (alpha=0.3)"),
    ("ets",                "Smoothing",    "ETS / Holt-Winters"),
    ("arima_111",          "ARIMA",        "ARIMA(1,1,1)"),
    ("arima_212",          "ARIMA",        "ARIMA(2,1,2)"),
    ("sarima_weekly",      "ARIMA",        "SARIMA (weekly)"),
    ("croston",            "Intermittent", "Croston"),
    ("sba",                "Intermittent", "SBA"),
    ("tsb",                "Intermittent", "TSB"),
    ("weekly_hurdle_12w",  "Hurdle",       "Weekly hurdle (12wk)"),
    ("prophet_plain",      "Prophet",      "Prophet (plain)"),
    ("prophet_cal",        "Prophet",      "Prophet (+calendar)"),
    ("extra_trees",        "ML",           "Extra Trees"),
    ("gradient_boosting",  "ML",           "Gradient Boosting"),
    ("lightgbm",           "ML",           "LightGBM"),
    ("random_forest",      "ML",           "Random Forest"),
    ("ridge",              "ML",           "Ridge"),
    ("xgboost",            "ML",           "XGBoost"),
]


def sheet_champion(wb, bench):
    ws = wb.create_sheet("Best Model per Category")
    r = title_block(ws, "Every Tested Model's MASE, per Category (27 Methods)",
                    ["Source: data/category_benchmark_folds.csv — trailing averages, smoothing, ARIMA, intermittent-demand,",
                     "hurdle, Prophet, and ML methods, all scored on the same test months per category.",
                     "Yellow = that category's best-scoring method. Green column = the model shipped in the app, for reference.",
                     "CAUTION: with 27 candidates and only ~12 test months per category, some method will look best in each category",
                     "by chance alone. The winners below are scattered across 8 unrelated method families with no consistent pattern,",
                     "which is the signature of that kind of noise, not a real effect — see the Overview tab. Shown in full for the",
                     "record, not as a basis for switching models per category."])
    tbl = {}
    for raw, _, _ in METHOD_INFO:
        g = bench[bench.method == raw]
        tbl[raw] = g.groupby("sku").apply(lambda x: x.abs_error.mean() / x.naive_scale.mean(),
                                          include_groups=False)
    full = pd.DataFrame(tbl).reindex(CATS)
    shipped_col_idx = [i for i, (raw, _, _) in enumerate(METHOD_INFO)
                       if raw == "RM6_6month_180d"][0] + 2   # +1 header, +1 category col

    headers = ["Category"] + [f"{fam}: {label}" if fam != label else label
                              for _, fam, label in METHOD_INFO]
    rows = [[cat] + [full.loc[cat, raw] for raw, _, _ in METHOD_INFO] for cat in CATS]
    end = write_table(ws, r, headers, rows, widths=[20] + [16] * len(METHOD_INFO),
                      number_formats={i: "0.00" for i in range(1, len(headers))})

    YELLOW = PatternFill("solid", fgColor="FFF2CC")
    for i, cat in enumerate(CATS):
        row_i = r + 1 + i
        best_j = int(np.argmin([full.loc[cat, raw] for raw, _, _ in METHOD_INFO])) + 2
        cell = ws.cell(row=row_i, column=best_j)
        cell.fill = YELLOW
        cell.font = BOLD_FONT
        # keep the shipped-model column visible even on a row where it also wins
        if best_j != shipped_col_idx:
            ws.cell(row=row_i, column=shipped_col_idx).fill = SHIP_FILL
        else:
            ws.cell(row=row_i, column=shipped_col_idx).fill = YELLOW

    ws.cell(row=r, column=shipped_col_idx).fill = HEAD_FILL  # keep header dark blue, not green

    ws.cell(row=end, column=1, value="Macro average").font = BOLD_FONT
    for j in range(2, len(headers) + 1):
        col = get_column_letter(j)
        ws.cell(row=end, column=j, value="=AVERAGE(%s%d:%s%d)" % (col, r + 1, col, end - 1))
        ws.cell(row=end, column=j).font = BOLD_FONT
        ws.cell(row=end, column=j).number_format = "0.00"
    for c in range(1, len(headers) + 1):
        ws.cell(row=end, column=c).border = BORDER

    ws.cell(row=end + 2, column=1,
           value="Best-model-per-category summary (for reference — see the caution note above):").font = NOTE_FONT
    ws.merge_cells(start_row=end + 2, start_column=1, end_row=end + 2, end_column=6)
    champ_headers = ["Category", "Best-scoring model", "Its MASE", "6-month average's MASE", "Gap"]
    champion = full.idxmin(axis=1)
    champion_label = {raw: (f"{fam}: {label}" if fam != label else label) for raw, fam, label in METHOD_INFO}
    rm6 = full["RM6_6month_180d"]
    champ_rows = [[cat, champion_label[champion[cat]], full.loc[cat].min(), rm6[cat],
                  rm6[cat] - full.loc[cat].min()] for cat in CATS]
    end2 = write_table(ws, end + 4, champ_headers, champ_rows, widths=[20, 24, 11, 20, 10],
                       number_formats={2: "0.00", 3: "0.00", 4: "0.00"})
    ws.cell(row=end2, column=1, value="Macro average").font = BOLD_FONT
    for j, col in [(3, "C"), (4, "D")]:
        ws.cell(row=end2, column=j, value="=AVERAGE(%s%d:%s%d)" % (col, end + 5, col, end2 - 1))
        ws.cell(row=end2, column=j).font = BOLD_FONT
        ws.cell(row=end2, column=j).number_format = "0.00"
    for c in range(1, 6):
        ws.cell(row=end2, column=c).border = BORDER
    ws.freeze_panes = ws.cell(row=r + 1, column=2).coordinate


def sheet_tuning(wb, tune):
    ws = wb.create_sheet("Parameter Tuning")
    r = title_block(ws, "Parameter Tuning — TSB and Prophet Settings",
                    ["Source: data/category_forecast_param_tuning.csv (scripts/tune_category_forecast_params.py).",
                     "Adjusts each model's OWN settings (not swapping the model) — 16 TSB alpha/beta combinations and 6 Prophet",
                     "trend-flexibility / seasonality combinations — scored on the shipped model's own category series and harness.",
                     "Tests whether the gap to the shipped 6-month average is a settings problem. Result: no — nothing beat it."])
    tune = tune.sort_values("macro_mase").reset_index(drop=True)
    headers = ["Configuration", "Macro-average MASE", "Categories with MASE < 1 (of 12)", "Categories beating ‘repeat last month’ (of 12)"]
    rows = tune[["method", "macro_mase", "below_1", "beats_naive_last30"]].values.tolist()
    shipped_idx = tune.index[tune["method"].str.contains("SHIPPED")].tolist()
    end = write_table(ws, r, headers, rows, widths=[32, 20, 26, 30],
                      highlight_rows=set(shipped_idx), number_formats={1: "0.00"})
    ws.cell(row=end + 2, column=1,
           value=("Best TSB setting: MASE 1.33 (vs. the 6-month average's 1.14). Best Prophet setting: MASE 1.22 (vs. the 6-month average's 1.14). "
                  "Neither beats the shipped model at any tested setting.")).font = NOTE_FONT
    ws.cell(row=end + 2, column=1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=end + 2, start_column=1, end_row=end + 2, end_column=4)
    ws.row_dimensions[end + 2].height = 30


def sheet_shape(wb, shape):
    ws = wb.create_sheet("Daily Shape Test")
    r = title_block(ws, "Day-by-Day Shape — Can Date-Aware Models Beat a Flat Line?",
                    ["Source: data/category_forecast_shape_test.csv (scripts/test_category_forecast_shape.py). 12 categories x 12 test months.",
                     "The 6-month average gives the validated 30-day TOTAL but draws a flat line. “6-month total x shape” keeps that total and lets a",
                     "date-aware model decide how it is spread across the days. “AS-IS” ships the date-aware model's own total and shape.",
                     "Daily error vs flat: below 1.00 = a better day-by-day pattern than a flat line. 30-day total error is the same for every hybrid by construction."])
    headers = ["Method", "Daily error vs flat line (all months)", "Category-months beating flat (%)",
               "Daily error vs flat, months with a semester break", "Daily error vs flat, other months",
               "30-day total error (WMAPE, %)"]
    rows = shape[["method", "daily_error_vs_flat", "folds_beating_flat_pct",
                  "daily_error_vs_flat_break_months", "daily_error_vs_flat_other_months",
                  "total_30d_wmape_pct"]].values.tolist()
    shipped = {i for i, row in enumerate(rows) if "Prophet: weekly + calendar (no yearly)" in row[0]
               and row[0].startswith("6-month total")}
    end = write_table(ws, r, headers, rows, widths=[62, 20, 20, 24, 20, 20],
                      highlight_rows=shipped,
                      number_formats={1: "0.000", 2: "0.0", 3: "0.000", 4: "0.000", 5: "0.0"})
    ws.cell(row=end + 2, column=1,
           value=("Shipped: “6-month total x shape: Prophet: weekly + calendar (no yearly)” (green row). Almost all of the gain is in months with a "
                  "semester break (4 of the 12 test windows); in ordinary months it is about a wash. One-off bulk orders cannot be predicted from dates. "
                  "Shipping Prophet AS-IS is worse than a flat line on both daily and total error.")).font = NOTE_FONT
    ws.cell(row=end + 2, column=1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=end + 2, start_column=1, end_row=end + 2, end_column=6)
    ws.row_dimensions[end + 2].height = 48


# ------------------------------------------------------------- item level
LEADER = "50/50: top-down + TSB"
OLD_PROPHET = "Prophet (as shipped)"


def write_block(ws, start_row, headers, rows, number_formats=None, fills=None):
    """A table inside a sheet that holds several. No column resize, no freeze."""
    fills = fills or {}
    for j, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=j, value=h)
    style_header(ws, start_row, len(headers))
    ws.row_dimensions[start_row].height = 45
    for i, row in enumerate(rows):
        for j, val in enumerate(row, start=1):
            cell = ws.cell(row=start_row + 1 + i, column=j, value=val)
            cell.font = BOLD_FONT if i in fills else BODY_FONT
            cell.border = BORDER
            if i in fills:
                cell.fill = fills[i]
            if number_formats and (j - 1) in number_formats:
                cell.number_format = number_formats[j - 1]
    return start_row + 1 + len(rows)


def sheet_item_methods(wb, t):
    ws = wb.create_sheet("Item Method Results")
    r = title_block(ws, "Item-Level Forecast — All 47 Methods on the 58 Fast Items",
                    ["Source: data/item_forecast_method_test.csv (scripts/test_item_forecast_methods.py). Same walk-forward test the app uses: up to 12 past 30-day periods per item,",
                     "each forecast using only data from before it. MASE below 1 = beat a simple benchmark; lower is better. Green = what the app uses now. Orange = the Prophet it replaced.",
                     "“Still selling” = the 45 items that sold anything in the last 360 days. The other 13 stopped selling (last sale May 2024 to May 2025), and a method that keeps forecasting them is penalised on every period.",
                     "Sorted by average MASE. 47 methods were compared, so the top group is within noise of each other; see Item Robustness."])
    t = t.sort_values("mean_MASE").reset_index(drop=True)
    headers = ["Method", "Type", "Average MASE (57 items)", "Typical (median) MASE", "Items with MASE below 1",
               "Error, % of units sold", "Over (+) / under (−) forecast, %", "Average MASE, 45 items still selling",
               "MASE, older 6 periods", "MASE, newer 6 periods", "Gap vs old Prophet (MASE; negative = better)",
               "95% range: low", "95% range: high", "Items better than old Prophet (of 57)"]
    rows = t[["method", "family", "mean_MASE", "median_MASE", "items_below_1", "pooled_WMAPE_pct", "bias_pct",
              "mean_MASE_still_selling", "mean_MASE_older_folds", "mean_MASE_newer_folds",
              "vs_prophet_mean_MASE_diff", "vs_prophet_ci_lo", "vs_prophet_ci_hi",
              "items_better_than_prophet"]].values.tolist()
    lead = t.index[t["method"] == LEADER].tolist()
    old = t.index[t["method"] == OLD_PROPHET].tolist()
    end = write_table(ws, r, headers, rows, widths=[54, 18, 13, 13, 13, 13, 14, 15, 12, 12, 17, 11, 11, 14],
                      highlight_rows=set(lead),
                      number_formats={2: "0.00", 3: "0.00", 5: "0.0", 6: "0.0", 7: "0.00", 8: "0.00", 9: "0.00",
                                      10: "0.00", 11: "0.00", 12: "0.00"})
    for i in old:
        for c in range(1, len(headers) + 1):
            ws.cell(row=r + 1 + i, column=c).fill = WARN_FILL
            ws.cell(row=r + 1 + i, column=c).font = BOLD_FONT
    ws.row_dimensions[r].height = 62
    ws.cell(row=end + 2, column=1,
            value=("The 95% range comes from re-sampling the items 4,000 times; a range that stays below zero means the method beat the old Prophet "
                   "however the items are drawn. Blends with Prophet in them scored worse than the same blend without it. Over/under is the average "
                   "size of the forecast against actual sales: the old Prophet over-forecast by 25%, the app's blend under-forecasts by 7%.")).font = NOTE_FONT
    ws.cell(row=end + 2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=end + 2, start_column=1, end_row=end + 2, end_column=10)
    ws.row_dimensions[end + 2].height = 48
    ws.freeze_panes = ws.cell(row=r + 1, column=2).coordinate


def sheet_item_robustness(wb, rob, sens, win):
    ws = wb.create_sheet("Item Robustness")
    r = title_block(ws, "Item-Level Forecast — Is the Winner Luck?",
                    ["Source: data/item_forecast_robustness.csv, item_forecast_sensitivity.csv, item_forecast_by_window.csv (scripts/test_item_forecast_methods.py).",
                     "Three checks on the 50/50 blend the app uses, because the best of 47 methods can look good by chance."])
    for col, w in zip("ABCDEFGHIJKL", [48, 46, 20, 14, 12, 12, 14, 16, 14, 14, 14, 14]):
        ws.column_dimensions[col].width = w

    ws.cell(row=r, column=1, value="1. Head-to-head, resampling the items (negative gap = the first method is better)").font = BOLD_FONT
    headers = ["Method", "Compared with", "Items", "Gap in average MASE", "95% range: low", "95% range: high",
               "Chance it is better", "Items where it is better", "Items compared"]
    rows = [[x.method, x.versus, x.cohort, x.mean_diff, x.ci_lo, x.ci_hi, x.p_better, x.items_better, x.n_items]
            for x in rob.itertuples()]
    end = write_block(ws, r + 1, headers, rows, number_formats={3: "0.00", 4: "0.00", 5: "0.00", 6: "0%"})
    ws.cell(row=end, column=1,
            value=("The blend beats the old Prophet however the items are drawn, on all items and on the 45 still selling. It beats TSB alone, but "
                   "not clearly the top-down part alone. The plain 6-month average does NOT reliably beat Prophet at item level, unlike category level.")).font = NOTE_FONT
    ws.cell(row=end, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=end, start_column=1, end_row=end, end_column=9)
    ws.row_dimensions[end].height = 32

    r2 = end + 2
    ws.cell(row=r2, column=1, value="2. Does the blend depend on its two settings? Average MASE, 57 items (a flat grid means no)").font = BOLD_FONT
    b = sens[sens.method.str.startswith("50/50")].copy()
    b["td"] = b.method.str.extract(r"top-down \((.*?)\)")[0]
    b["tsb"] = b.method.str.extract(r"\+ (TSB .*)$")[0]
    grid = b.pivot(index="td", columns="tsb", values="mean_MASE")
    order = ["level 180d, share 14d", "level 180d, share 30d", "level 180d, share 60d", "level 180d, share 90d",
             "level 90d, share 30d", "level 365d, share 30d"]
    grid = grid.reindex(order)
    headers = ["Top-down setting (category level window, item share window)"] + list(grid.columns)
    rows = [[k] + [float(v) for v in vals] for k, vals in zip(grid.index, grid.values)]
    lead_row = order.index("level 180d, share 30d")
    end = write_block(ws, r2 + 1, headers, rows, number_formats={1: "0.00", 2: "0.00", 3: "0.00"},
                      fills={lead_row: SHIP_FILL})
    ws.cell(row=end, column=1,
            value=("Green = the setting the app uses (180-day category level, 30-day share, TSB 0.05). All 18 combinations land between "
                   "1.68 and 1.90, against 2.52 for the old Prophet. The 365-day level scores a little better but under-forecasts by 12 to 14%.")).font = NOTE_FONT
    ws.cell(row=end, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=end, start_column=1, end_row=end, end_column=5)
    ws.row_dimensions[end].height = 46

    r3 = end + 2
    ws.cell(row=r3, column=1, value="3. Error and bias in each past 30-day window (the same window one year ago is 2025-07-14 to 2025-08-12)").font = BOLD_FONT
    headers = ["Window start", "Window end", "Enrollment days", "Semester-break days", "Exam-week days", "Event days",
               "Units sold", "Old Prophet: error %", "Old Prophet: over/under %", "Blend: error %", "Blend: over/under %",
               "Repeat last month: error %"]
    rows = [[str(x.start), str(x.end), x.enrollment_days, x.break_days, x.exam_days, x.event_days, x.actual_units,
             x.prophet_WMAPE_pct, x.prophet_bias_pct, x.blend_WMAPE_pct, x.blend_bias_pct, x.repeat_last_month_WMAPE_pct]
            for x in win.itertuples()]
    end = write_block(ws, r3 + 1, headers, rows,
                      number_formats={6: "#,##0", 7: "0", 8: "0", 9: "0", 10: "0", 11: "0"})
    ws.cell(row=end, column=1,
            value=("The blend has a lower error than the old Prophet in all 12 windows. It has no seasonal term, so it under-forecasts busy windows "
                   "(about 40% low in the window a year ago that matches the current forecast window) and over-forecasts quiet ones "
                   "(December holidays, February to April). Only 3 windows contain enrollment days, so enrollment effects cannot be tested.")).font = NOTE_FONT
    ws.cell(row=end, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=end, start_column=1, end_row=end, end_column=10)
    ws.row_dimensions[end].height = 46


def sheet_item_shape(wb, shp):
    ws = wb.create_sheet("Item Daily Shape")
    r = title_block(ws, "Item-Level Day-by-Day Shape — Which Pattern Beats a Flat Line?",
                    ["Source: data/item_forecast_shape_test.csv (scripts/test_item_forecast_methods.py). The item's 30-day total (the 50/50 blend) is fixed; only how it is spread across the 30 days changes.",
                     "Daily error vs flat line: below 1.00 = a better day-by-day pattern than a flat line. An item's own sales are mostly zero days, so its own pattern is noisy; its category has enough sales to show a rhythm.",
                     "The 30-day total (and so the accuracy on the Item Method Results tab) is identical for every pattern."])
    g = shp[shp.cohort == "all Fast items"].reset_index(drop=True)
    headers = ["Pattern", "Daily error vs flat line (all months)", "Item-months beating flat (%)",
               "Daily error vs flat, months with a semester break", "Daily error vs flat, other months"]
    rows = g[["shape", "daily_error_vs_flat", "item_folds_beating_flat_pct", "break_months", "other_months"]].values.tolist()
    lead = g.index[g["shape"] == "Category's Prophet pattern"].tolist()
    end = write_table(ws, r, headers, rows, widths=[38, 20, 20, 24, 20], highlight_rows=set(lead),
                      number_formats={1: "0.000", 2: "0.0", 3: "0.000", 4: "0.000"})
    ws.cell(row=end + 2, column=1,
            value=("Green = what the app uses: each item takes its category's Prophet pattern (weekly rhythm, store closures, school calendar). "
                   "The gain is small (about 2% of daily error, about 6% in semester-break months) and is a rhythm, not a forecast of individual "
                   "spikes or bulk orders. The results are the same on the 45 items still selling.")).font = NOTE_FONT
    ws.cell(row=end + 2, column=1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=end + 2, start_column=1, end_row=end + 2, end_column=5)
    ws.row_dimensions[end + 2].height = 48


def sheet_lead_time_gap(wb, df):
    ws = wb.create_sheet("Lead-Time Validation Check")
    r = title_block(ws, "Disclosure Check — Validating With the Real Reorder Delay",
                    ["Source: data/lead_time_gap_test.csv (scripts/test_lead_time_gap.py). NOT the headline number the app and the other tabs report — a disclosure check, kept",
                     "alongside them in case it is asked about. Every other tab tests a forecast as if it could be acted on the very next day. A real reorder cannot: it does not restock the",
                     "shelf for lead_time_days (14, 18 or 28 days here). This re-tests the SAME two shipped models on the SAME past periods, but only using the sales data that would",
                     "genuinely have been known that many days earlier — the honest test of what a forecast used for reordering can actually achieve.",
                     "A bigger gap forecasts further ahead from less data, so a worse score is the expected cost of this check, not a fault. Green = the current (no-gap) numbers already shown elsewhere."])
    headers = ["Level", "Test", "Items / categories tested", "Average MASE", "Typical (median) MASE",
               "Below 1 (better than benchmark)", "Beats “repeat last 30 days”", "Error, % of units sold"]
    rows, fills = [], {}
    for level, label_gap0 in (("item", "current (gap=0)"), ("category", "current (gap=0)")):
        g = df[df.level == level]
        for label in (label_gap0, "with lead-time gap"):
            sub = g[g.gap_label == label]
            n = sub["name"].nunique()
            below1 = int((sub.mase < 1).sum())
            beats = int(sub.beats_naive.sum())
            wmape = 100 * (sub.mae * sub.n_folds).sum() / (sub.mean_actual * sub.n_folds).sum()
            test_name = "Current test (forecast usable the next day)" if label.startswith("current")                 else "With the real reorder delay (14-28 days, per item / category)"
            if label.startswith("current"):
                fills[len(rows)] = SHIP_FILL
            rows.append([level.capitalize(), test_name, n, sub.mase.mean(), sub.mase.median(),
                        f"{below1} of {n}", f"{beats} of {n}", wmape])
    end = write_table(ws, r, headers, rows, widths=[11, 46, 20, 14, 16, 24, 22, 18],
                      highlight_rows=set(fills), number_formats={3: "0.00", 4: "0.00", 7: "0.0"})
    ws.cell(row=end + 2, column=1,
            value=("Measured on the models before the calendar adjustment was added. "
                   "Item forecasts get meaningfully less accurate under the real delay (average MASE 1.71 to 2.14) — expected, since the model has less "
                   "data and is forecasting further out. Category forecasts do not clearly get worse (1.14 to 1.03), but it is a mixed result, not a clean win: "
                   "2 of the 12 categories (Lanyards & IDs, Shirts & Tops) flip from beating “repeat last 30 days” to not, because that benchmark is just as "
                   "sensitive to the delay as the model is. This check does not change what the app forecasts or displays — only how the accuracy claim is tested.")
           ).font = NOTE_FONT
    ws.cell(row=end + 2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=end + 2, start_column=1, end_row=end + 2, end_column=8)
    ws.row_dimensions[end + 2].height = 60


def sheet_calendar_adjustment(wb, res, win):
    ws = wb.create_sheet("Calendar Adjustment")
    r = title_block(ws, "Calendar Adjustment — Lowering the Total When Quieter Days Are Coming",
                    ["Source: data/calendar_adjustment_test.csv and calendar_adjustment_by_window.csv (scripts/test_calendar_adjustment.py).",
                     "A trailing average repeats one daily rate, so a month full of semester-break days got the same total as an ordinary month.",
                     "The school calendar (published in advance) says when breaks and exams are coming. “Lowered” scales the total down to match and",
                     "never raises it (green = what the app uses). “Up or down” is the version allowed to go both ways, kept for comparison.",
                     "Caution: the “only down” rule was chosen after seeing the up-or-down version fail after breaks, so the gain may be a little optimistic."])
    label = {"base": "Before (no adjustment)", "uncapped": "Up or down", "capped": "Lowered only (shipped)"}
    headers = ["Level", "Test months", "Version", "Average MASE", "MASE below 1",
               "Error, % of units sold", "Over (+) / under (−) forecast, %"]
    rows, fills = [], {}
    for x in res.itertuples():
        if x.version == "capped":
            fills[len(rows)] = SHIP_FILL
        rows.append([x.level.capitalize(), x.folds, label[x.version], x.mean_MASE,
                     f"{int(x.n_below_1)} of {int(x.n)}", x.pooled_WMAPE_pct, x.bias_pct])
    end = write_block(ws, r, headers, rows, number_formats={3: "0.00", 5: "0.0", 6: "0.0"}, fills=fills)
    for col, w in zip("ABCDEFG", [11, 14, 24, 13, 13, 14, 16]):
        ws.column_dimensions[col].width = w

    cap = res[(res.folds == "all folds") & (res.version == "capped")]
    txt = []
    for x in cap.itertuples():
        txt.append(f"{x.level}: better for {int(x.n_better)} of {int(x.n)}, gap {x.gap_vs_base:+.3f} MASE "
                   f"(95% range {x.gap_ci_lo:+.3f} to {x.gap_ci_hi:+.3f})")
    ws.cell(row=end + 1, column=1,
            value=("Lowered vs before, resampling the categories / items: " + "; ".join(txt) + ". "
                   "Better in both the older and the newer test months at both levels. The cost is more under-forecasting (about -7% to -15%), "
                   "mostly in break months, when the store is quiet.")).font = NOTE_FONT
    ws.cell(row=end + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=end + 1, start_column=1, end_row=end + 1, end_column=7)
    ws.row_dimensions[end + 1].height = 48

    r2 = end + 3
    ws.cell(row=r2, column=1, value="Error in each past 30-day window (%, lower is better)").font = BOLD_FONT
    headers = ["Level", "Window start", "Before", "Up or down", "Lowered only", "Lowered only: over/under %"]
    rows = [[x.level.capitalize(), str(x.window_start), x.base_error_pct, x.uncapped_error_pct,
             x.capped_error_pct, x.capped_bias_pct] for x in win.itertuples()]
    end2 = write_block(ws, r2 + 1, headers, rows, number_formats={2: "0", 3: "0", 4: "0", 5: "0"})
    ws.cell(row=end2 + 1, column=1,
            value=("The biggest gain is the December window (21 semester-break days). The current forecast window (Jul 9 to Aug 7, 2026) has 7 break "
                   "and 5 exam days, so every category is lowered 4-15%. In the same season last year Shirts & Tops and Lanyards & IDs were already "
                   "under-forecast by 50-68% before any adjustment (likely the start of the school year, which the calendar does not flag), so those "
                   "two may read low this month.")).font = NOTE_FONT
    ws.cell(row=end2 + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=end2 + 1, start_column=1, end_row=end2 + 1, end_column=7)
    ws.row_dimensions[end2 + 1].height = 60
    ws.freeze_panes = None


def sheet_transformations(wb, res, chk):
    ws = wb.create_sheet("Transformations")
    r = title_block(ws, "Transformations — Do Log, Square-Root or Yeo-Johnson Help?",
                    ["Source: data/transformation_test.csv and transformation_checks.csv (scripts/test_transformations.py). Same 12 past 30-day periods as every other tab.",
                     "Each method is fitted on transformed daily sales, then its forecast is turned back into units: “naive” undoes the transform directly;",
                     "“smeared” uses the standard correction (average of the undone prediction plus each past error), because undoing an average is not the same as averaging.",
                     "Green = the model the app uses. Orange = Yeo-Johnson. Gap vs shipped: negative = better than the app's model; a 95% range that crosses 0 = no real difference."])
    headers = ["Level", "Method", "Average MASE", "MASE below 1", "Error, % of units sold",
               "Over (+) / under (−) forecast, %", "Gap vs shipped (MASE)", "95% range: low", "95% range: high"]
    rows, fills = [], {}
    for level in ("category", "item"):
        g = res[res.level == level].sort_values("macro_MASE")
        for x in g.itertuples():
            if x.method.startswith("REF shipped"):
                fills[len(rows)] = SHIP_FILL
            elif ", yj" in x.method or " yj " in x.method:
                fills[len(rows)] = WARN_FILL
            is_ship = x.method.startswith("REF shipped")
            rows.append([level.capitalize(), x.method.replace("REF ", "Reference: "), x.macro_MASE,
                         f"{int(x.below_1)} of {int(x.n)}", x.pooled_WMAPE_pct, x.bias_pct,
                         None if is_ship else x.vs_shipped_gap,
                         None if is_ship else x.vs_shipped_ci_lo,
                         None if is_ship else x.vs_shipped_ci_hi])
    end = write_block(ws, r, headers, rows,
                      number_formats={2: "0.00", 4: "0.0", 5: "0.0", 6: "+0.00;-0.00", 7: "+0.00;-0.00", 8: "+0.00;-0.00"},
                      fills=fills)
    for col, w in zip("ABCDEFGHI", [10, 52, 12, 12, 13, 15, 14, 12, 12]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=r + 1, column=3).coordinate

    cat = res[res.level == "category"].set_index("method")
    item = res[res.level == "item"].set_index("method")
    sp = cat.loc["Prophet, sqrt, smeared"]
    ship_c = cat.loc["REF shipped (6-month + calendar)"]
    ship_i = item.loc["REF shipped (blend + calendar)"]
    best_i = item[~item.is_reference.astype(bool)].sort_values("macro_MASE").iloc[0]
    c = chk.iloc[0]
    notes = [
        ("Averages cannot be helped by a transform.",
         f"Undone with the correction, every transformed 6-month average gives back the plain average exactly "
         f"(MASE {cat.loc['6-month avg, log1p, smeared back-transform', 'macro_MASE']:.2f}, same as the untransformed "
         f"{cat.loc['REF 6-month average (unadjusted)', 'macro_MASE']:.2f}). Undone the naive way, it predicts the typical day "
         f"rather than the average day and under-forecasts by 59-78%."),
        ("Square-root Prophet ties the app's category model, it does not beat it.",
         f"MASE {sp.macro_MASE:.2f} against {ship_c.macro_MASE:.2f}; gap {sp.vs_shipped_gap:+.3f} "
         f"(95% range {sp.vs_shipped_ci_lo:+.3f} to {sp.vs_shipped_ci_hi:+.3f}), better in {int(sp.vs_shipped_n_better)} of 12 categories. "
         f"It under-forecasts much less ({sp.bias_pct:+.0f}% against {ship_c.bias_pct:+.0f}%), so it is the version of Prophet to use "
         f"if Prophet is required. It is also the best of about 27 variants, so its score is probably slightly flattering."),
        ("At item level nothing comes close.",
         f"The best transformed option ({best_i.name}) scores {best_i.macro_MASE:.2f} against the app's {ship_i.macro_MASE:.2f}; "
         f"transforms improve item Prophet (2.52 to about 2.2) but not enough."),
        ("Yeo-Johnson is the wrong tool for this data.",
         f"{c.category_zero_day_pct:.0f}% of category-days are zero. Fitting Yeo-Johnson to make that look like a bell curve gives "
         f"extreme negative settings (lambda {c.yj_lambda_min:.1f} to {c.yj_lambda_max:.1f}, median {c.yj_lambda_median:.1f}, over "
         f"{int(c.n_category_slices)} training periods), and undoing them either wipes out the forecast or makes it explode "
         f"(the orange rows)."),
    ]
    rr = end + 1
    for head, body in notes:
        ws.cell(row=rr, column=1, value=head).font = BOLD_FONT
        ws.cell(row=rr + 1, column=1, value=body).font = NOTE_FONT
        ws.cell(row=rr + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=rr + 1, start_column=1, end_row=rr + 1, end_column=9)
        ws.row_dimensions[rr + 1].height = 46
        rr += 3


def main():
    folds = pd.read_csv(os.path.join(DATA, "category_forecast_method_folds.csv"))
    bench = pd.read_csv(os.path.join(DATA, "category_benchmark_folds.csv"))
    tune = pd.read_csv(os.path.join(DATA, "category_forecast_param_tuning.csv"))
    shape = pd.read_csv(os.path.join(DATA, "category_forecast_shape_test.csv"))
    item_t = pd.read_csv(os.path.join(DATA, "item_forecast_method_test.csv"))
    item_rob = pd.read_csv(os.path.join(DATA, "item_forecast_robustness.csv"))
    item_sens = pd.read_csv(os.path.join(DATA, "item_forecast_sensitivity.csv"))
    item_win = pd.read_csv(os.path.join(DATA, "item_forecast_by_window.csv"))
    item_shape = pd.read_csv(os.path.join(DATA, "item_forecast_shape_test.csv"))
    lead_gap = pd.read_csv(os.path.join(DATA, "lead_time_gap_test.csv"))
    cal_res = pd.read_csv(os.path.join(DATA, "calendar_adjustment_test.csv"))
    cal_win = pd.read_csv(os.path.join(DATA, "calendar_adjustment_by_window.csv"))
    tf_res = pd.read_csv(os.path.join(DATA, "transformation_test.csv"))
    tf_chk = pd.read_csv(os.path.join(DATA, "transformation_checks.csv"))
    shipped = load_shipped()

    wb = Workbook()
    sheet_overview(wb)
    sheet_shipped(wb, shipped)
    sheet_method_comparison(wb, folds)
    sheet_prophet_correction(wb, folds)
    sheet_intermittent(wb, bench)
    sheet_champion(wb, bench)
    sheet_tuning(wb, tune)
    sheet_shape(wb, shape)
    sheet_item_methods(wb, item_t)
    sheet_item_robustness(wb, item_rob, item_sens, item_win)
    sheet_item_shape(wb, item_shape)
    sheet_calendar_adjustment(wb, cal_res, cal_win)
    sheet_transformations(wb, tf_res, tf_chk)
    sheet_lead_time_gap(wb, lead_gap)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False

    os.makedirs(DATA, exist_ok=True)
    wb.save(OUT)
    print("Wrote", OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
