"""
scripts/build_current_results_workbook.py
------------------------------------------------------------------
Builds data/USTore_Current_Model_Results.xlsx: a short workbook showing only
the models the app uses right now and how accurate they are, for the
adviser. (The long list of every method tried is in
USTore_Forecast_Testing_Summary.xlsx.)

  Read Me            in plain words: what is forecast, how accurate it is, how it was
                     checked (and why the scores are slightly flattering), what each word means
  How It Works       the two models side by side, and the biggest item's forecast for the
                     last past window worked through step by step (checked against the model)
  Summary            one row per level (category, all items, items still selling)
  Categories         MASE, MAPE, error %, bias etc. for each of the 12 categories
  Items              the same for each of the 58 forecast items
  Category by Month  forecast vs actual for every category and past 30-day window
  Item by Month      forecast vs actual for every item and past 30-day window
  CICS vs Current    the store's own item list ("FOR CICS STUDENTS.xlsx") as the
                     grouping, against the current 12 categories
  CICS Per Item      that comparison per Fast and Slow item
  CICS Mapping       which CICS section each product was placed in

  More History Test  accuracy when the models see only the last 3-18 months
  What-If Scenarios  error if bulk orders / closures were known in advance
  Calendar Check     the school-calendar rule at 30 different window start dates
  Grouping Search    whether fewer / re-mixed categories forecast better

The three CICS tabs are read from data/cics_*.csv - run
scripts/compare_cics_categorization.py first. More History Test is read from
data/history_length_*.csv - run scripts/test_history_length.py first, and
What-If Scenarios from data/what_if_test.csv - run scripts/test_what_if.py,
Calendar Check from data/calendar_adjustment_starts.csv - run
scripts/test_calendar_adjustment_starts.py, and Grouping Search from
data/grouping_search_*.csv - run scripts/search_groupings.py.

Scores are recomputed with the production code (the same walk-forward test
step4 / step4c use, via scripts/test_calendar_adjustment.py) and checked
against the MAE / MASE stored in ustore.db; the script stops if they differ.
Run step4 and step4c first.

Run:  python scripts/build_current_results_workbook.py
------------------------------------------------------------------
"""
import os
import sqlite3
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))
import test_calendar_adjustment as tca
from forecasting.calendar_adjust import calendar_multiplier, load_day_types
from forecasting.evaluate import make_folds
from forecasting.intermittent import tsb
from forecasting.topdown import LEVEL_WINDOW, SHARE_WINDOW, TSB_ALPHA, TSB_BETA

DB_PATH = os.path.join(ROOT, "ustore.db")
OUT = os.path.join(ROOT, "data", "USTore_Current_Model_Results.xlsx")
DISCONTINUED_DAYS = 365   # same rule as backend/app.py

CAT_MODEL = "6-month average per category, calendar-adjusted"
ITEM_MODEL = "Category share + recent-sales blend, calendar-adjusted"
CAT_MODEL_ID = "RM6_6month_180d+calendar+prophet_shape"
ITEM_MODEL_ID = "topdown_tsb+calendar+prophet_shape"

FONT = "Calibri"
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT, size=14, bold=True, color="1F4E78")
NOTE_FONT = Font(name=FONT, size=10, italic=True, color="595959")
BODY_FONT = Font(name=FONT, size=11)
BOLD_FONT = Font(name=FONT, size=11, bold=True)
GOOD_FILL = PatternFill("solid", fgColor="E2EFDA")
BAD_FILL = PatternFill("solid", fgColor="FCE4D6")
GREY_FILL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------- scoring
def score(df):
    """Per-key accuracy of the shipped ("capped") version over the past windows."""
    out = []
    for key, g in df.groupby("key"):
        err = g.pred - g.actual
        sc = g.scale[np.isfinite(g.scale) & (g.scale > 0)]
        nz = g.actual > 0
        out.append(dict(
            key=key, windows=len(g),
            avg_actual=g.actual.mean(), avg_forecast=g.pred.mean(),
            mae=err.abs().mean(), rmse=np.sqrt((err ** 2).mean()),
            mase=err.abs().mean() / sc.mean() if len(sc) else np.nan,
            mape=100 * (err[nz].abs() / g.actual[nz]).mean() if nz.any() else np.nan,
            wmape=100 * err.abs().sum() / g.actual.sum() if g.actual.sum() else np.nan,
            bias=100 * err.sum() / g.actual.sum() if g.actual.sum() else np.nan))
    return pd.DataFrame(out).set_index("key")


MAPE_TARGET = 20.0          # the manuscript's acceptance criterion (section 3.3.4)

# Summary columns, plain-language first: (key, header, number format).
SUMMARY_COLS = [
    ("level", "Level", None),
    ("model", "Model in use", None),
    ("n", "How many", None),
    ("wmape", "How far off, % of units sold (WMAPE)", "0.0"),
    ("bias", "Too high (+) / too low (-), %", "+0.0;-0.0"),
    ("beats", "Better than 'same as last 30 days'", None),
    ("target", "Meet the paper's target (MAPE 20% or less)", None),
    ("mape_med", "Typical % miss per window (median MAPE)", "0.0"),
    ("mape_avg", "Average % miss per window (MAPE)", "0.0"),
    ("mase_avg", "Average MASE", "0.00"),
    ("mase_med", "Median MASE", "0.00"),
    ("mase_below", "MASE below 1", None),
]


def summary_row(level, model, s, df):
    """One Summary line: `s` = per-category or per-item scores, `df` = their past windows."""
    err = df.pred - df.actual
    return dict(level=level, model=model, n=len(s),
                wmape=100 * err.abs().sum() / df.actual.sum(), bias=100 * err.sum() / df.actual.sum(),
                beats=f"{int(s.beats_naive.sum())} of {s.beats_naive.notna().sum()}",
                target=f"{int((s.mape <= MAPE_TARGET).sum())} of {len(s)}",
                mape_med=s.mape.median(), mape_avg=s.mape.mean(),
                mase_avg=s.mase.mean(), mase_med=s.mase.median(),
                mase_below=f"{int((s.mase < 1).sum())} of {s.mase.notna().sum()}")


def wmape(df):
    return 100 * (df.pred - df.actual).abs().sum() / df.actual.sum()


def check_against_db(s, stored, label):
    j = s[["mae", "mase"]].join(stored[["mae", "mase"]], rsuffix="_db", how="outer")
    def differs(a, b):   # both blank (e.g. MASE of an item with no sales in any window) counts as equal
        return ((a - b).abs() > 1e-6) | (a.isna() != b.isna())
    bad = j[differs(j.mae, j.mae_db) | differs(j.mase, j.mase_db)]
    if len(bad):
        print(f"{label}: recomputed scores differ from ustore.db for {len(bad)} rows:")
        print(bad)
        raise SystemExit(1)
    print(f"{label}: {len(j)} rows match ustore.db (MAE and MASE)")


# ---------------------------------------------------------------- styling
def title(ws, text, notes):
    ws.cell(row=1, column=1, value=text).font = TITLE_FONT
    for i, n in enumerate(notes):
        ws.cell(row=2 + i, column=1, value=n).font = NOTE_FONT
    return 3 + len(notes)


def table(ws, r0, headers, rows, formats, widths, fills=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 45
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if isinstance(v, float) and not np.isfinite(v):
                v = None
            c = ws.cell(row=r0 + 1 + i, column=j + 1, value=v)
            c.font, c.border = BODY_FONT, BORDER
            if j in formats:
                c.number_format = formats[j]
            if fills and (i, j) in fills:
                c.fill = fills[(i, j)]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=r0 + 1, column=2).coordinate
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    return r0 + len(rows)


def mase_fills(rows, col, fills=None):
    fills = fills if fills is not None else {}
    for i, row in enumerate(rows):
        v = row[col]
        if v is not None and np.isfinite(v):
            fills[(i, col)] = GOOD_FILL if v < 1 else BAD_FILL if v > 2 else None
    return {k: v for k, v in fills.items() if v is not None}


# ---------------------------------------------------------------- sheets
GLOSSARY = [
    ("How far off (WMAPE)", "All units the forecast missed by, divided by all units sold. The main accuracy number in "
                            "this workbook. Big sellers count more, so it reflects stock and money."),
    ("Too high / too low", "All units forecast minus all units sold, as a % of units sold. Negative = the forecast was "
                           "too low overall."),
    ("Better than 'same as last 30 days'", "The forecast's average miss, in units, was smaller than simply repeating "
                                           "the last 30 days' sales."),
    ("MAPE", "The % miss in each 30-day window, averaged. Windows with no sales are skipped. Very high for items that "
             "sell only a few units (forecasting 6 when 1 sold is a 500% miss). The paper's 20% target uses this measure."),
    ("MASE", "The forecast's average miss divided by how much the item's (or category's) sales usually changed from "
             "one 30-day block to the next before the forecast was made. Below 1 = the forecast missed by less than "
             "sales usually change. Green below 1, orange above 2. It is not the same test as 'Better than same as "
             "last 30 days', so the two can disagree."),
    ("Average / median", "Average = every category or item counts the same, so a few extreme ones can pull it. "
                         "Median = the middle one when sorted: the typical category or item."),
]


def sheet_readme(wb, f):
    """`f` holds every number the text quotes, computed in main()."""
    ws = wb.active
    ws.title = "Read Me"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    ws.cell(row=1, column=1, value="USTore Forecasting: What the App Forecasts and How Accurate It Is").font = TITLE_FONT
    ws.cell(row=2, column=1, value=f"Data up to {f['snapshot']}. The current forecast covers {f['fc_start']} to "
                                   f"{f['fc_end']}.").font = NOTE_FONT
    cat, item, live = f["cat"], f["item"], f["live"]
    lines = [
        ("IN SHORT", None),
        ("What it does", f"The app forecasts the next 30 days of sales for each of the {cat['n']} product categories and "
                         f"for each of the {item['n']} items classed as fast-moving (the top 20% by average daily sales "
                         f"over all their history). The other {f['n_other']} items (slow or non-moving) are not forecast."),
        ("How accurate", f"Checked on 12 past months: category forecasts were off by {cat['wmape']:.0f}% of units sold, and "
                         f"item forecasts by {item['wmape']:.0f}%. Both came in too low overall: categories by "
                         f"{abs(cat['bias']):.0f}%, items by {abs(item['bias']):.0f}%."),
        ("Why the error is high", f"Sales come in bursts: bulk orders and school-calendar rushes. Logging bulk orders in "
                                  f"advance would cut the item error from {item['wmape']:.0f}% to about "
                                  f"{f['bulk_lo']:.0f}-{f['bulk_hi']:.0f}% (What-If tab). Giving the forecast more past "
                                  "sales would not lower it (More History tab)."),
        ("The paper's target", f"The paper's target is a MAPE of {MAPE_TARGET:.0f}% or less. {cat['target']} categories and "
                               f"{item['target']} items meet it."),
        ("", None),
        ("HOW THE FORECAST WORKS", None),
        ("Per category", "The category's average daily sales over the last 6 months, times 30. "
                         "(Technical name: 6-month moving average.)"),
        ("Per item", "Half from the item's share of its category in the last 30 days, times the category's 6-month "
                     "level. Half from how often the item has sold lately and how many it sells each time; this part "
                     "fades toward zero if the item stops selling. (Technical name: category share + TSB blend.)"),
        ("School calendar", "Both are lowered when more semester-break or exam days are coming than in the recent past. "
                            "They are never raised."),
        ("Day by day", "The 30-day total is spread over the days using the category's usual weekly pattern (from "
                       "Prophet). This does not change the total."),
        ("Worked example", f"The How It Works tab follows {f['ex_item']} through every step with real numbers."),
        ("", None),
        ("HOW IT WAS CHECKED", None),
        ("Walk-forward test", f"Pretend it is an earlier date, forecast the next 30 days using only sales before it, then "
                              f"compare with what really sold. Repeated for 12 past 30-day windows ({f['first_win']} to "
                              f"{f['last_win']})."),
        ("Fair warning", f"The methods were picked by comparing many options on nearly the same past months (47 for "
                         f"items), and the calendar rule was set after seeing two of them, {f['cal_effect']}. "
                         f"Without the calendar rule, items are off by {f['item_base']:.1f}% instead of "
                         f"{item['wmape']:.1f}%, and categories by {f['cat_base']:.1f}% instead of {cat['wmape']:.1f}%."),
        ("Calendar rule, fairly", "Which 12 windows are tested decides whether they straddle a break or a rush, so one "
                                  "set of windows can make the calendar rule look good or bad by chance. Re-run with the "
                                  f"windows starting on {f['starts']['item']['n']} different days, it lowers the error on "
                                  f"average: items {f['starts']['item']['base']:.1f}% to {f['starts']['item']['capped']:.1f}% "
                                  f"(better at {f['starts']['item']['helped']} of {f['starts']['item']['n']}), categories "
                                  f"{f['starts']['category']['base']:.1f}% to {f['starts']['category']['capped']:.1f}% "
                                  f"(better at {f['starts']['category']['helped']} of {f['starts']['category']['n']}). "
                                  "Quote this average, not one set of windows (Calendar Check tab)."),
        ("Items that stopped selling", f"{item['n'] - live['n']} of the {item['n']} items have not sold in a year. They "
                                       f"are forecast near zero and sold nothing, so their MASE looks perfect. Average "
                                       f"item MASE is {item['mase_avg']:.2f} with them and {live['mase_avg']:.2f} for the "
                                       f"{live['n']} still selling; the Summary tab shows both."),
        ("", None),
        ("WHAT THE WORDS MEAN", None),
    ] + GLOSSARY + [
        ("", None),
        ("OTHER TABS", None),
        ("Summary, Categories, Items", "The scores: overall, per category, per item."),
        ("By Month tabs", "Forecast against actual sales in every past 30-day window."),
        ("CICS tabs", "The same models re-run with the store's own item list (FOR CICS STUDENTS.xlsx) as the "
                      "categories, to see whether it forecasts better. Slow items are included there for this "
                      "comparison only."),
        ("More History Test", "Whether more past sales make the forecasts better: the same models re-run seeing only "
                              "the last 3 to 18 months. Real data only."),
        ("What-If Scenarios", "How low the error could go if the store logged bulk orders or closures in advance. "
                              "Not what the app does."),
        ("Calendar Check", "The school-calendar rule tested with the 12 windows starting on 30 different days, with and "
                           "without the rule."),
        ("Grouping Search", "Whether fewer or re-mixed categories forecast better, chosen on older months and scored on "
                            "newer ones the search never saw."),
        ("All methods tried", "USTore_Forecast_Testing_Summary.xlsx."),
    ]
    r = 4
    for a, b in lines:
        ca = ws.cell(row=r, column=1, value=a)
        if b is None:
            ca.font = Font(name=FONT, size=11, bold=True, color="1F4E78")
        else:
            ca.font = BOLD_FONT
            cb = ws.cell(row=r, column=2, value=b)
            cb.font = BODY_FONT
            cb.alignment = Alignment(wrap_text=True, vertical="top")
            ca.alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[r].height = 15 * max(1, -(-len(b) // 105))
        r += 1


def sheet_summary(wb, summ, n_dead):
    ws = wb.create_sheet("Summary")
    r = title(ws, "Summary: How Accurate the Forecasts Were", [
        "Checked on 12 past 30-day windows. Read the 'How far off' column first; the rest is detail.",
        f"'Items still selling' leaves out the {n_dead} items with no sale in a year: they are forecast near zero, "
        "sold nothing, and so make MASE look better than it is for the items that matter."])
    headers = [h for _, h, _ in SUMMARY_COLS]
    rows = [[s[k] for k, _, _ in SUMMARY_COLS] for s in summ]
    formats = {j: fmt for j, (_, _, fmt) in enumerate(SUMMARY_COLS) if fmt}
    r = table(ws, r, headers, rows, formats, [26, 44, 9, 13, 12, 13, 13, 12, 12, 10, 10, 10])
    ws.freeze_panes = None
    ws.auto_filter.ref = None

    r += 2
    ws.cell(row=r, column=1, value="WHAT THE WORDS MEAN").font = Font(name=FONT, size=11, bold=True, color="1F4E78")
    for term, text in GLOSSARY:
        r += 1
        ca = ws.cell(row=r, column=1, value=term)
        ca.font, ca.alignment = BOLD_FONT, Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=len(headers))
        cb = ws.cell(row=r, column=2, value=text)
        cb.font, cb.alignment = BODY_FONT, Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 15 * max(1, -(-len(text) // 130))


def sheet_categories(wb, cs):
    ws = wb.create_sheet("Categories")
    r = title(ws, f"Per Category: {CAT_MODEL}", [
        "One row per category, biggest first, checked on 12 past 30-day windows. 'How far off' = units missed / units sold.",
        "MASE: green = below 1 (missed by less than sales usually change month to month), orange = above 2. Words are explained on the Read Me tab.",
        "Next 30 days = the forecast the app currently shows."])
    cs = cs.sort_values("avg_actual", ascending=False)
    headers = ["Category", "Avg units sold per 30 days", "Avg forecast per 30 days", "How far off, % of units sold",
               "Too high (+) / too low (-), %", "Better than 'same as last 30 days'", "MASE",
               "Next 30 days forecast (units)"]
    rows = [[k, x.avg_actual, x.avg_forecast, x.wmape, x.bias,
             "Yes" if x.beats_naive == 1 else "No", x.mase, x.next_30] for k, x in cs.iterrows()]
    table(ws, r, headers, rows, {1: "0.0", 2: "0.0", 3: "0.0", 4: "+0.0;-0.0", 6: "0.00", 7: "0"},
          [22, 12, 12, 12, 12, 13, 9, 13], mase_fills(rows, 6))


def sheet_items(wb, its):
    ws = wb.create_sheet("Items")
    n_no_mase = int(its.mase.isna().sum())
    r = title(ws, f"Per Item: {ITEM_MODEL}", [
        "One row per forecast item, biggest sellers first, checked on 12 past 30-day windows. 'How far off' = units missed / units sold.",
        "MASE: green = below 1 (missed by less than sales usually change month to month), orange = above 2. Words are explained on the Read Me tab.",
        f"Grey rows: no sale in the last {DISCONTINUED_DAYS} days (probably no longer stocked). They are forecast near zero and sold "
        "nothing, so their MASE is near 0 and pulls the item average down. The Summary tab also shows the items still selling.",
        f"{n_no_mase} item(s) sold nothing in their scoring history, so MASE cannot be worked out for them; MASE covers "
        f"{int(its.mase.notna().sum())} items."])
    its = its.sort_values("avg_actual", ascending=False)
    headers = ["Item", "Category", "Avg units sold per 30 days", "Avg forecast per 30 days",
               "How far off, % of units sold", "Too high (+) / too low (-), %", "Better than 'same as last 30 days'",
               "MASE", "Last sale", "Next 30 days forecast (units)"]
    rows, fills = [], {}
    for i, (k, x) in enumerate(its.iterrows()):
        rows.append([x.item_name, x.category, x.avg_actual, x.avg_forecast, x.wmape, x.bias,
                     "Yes" if x.beats_naive == 1 else "No", x.mase, x.last_sale, x.next_30])
        if x.discontinued:
            for j in range(len(headers)):
                fills[(i, j)] = GREY_FILL
    fills = mase_fills(rows, 7, fills)
    table(ws, r, headers, rows, {2: "0.0", 3: "0.0", 4: "0.0", 5: "+0.0;-0.0", 7: "0.00", 9: "0"},
          [40, 20, 11, 11, 12, 11, 13, 8, 11, 12], fills)


def sheet_by_month(wb, name, heading, df, lead_cols, lead_widths, what):
    ws = wb.create_sheet(name)
    r = title(ws, heading, [
        "Each row: one past 30-day window. Forecast made using only data before the window started.",
        "Error = forecast minus actual (positive = forecast too high). % error is blank when nothing sold. Use the filter arrows to pick one " + what + " or window."])
    headers = lead_cols + ["Window start", "Window end", "Actual units sold", "Forecast units", "Error (units)", "% error"]
    rows = []
    for x in df.itertuples():
        start = pd.Timestamp(x.window_start)
        e = x.pred - x.actual
        rows.append(list(x.lead) + [start.date().isoformat(), (start + pd.Timedelta(days=29)).date().isoformat(),
                                    x.actual, x.pred, e, 100 * e / x.actual if x.actual > 0 else None])
    n = len(lead_cols)
    table(ws, r, headers, rows, {n + 2: "0", n + 3: "0.0", n + 4: "+0.0;-0.0", n + 5: "+0;-0"},
          lead_widths + [12, 12, 11, 11, 11, 10])


def load_cics():
    paths = {k: os.path.join(ROOT, "data", f"cics_{k}.csv") for k in
             ("vs_current_summary", "vs_current_category", "vs_current_items", "category_mapping")}
    missing = [p for p in paths.values() if not os.path.exists(p)]
    if missing:
        raise SystemExit(f"missing {missing} - run scripts/compare_cics_categorization.py first")
    return {k: pd.read_csv(p) for k, p in paths.items()}


def better_fills(rows, col, fills):
    for i, row in enumerate(rows):
        fills[(i, col)] = {"CICS": GOOD_FILL, "Current": BAD_FILL}.get(row[col])
    return {k: v for k, v in fills.items() if v is not None}


def sheet_cics_summary(wb, cics):
    ws = wb.create_sheet("CICS vs Current")
    r = title(ws, "Store Item List (CICS) vs Current 12 Categories: Same Models, Same 12 Windows", [
        "CICS = the store's own sections from FOR CICS STUDENTS.xlsx (17 with sales; Extras/Packaging has no item in our data).",
        "Category level: each grouping's own categories forecast with the 6-month average + calendar. Smaller groups are harder to forecast,",
        "so the fair head-to-head is the ITEM level: the same items, the same actual sales, only the category each item borrows from changes.",
        "Slow items are not forecast by the app; they are scored here the same way as the fast items, for this comparison only.",
        "Difference = CICS minus current (negative = CICS better). 95% range from resampling items: if it spans 0, the two are a tie."])
    s = cics["vs_current_summary"]
    headers = ["Level", "Group of", "How many (current / CICS)", "Current: average MASE", "CICS: average MASE",
               "Current: median MASE", "CICS: median MASE", "Current: error % of units (WMAPE)", "CICS: error % of units (WMAPE)",
               "MASE below 1 (current / CICS)", "Items better with CICS", "Items better with current",
               "Difference in average MASE", "95% range: low", "95% range: high", "Verdict"]
    rows, fills = [], {}
    for cohort, g in s.groupby("cohort", sort=False):
        c, n = g[g.scheme == "current"].iloc[0], g[g.scheme == "cics"].iloc[0]
        if c.level == "category":
            verdict = "Current better" if n.mean_mase > c.mean_mase and n.wmape_pct > c.wmape_pct else "Mixed"
            rows.append(["Category", "categories", f"{int(c.n)} / {int(n.n)}", c.mean_mase, n.mean_mase,
                         c.median_mase, n.median_mase, c.wmape_pct, n.wmape_pct,
                         f"{int(c.mase_below_1)} / {int(n.mase_below_1)}", None, None, n.mean_mase - c.mean_mase,
                         None, None, verdict])
        else:
            tie = c.gap_ci_lo < 0 < c.gap_ci_hi
            verdict = "Tie" if tie else ("CICS better" if c.gap_cics_minus_current < 0 else "Current better")
            rows.append(["Item", cohort, f"{int(c.n)}", c.mean_mase, n.mean_mase, c.median_mase, n.median_mase,
                         c.wmape_pct, n.wmape_pct, f"{int(c.mase_below_1)} / {int(n.mase_below_1)}",
                         int(c.cics_better), int(c.cics_worse), c.gap_cics_minus_current, c.gap_ci_lo, c.gap_ci_hi,
                         verdict])
    r = table(ws, r, headers, rows, {3: "0.000", 4: "0.000", 5: "0.000", 6: "0.000", 7: "0.0", 8: "0.0",
                                     12: "+0.000;-0.000", 13: "+0.000;-0.000", 14: "+0.000;-0.000"},
              [10, 24, 11, 10, 10, 10, 10, 12, 12, 11, 10, 10, 11, 10, 10, 14])
    ws.freeze_panes = None
    ws.auto_filter.ref = None

    # Per-category tables, one per grouping, side by side in reading order.
    cat = cics["vs_current_category"]
    r += 2
    ws.cell(row=r, column=1, value="Per category, each grouping (biggest first). Green MASE below 1, orange above 2.").font = BOLD_FONT
    r += 1
    headers = ["Grouping", "Category", "Avg units sold per 30 days", "MASE", "Error % of units sold (WMAPE)",
               "MAE (units)", "Simple guess MAE (units)", "Beats simple guess"]
    rows = []
    for scheme, label in (("current", "Current"), ("cics", "CICS")):
        g = cat[cat.scheme == scheme].sort_values("mean_actual_30d", ascending=False)
        rows += [[label, x.group, x.mean_actual_30d, x.mase, x.wmape_pct, x.mae, x.naive_mae,
                  "Yes" if x.beats_naive == 1 else "No"] for x in g.itertuples()]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=r, column=j, value=h)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 45
    fills = mase_fills(rows, 3)
    fmt = {2: "0.0", 3: "0.00", 4: "0.0", 5: "0.0", 6: "0.0"}
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(row=r + 1 + i, column=j + 1, value=None if isinstance(v, float) and not np.isfinite(v) else v)
            c.font, c.border = BODY_FONT, BORDER
            if j in fmt:
                c.number_format = fmt[j]
            if (i, j) in fills:
                c.fill = fills[(i, j)]
    ws.column_dimensions["B"].width = 38


def sheet_cics_items(wb, cics):
    ws = wb.create_sheet("CICS Per Item")
    r = title(ws, "Per Item: Current Category vs CICS Section (Fast and Slow Items)", [
        "Same item model the app uses (category share + recent-sales blend, calendar-adjusted); only the category it borrows from changes.",
        "Better with: which grouping gave the lower MASE (green = CICS, orange = current). 'Same' when the category's members did not change,",
        "e.g. Drinkware and Mugs & Tumblers hold exactly the same items. Slow items are not forecast by the app; scored here for comparison only.",
        "Items with no MASE sold nothing in their scoring history (or had too little of it) and are left out of the averages."])
    it = cics["vs_current_items"].copy()
    it["avg"] = it.actual_sum / it.n_folds.where(it.n_folds > 0)
    it = it.sort_values(["fsn_class", "avg"], ascending=[True, False])
    headers = ["Item", "Speed", "Current category", "CICS section", "Category members changed",
               "Test windows", "Avg units sold per 30 days", "Current: MASE", "CICS: MASE", "Difference (CICS - current)",
               "Current: MAE (units)", "CICS: MAE (units)", "Current: over/under, %", "CICS: over/under, %", "Better with"]
    rows = []
    for x in it.itertuples():
        d = x.cics_mase - x.current_mase
        better = ("" if not np.isfinite(d) else "Same" if abs(d) < 1e-9 else "CICS" if d < 0 else "Current")
        bias = lambda b: 100 * b / x.actual_sum if x.actual_sum > 0 else np.nan
        rows.append([x.item_name, "Fast" if x.fsn_class == "F" else "Slow", x.current_category, x.cics_category,
                     "Yes" if x.category_changed else "No", int(x.n_folds), x.avg, x.current_mase, x.cics_mase, d,
                     x.current_mae, x.cics_mae, bias(x.current_bias), bias(x.cics_bias), better])
    fills = better_fills(rows, 14, mase_fills(rows, 8, mase_fills(rows, 7)))
    table(ws, r, headers, rows, {6: "0.0", 7: "0.00", 8: "0.00", 9: "+0.00;-0.00", 10: "0.0", 11: "0.0",
                                 12: "+0.0;-0.0", 13: "+0.0;-0.0"},
          [40, 7, 20, 30, 10, 8, 10, 9, 9, 11, 10, 10, 10, 10, 10], fills)


def sheet_cics_mapping(wb, cics):
    ws = wb.create_sheet("CICS Mapping")
    r = title(ws, "Where Each Product Goes in the CICS Item List", [
        "On CICS list: the item appears in FOR CICS STUDENTS.xlsx, often under a different spelling (e.g. 'C. Keychain' = UST College Keychain).",
        "By name: not on the list; section chosen from the name (price and sale dates used where the name alone was unclear - see Note).",
        "Uncategorized: no clue in the name. Both such products have never sold. Use the filter arrows to review one section."])
    m = cics["category_mapping"].sort_values(["cics", "units_sold"], ascending=[True, False])
    basis = {"list": "On CICS list", "name": "By name", "none": "Uncategorized"}
    speed = {"F": "Fast", "S": "Slow", "N": "Non-moving"}
    rows = [[x.item_name, speed.get(x.fsn_class, x.fsn_class), x.units_sold, x.current, x.cics,
             basis[x.basis], x.note if isinstance(x.note, str) else None] for x in m.itertuples()]
    fills = {(i, 5): BAD_FILL for i, row in enumerate(rows) if row[5] != "On CICS list"}
    table(ws, r, ["Item", "Speed", "Units sold (all history)", "Current category", "CICS section", "How placed", "Note"],
          rows, {2: "0"}, [44, 11, 11, 20, 36, 14, 60], fills)


def load_history():
    paths = [os.path.join(ROOT, "data", f) for f in ("history_length_test.csv", "history_length_series.csv")]
    missing = [p for p in paths if not os.path.exists(p)]
    if missing:
        raise SystemExit(f"missing {missing} - run scripts/test_history_length.py first")
    return pd.read_csv(paths[0]), pd.read_csv(paths[1])


def sheet_history(wb, summ, per, span_months):
    ws = wb.create_sheet("More History Test")
    arms = list(dict.fromkeys(summ.arm))
    cohorts = ["Categories", "Fast items", "Slow items"]
    s = summ.set_index(["cohort", "arm"])
    all_arm = arms[-1]

    def wm(c, a):
        return s.loc[(c, a), "wmape_pct"]
    best = {c: min(arms, key=lambda a: (round(wm(c, a), 1), arms.index(a))) for c in cohorts}
    r = title(ws, "Does More Sales History Improve the Forecasts? (Real Data Only)", [
        "Each past 30-day window was forecast again, but the model was only allowed to see the last 3, 6, 9, 12 or 18 months "
        "of sales before it, or all of it (what the app does).",
        "Same windows, same actual sales, same models as the app - only the amount of history changes. No synthetic data is used.",
        "Months actually used: the earliest test windows had only about 14 months of history in total, so the longer settings "
        "average less than their label.",
        "Lower MASE and lower error % are better. Green = the best setting for that group."])

    headers = ["History given to the model", "Months actually used (average)"]
    for c in cohorts:
        headers += [f"{c}: average MASE", f"{c}: error % of units sold"]
    rows, fills = [], {}
    for i, a in enumerate(arms):
        row = [a, s.loc[(cohorts[0], a), "months_used"]]
        for j, c in enumerate(cohorts):
            row += [s.loc[(c, a), "mean_mase"], wm(c, a)]
            if a == best[c]:
                fills[(i, 3 + 2 * j)] = GOOD_FILL
        rows.append(row)
    fmt = {1: "0.0"}
    for j in range(len(cohorts)):
        fmt[2 + 2 * j], fmt[3 + 2 * j] = "0.000", "0.0"
    r = table(ws, r, headers, rows, fmt, [24, 13, 12, 12, 12, 12, 12, 12], fills)
    ws.freeze_panes = None
    ws.auto_filter.ref = None

    def change(c, a, b):
        return wm(c, b) - wm(c, a)
    first = arms[0]
    six = arms[1]
    year = next(a for a in arms if a.startswith("12"))
    findings = [
        ("More history helps at first", f"Going from 3 to 6 months of history lowers the error from {wm('Categories', first):.1f}% "
                                        f"to {wm('Categories', six):.1f}% for categories ({change('Categories', first, six):+.1f} points) "
                                        f"and from {wm('Fast items', first):.1f}% to {wm('Fast items', six):.1f}% for fast items "
                                        f"({change('Fast items', first, six):+.1f})."),
        ("Then it levels off", f"Past about 9-12 months the numbers stop changing: 12 months, 18 months and all history give "
                               f"{wm('Categories', year):.1f}% for categories either way. The models in use mostly look at "
                               "the last 6 months (the average) and the last year (the calendar adjustment), so older sales "
                               "barely change their forecast."),
        ("What more years would need", "To gain from several years of data, a model has to learn from it - for example a "
                                       "yearly pattern such as 'this month last year'. That needs at least 2-3 years of real "
                                       f"sales, which the store does not have yet (about {span_months} months). Synthetic history cannot "
                                       "stand in for it: it only repeats patterns already in the data (tested before, see "
                                       "docs/SPARSE_DEMAND_EXPERIMENTS.md section 4)."),
        ("In one line", "About 6 to 12 months of history is enough for the current models; more years pay off only "
                        "together with a model that uses them."),
    ]
    r += 2
    ws.cell(row=r, column=1, value="WHAT THIS SHOWS").font = Font(name=FONT, size=11, bold=True, color="1F4E78")
    for head, text in findings:
        r += 1
        ca = ws.cell(row=r, column=1, value=head)
        ca.font, ca.alignment = BOLD_FONT, Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        cb = ws.cell(row=r, column=2, value=text)
        cb.font, cb.alignment = BODY_FONT, Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 15 * max(2, -(-len(text) // 95))

    # Per category: MASE for each history length.
    r += 2
    ws.cell(row=r, column=1, value="Per category: MASE for each amount of history (green = best for that category)").font = BOLD_FONT
    r += 1
    cat = per[per.level == "category"].pivot(index="name", columns="arm", values="mase")[arms]
    cat = cat.loc[per[(per.level == "category") & (per.arm == all_arm)].set_index("name").actual
                  .sort_values(ascending=False).index]
    for j, h in enumerate(["Category"] + arms, start=1):
        c = ws.cell(row=r, column=j, value=h)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (k, vals) in enumerate(cat.iterrows(), start=1):
        lo = vals.round(3).min()
        c = ws.cell(row=r + i, column=1, value=k)
        c.font, c.border = BODY_FONT, BORDER
        for j, a in enumerate(arms, start=2):
            c = ws.cell(row=r + i, column=j, value=float(vals[a]))
            c.font, c.border, c.number_format = BODY_FONT, BORDER, "0.000"
            if round(vals[a], 3) == lo:
                c.fill = GOOD_FILL


def sheet_what_if(wb, wi, n_closed):
    ws = wb.create_sheet("What-If Scenarios")
    scen = list(dict.fromkeys(wi.scenario))
    cohorts = ["Categories", "Fast items", "Slow items"]
    s = wi.set_index(["cohort", "scenario"])
    base = scen[0]
    r = title(ws, "What-If: How Low Could the Error Go if the Store Removed Causes No Model Can Predict?", [
        "These are NOT what the app does. Each row asks: if the store changed one thing about its operations, how much of "
        "today's error would disappear? Same models, same 12 past windows.",
        "Bulk orders known = large one-day purchases (organisation / bulk orders, pre-orders) are logged ahead and added to "
        "the forecast exactly; the model forecasts only regular walk-in sales.",
        "A 'bulk' day is one that sold more than 20x (strict) or 10x (broader) the item's usual selling day. Some of these are "
        "enrollment rushes rather than orders, so treat the bulk rows as the most that logging orders could achieve.",
        "Error % always counts ALL units sold, bulk included, so removing bulk days does not shrink the total the error is "
        "measured against. Green = lower error than the app today."])
    headers = ["Scenario", "Units treated as known in advance, % of all sold"]
    for c in cohorts:
        headers += [f"{c}: error % of units sold", f"{c}: change vs today (points)", f"{c}: average MASE"]
    rows, fills = [], {}
    for i, sc in enumerate(scen):
        row = [sc, s.loc[(cohorts[0], sc), "units_known_in_advance_pct"]]
        for j, c in enumerate(cohorts):
            w, w0 = s.loc[(c, sc), "wmape_pct"], s.loc[(c, base), "wmape_pct"]
            row += [w, w - w0, s.loc[(c, sc), "mean_mase"]]
            if w < w0 - 0.05:
                fills[(i, 2 + 3 * j)] = GOOD_FILL
        rows.append(row)
    fmt = {1: "0.0"}
    for j in range(len(cohorts)):
        fmt[2 + 3 * j], fmt[3 + 3 * j], fmt[4 + 3 * j] = "0.0", "+0.0;-0.0;0.0", "0.00"
    r = table(ws, r, headers, rows, fmt, [38, 14] + [11, 11, 10] * len(cohorts), fills)
    ws.freeze_panes = None
    ws.auto_filter.ref = None

    def w(c, sc):
        return s.loc[(c, sc), "wmape_pct"]
    strict = next(x for x in scen if "20x" in x)
    broad = next(x for x in scen if "10x" in x and "closures" not in x)
    closed = next(x for x in scen if x.startswith("Closures"))
    findings = [
        ("Largest cause tested", f"Large one-day purchases. They are {s.loc[(cohorts[0], strict), 'units_known_in_advance_pct']:.0f}% "
                                 f"(strict) to {s.loc[(cohorts[0], broad), 'units_known_in_advance_pct']:.0f}% (broader) of all "
                                 "units sold, and the sales history gives no warning of them. Some flagged days are "
                                 "enrollment rushes rather than orders, so this is the most logging orders could achieve."),
        ("If the store logged them", f"Category error would fall from {w('Categories', base):.1f}% to {w('Categories', strict):.1f}% "
                                     f"(strict) or {w('Categories', broad):.1f}% (broader); fast items from "
                                     f"{w('Fast items', base):.1f}% to {w('Fast items', strict):.1f}% or {w('Fast items', broad):.1f}%. "
                                     f"Average category MASE would drop from {s.loc[('Categories', base), 'mean_mase']:.2f} to "
                                     f"{s.loc[('Categories', strict), 'mean_mase']:.2f}, below 1."),
        ("Store closures", f"Knowing closures in advance helps very little ({w('Categories', closed) - w('Categories', base):+.1f} "
                           f"points for categories): the store closes rarely, {n_closed} days in the 12 test windows."),
        ("Recommendation", "Record bulk / organisation orders and pre-orders separately from walk-in sales. The model then "
                           "forecasts walk-in demand, and known orders are added on top."),
    ]
    r += 2
    ws.cell(row=r, column=1, value="WHAT THIS SHOWS").font = Font(name=FONT, size=11, bold=True, color="1F4E78")
    for head, text in findings:
        r += 1
        ca = ws.cell(row=r, column=1, value=head)
        ca.font, ca.alignment = BOLD_FONT, Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)
        cb = ws.cell(row=r, column=2, value=text)
        cb.font, cb.alignment = BODY_FONT, Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 15 * max(2, -(-len(text) // 120))


def sheet_calendar_starts(wb, starts):
    """The calendar rule scored at 30 window alignments (scripts/test_calendar_adjustment_starts.py)."""
    ws = wb.create_sheet("Calendar Check")
    r = title(ws, "Does the School-Calendar Rule Help? Tested at 30 Different Start Dates", [
        "The usual test uses 12 past 30-day windows ending on the last day of data. Where they start decides which of them "
        "straddle a break or an enrollment rush, so the result moves with it.",
        "Here the same test is repeated with 0 to 29 days cut off the end of the data, so the windows start on 30 "
        "different days. Error % = all units missed / all units sold. Green = the rule lowered the error."])
    rows, fills = [], {}
    wide = starts.pivot(index=["days_cut", "data_end"], columns="level",
                        values=["first_window", "base_wmape", "capped_wmape"]).reset_index()
    for i, (_, x) in enumerate(wide.iterrows()):
        row = [str(x[("data_end", "")]), str(x[("first_window", "category")])]
        for j, lvl in enumerate(("category", "item")):
            base, cap = x[("base_wmape", lvl)], x[("capped_wmape", lvl)]
            row += [base, cap, cap - base]
            if cap < base:
                fills[(i, 3 + 3 * j)] = GOOD_FILL
        rows.append(row)
    headers = ["Data ends", "First window starts",
               "Categories: error % without rule", "Categories: with rule", "Categories: change (points)",
               "Items: error % without rule", "Items: with rule", "Items: change (points)"]
    fmt = {2: "0.0", 3: "0.0", 4: "+0.0;-0.0;0.0", 5: "0.0", 6: "0.0", 7: "+0.0;-0.0;0.0"}
    r = table(ws, r, headers, rows, fmt, [13, 15, 14, 12, 12, 14, 12, 12], fills)

    r += 2
    ws.cell(row=r, column=1, value="AVERAGE OVER ALL START DATES").font = Font(name=FONT, size=11, bold=True, color="1F4E78")
    for lvl, label in (("category", "Categories"), ("item", "Items")):
        g = starts[starts.level == lvl]
        r += 1
        ws.cell(row=r, column=1, value=label).font = BOLD_FONT
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=len(headers))
        ws.cell(row=r, column=2, value=f"{g.base_wmape.mean():.1f}% without the rule, {g.capped_wmape.mean():.1f}% with it; "
                                       f"the rule lowered the error at {int(g.helped.sum())} of {len(g)} start dates.").font = BODY_FONT


def sheet_grouping_search(wb, cand, path):
    ws = wb.create_sheet("Grouping Search")
    r = title(ws, "Would Fewer or Different Categories Forecast Better?", [
        f"Many groupings were tried: hand-made ones, and a search that merged categories step by step "
        f"({int(path.n_groups.max())} small groups down to 1).",
        "Scored on item accuracy - every grouping forecasts the same items, so the comparison is fair. Each item's forecast borrows "
        "its category's sales, which is where the grouping matters.",
        "To avoid picking a lucky winner, groupings were CHOSEN on the older 6 past months and are SCORED here on the newer 6, "
        "which the search never saw.",
        "Lower MASE is better. Green = best on the unseen months. Category error % always falls with fewer groups (bigger totals "
        "are smoother), so it cannot decide how many groups to have."])
    base = cand.set_index("grouping").loc["Current categories"]
    headers = ["Grouping", "Number of groups", "Chosen-on months: item MASE", "Unseen months: item MASE (all)",
               "Unseen: fast items MASE", "Unseen: slow items MASE", "Change vs current (unseen, all items)",
               "Unseen: category error % of units", "Unseen: category average MASE"]
    rows, fills = [], {}
    best = cand.test_all_mase.min()
    for i, x in enumerate(cand.itertuples()):
        rows.append([x.grouping, int(x.n_groups), x.sel_all_mase, x.test_all_mase, x.test_fast_mase,
                     x.test_slow_mase, x.test_all_mase - base.test_all_mase, x.test_cat_wmape, x.test_cat_mase])
        if abs(x.test_all_mase - best) < 1e-12:
            fills[(i, 3)] = GOOD_FILL
    r = table(ws, r, headers, rows, {2: "0.000", 3: "0.000", 4: "0.000", 5: "0.000", 6: "+0.000;-0.000;0.000",
                                     7: "0.0", 8: "0.00"},
              [34, 9, 12, 12, 11, 11, 12, 12, 12], fills)
    ws.freeze_panes = None
    ws.auto_filter.ref = None

    pick = cand.iloc[-1]
    lo, hi = pick.test_vs_current_ci_lo, pick.test_vs_current_ci_hi
    ci_note = ("includes zero" if lo < 0 < hi else
               "is entirely above zero, so it is genuinely worse there" if lo >= 0 else "is entirely below zero")
    findings = [
        ("Answer", "No. Reducing or re-mixing the categories did not make the item forecasts better on the months the search "
                   "did not see. The current categories scored best, though every sensible grouping is within a narrow band."),
        ("The search result", f"The best grouping on the older months ({int(pick.n_groups)} groups) looked clearly better there "
                              f"({pick.sel_all_mase:.2f} vs {base.sel_all_mase:.2f}), but on the newer months it was slightly worse "
                              f"({pick.test_all_mase:.3f} vs {base.test_all_mase:.3f}; the 95% range of the difference, "
                              f"{lo:+.3f} to {hi:+.3f}, {ci_note}). Its "
                              "advantage was luck, which is exactly what testing on unseen months is for."),
        ("Fewer groups", "Two groups (apparel / non-apparel) or one for the whole store make item forecasts worse, and the "
                         "dashboard would lose detail. Their lower category error % only reflects bigger, smoother totals."),
        ("Worth considering", "The store's own sections with the tiny ones folded in (11 groups) forecast items about as well as "
                              "the current categories and use the store's own names - an option if the client prefers them."),
    ]
    r += 2
    ws.cell(row=r, column=1, value="WHAT THIS SHOWS").font = Font(name=FONT, size=11, bold=True, color="1F4E78")
    for head, text in findings:
        r += 1
        ca = ws.cell(row=r, column=1, value=head)
        ca.font, ca.alignment = BOLD_FONT, Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        cb = ws.cell(row=r, column=2, value=text)
        cb.font, cb.alignment = BODY_FONT, Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 15 * max(2, -(-len(text) // 110))

    r += 2
    ws.cell(row=r, column=1, value=f"The step-by-step search: accuracy as categories are merged ({int(path.n_groups.max())} groups down to 1)").font = BOLD_FONT
    r += 1
    heads = ["Number of groups", "Chosen-on months: item MASE", "Unseen months: item MASE", "Unseen: category error % of units"]
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=r, column=j, value=h)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 45
    for i, x in enumerate(path.itertuples(), start=1):
        for j, (v, f) in enumerate(((int(x.n_groups), "0"), (x.sel_all_mase, "0.000"),
                                    (x.test_all_mase, "0.000"), (x.test_cat_wmape, "0.0")), start=1):
            c = ws.cell(row=r + i, column=j, value=v)
            c.font, c.border, c.number_format = BODY_FONT, BORDER, f


# ---------------------------------------------------------------- worked example
def worked_example(con, its, cat_all, item_all):
    """The biggest still-selling item's forecast for the last past window, rebuilt step by
    step from the parts of the production model, plus its category's. Stops if the steps do
    not add up to the forecasts the walk-forward test scored, so the tab cannot drift from
    the model it describes."""
    key = its[~its.discontinued].avg_actual.idxmax()
    item_name, cat_name = its.item_name[key], its.category[key]
    fact = pd.read_sql_query("""SELECT f.product_id, d.calendar_date, f.quantity_sold FROM Fact_Sales f
        JOIN Dim_Date d ON d.date_id = f.date_id""", con, parse_dates=["calendar_date"])
    cat_of = (pd.read_sql_query("SELECT product_id, forecast_category FROM Dim_Product", con)
              .set_index("product_id").forecast_category.fillna(tca.RESIDUE))
    idx = pd.date_range(fact.calendar_date.min(), fact.loc[fact.quantity_sold > 0, "calendar_date"].max(), freq="D")
    daily = (fact.groupby(["product_id", "calendar_date"]).quantity_sold.sum().unstack(0)
             .reindex(idx, fill_value=0.0).fillna(0.0).astype(float))
    v = daily[key].to_numpy()
    cv = daily[[p for p in daily.columns if cat_of.get(p) == cat_name]].sum(axis=1).to_numpy()
    h = tca.H
    types = load_day_types(con, idx, h)
    fold = make_folds(v.size, h, tca.MIN_F, tca.MAX_F, tca.MIN_T)[-1]
    n = fold.train_end
    t, c = v[:n], cv[:n]
    level = c[-LEVEL_WINDOW:].mean() * h
    item30, cat30 = t[-SHARE_WINDOW:].sum(), c[-SHARE_WINDOW:].sum()
    share = item30 / cat30 if cat30 > 0 else 0.0
    r = tsb(t, TSB_ALPHA, TSB_BETA)
    topdown, tsb_total = level * share, r.point_forecast * h
    blend = 0.5 * (topdown + tsb_total)
    mult = calendar_multiplier(c, types, n, h)
    raw_mult = calendar_multiplier(c, types, n, h, cap=False)
    start = idx[fold.test_start].date()

    def scored(df, k, version):
        g = df[(df.key == k) & (df.version == version) & (df.window_start == start)]
        if len(g) != 1:
            raise SystemExit(f"worked example: no single scored window for {k} starting {start}")
        return g.iloc[0]
    for got, want in ((scored(item_all, key, "base").pred, blend),
                      (scored(item_all, key, "capped").pred, blend * mult),
                      (scored(cat_all, cat_name, "base").pred, level),
                      (scored(cat_all, cat_name, "capped").pred, level * mult)):
        if abs(got - want) > 1e-6:
            raise SystemExit(f"worked example does not reproduce the model: {got} vs {want}")
    it, ct = scored(item_all, key, "capped"), scored(cat_all, cat_name, "capped")
    return dict(item=item_name, category=cat_name, made=idx[n - 1].date(), start=start,
                end=idx[fold.test_end - 1].date(), level=level, item30=item30, cat30=cat30, share=share,
                topdown=topdown, p=r.probability_estimate, z=r.size_estimate, tsb=tsb_total, blend=blend,
                mult=mult, raw_mult=raw_mult, forecast=it.pred, actual=it.actual,
                cat_forecast=ct.pred, cat_actual=ct.actual)


def text_table(ws, r0, headers, rows, chars):
    """A table of wrapped text: first column bold, row height from the longest cell."""
    for j, hd in enumerate(headers, start=1):
        c = ws.cell(row=r0, column=j, value=hd)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 30
    for i, row in enumerate(rows, start=1):
        lines = 1
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=r0 + i, column=j, value=v)
            c.font, c.border = (BOLD_FONT if j == 1 else BODY_FONT), BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
            lines = max(lines, -(-len(str(v)) // chars[j - 1]))
        ws.row_dimensions[r0 + i].height = 15 * lines
    return r0 + len(rows)


def sheet_how_it_works(wb, e, cat, item):
    ws = wb.create_sheet("How It Works")
    for col, w in zip("ABC", (34, 60, 60)):
        ws.column_dimensions[col].width = w
    chars = (30, 55, 55)
    r = title(ws, "How the Forecast Works", [
        "The app makes two forecasts of the next 30 days: one per category and one per fast-moving item.",
        "Below: the two models side by side, then one real forecast worked through step by step."])
    models = [
        ("What it forecasts", f"Total sales of each of the {cat['n']} categories, all their items included",
         f"Sales of each of the {item['n']} items classed as fast-moving"),
        ("Model (technical name)", "6-month moving average, calendar-adjusted",
         "Category share + TSB (Teunter-Syntetos-Babai), half and half, calendar-adjusted"),
        ("In plain words", "The category's average daily sales over the last 6 months, times 30.",
         "Half: the item's share of its category in the last 30 days, times the category's 6-month level. "
         "Half: how often the item has sold lately, times how many it sells each time; this part fades toward "
         "zero if the item stops selling."),
        ("School calendar", "Lowered when more semester-break or exam days are coming than in the recent past. "
                            "Never raised.", "Same, using its category's sales."),
        ("Day by day", "Spread over the 30 days by the category's usual weekly pattern (from Prophet). The total "
                       "does not change.", "Same."),
        ("How far off (12 past months)", f"{cat['wmape']:.1f}% of units sold", f"{item['wmape']:.1f}% of units sold"),
    ]
    r = text_table(ws, r + 1, ["", "Per category", "Per item"], models, chars)

    r += 2
    ws.cell(row=r, column=1, value=f"Example: {e['item']} and its category, {e['category']}").font = \
        Font(name=FONT, size=12, bold=True, color="1F4E78")
    ws.cell(row=r + 1, column=1, value=f"Forecast made with sales up to {e['made']}, for {e['start']} to {e['end']}: "
                                       "the last past month that can be checked against real sales.").font = NOTE_FONT
    ws.cell(row=r + 2, column=1, value="Every number is recomputed from the model's own parts; the workbook build stops "
                                       "if they do not add up to the forecast the test scored.").font = NOTE_FONT
    if e["raw_mult"] > 1 + 1e-9:
        cal = f"x {e['mult']:.2f}: the calendar pointed {e['raw_mult'] - 1:+.0%}, but a forecast is never raised"
    elif e["mult"] < 1 - 1e-9:
        cal = f"x {e['mult']:.2f}: more quiet days are coming, so the forecast is lowered"
    else:
        cal = f"x {e['mult']:.2f}"

    def off(f_, a):
        return f"{abs(f_ - a):,.0f} units ({abs(f_ - a) / a:.0%} of what sold)" if a > 0 else f"{abs(f_ - a):,.0f} units"
    steps = [
        ("1. Category's average over the last 6 months", f"{e['level'] / 30:.1f} a day x 30 = {e['level']:,.0f}",
         f"{e['level'] / 30:.1f} a day x 30 = {e['level']:,.0f}"),
        ("2. Item's share of the category, last 30 days",
         f"{e['item30']:,.0f} of {e['cat30']:,.0f} units = {e['share']:.1%}, so {e['level']:,.0f} x {e['share']:.1%} "
         f"= {e['topdown']:,.0f}", "Not used"),
        ("3. Item's own recent sales (TSB)", f"Sells on about {e['p']:.0%} of days, about {e['z']:.1f} units each "
                                             f"time: about {e['tsb']:,.0f} in 30 days", "Not used"),
        ("4. Average of steps 2 and 3", f"({e['topdown']:,.0f} + {e['tsb']:,.0f}) / 2 = {e['blend']:,.0f}", "Not used"),
        ("5. School-calendar check", cal, cal),
        ("Forecast", f"{e['forecast']:,.0f}", f"{e['cat_forecast']:,.0f}"),
        ("Actually sold", f"{e['actual']:,.0f}", f"{e['cat_actual']:,.0f}"),
        ("Off by", off(e["forecast"], e["actual"]), off(e["cat_forecast"], e["cat_actual"])),
    ]
    r = text_table(ws, r + 4, ["Step", f"{e['item']} (item)", f"{e['category']} (its category)"], steps, chars)
    ws.cell(row=r + 2, column=1, value="Every item goes through the same steps every time the forecast runs; only the "
                                       "numbers change.").font = NOTE_FONT


# ---------------------------------------------------------------- main
def main():
    cics = load_cics()
    hist_summ, hist_per = load_history()
    wi_path = os.path.join(ROOT, "data", "what_if_test.csv")
    if not os.path.exists(wi_path):
        raise SystemExit(f"missing {wi_path} - run scripts/test_what_if.py first")
    what_if = pd.read_csv(wi_path)
    gs = [os.path.join(ROOT, "data", f) for f in ("grouping_search_candidates.csv", "grouping_search_path.csv")]
    if not all(os.path.exists(p) for p in gs):
        raise SystemExit("missing data/grouping_search_*.csv - run scripts/search_groupings.py first")
    gs_cand, gs_path = pd.read_csv(gs[0]), pd.read_csv(gs[1])
    starts_path = os.path.join(ROOT, "data", "calendar_adjustment_starts.csv")
    if not os.path.exists(starts_path):
        raise SystemExit(f"missing {starts_path} - run scripts/test_calendar_adjustment_starts.py first")
    starts = pd.read_csv(starts_path)
    con = sqlite3.connect("file:%s?mode=ro" % DB_PATH, uri=True)
    # All three versions are kept: "capped" is what the app ships; "base" (no calendar
    # rule) is quoted on the Read Me so the rule's in-sample gain is visible.
    cat_all = pd.DataFrame(tca.category_rows(con))
    item_all = pd.DataFrame(tca.item_rows(con))
    cat, item = cat_all[cat_all.version == "capped"], item_all[item_all.version == "capped"]

    models = {r[0] for r in con.execute("SELECT DISTINCT model_type FROM Result_Forecast")}
    cmodels = {r[0] for r in con.execute("SELECT DISTINCT model_type FROM Result_Category_Forecast")}
    if models != {ITEM_MODEL_ID} or cmodels != {CAT_MODEL_ID}:
        print("ustore.db holds a different model than this workbook describes:", models, cmodels)
        return 1

    cstored = pd.read_sql_query("""SELECT forecast_category AS key, mae, mase, naive_mae, beats_naive_mae
        FROM Result_Category_Forecast_Metrics WHERE period_scope = 'overall'""", con).set_index("key")
    istored = pd.read_sql_query("""SELECT product_id AS key, mae, mase, naive_mae, beats_naive_mae
        FROM Result_Forecast_Metrics WHERE period_scope = 'overall'""", con).set_index("key")
    cnext = pd.read_sql_query("""SELECT forecast_category AS key, SUM(yhat) AS next_30, MIN(forecast_date) AS s,
        MAX(forecast_date) AS e, MAX(snapshot_date) AS snap FROM Result_Category_Forecast GROUP BY 1""", con).set_index("key")
    inext = pd.read_sql_query("SELECT product_id AS key, SUM(yhat) AS next_30 FROM Result_Forecast GROUP BY 1",
                              con).set_index("key")
    prod = pd.read_sql_query("SELECT product_id AS key, item_name, forecast_category FROM Dim_Product", con).set_index("key")
    last = pd.read_sql_query("""SELECT f.product_id AS key, MAX(d.calendar_date) AS last_sale FROM Fact_Sales f
        JOIN Dim_Date d ON d.date_id = f.date_id WHERE f.quantity_sold > 0 GROUP BY 1""", con).set_index("key")
    n_products = con.execute("SELECT COUNT(*) FROM Dim_Product").fetchone()[0]
    first_sale, last_sale = con.execute("""SELECT MIN(d.calendar_date), MAX(d.calendar_date) FROM Fact_Sales f
        JOIN Dim_Date d ON d.date_id = f.date_id WHERE f.quantity_sold > 0""").fetchone()
    span_months = round((pd.Timestamp(last_sale) - pd.Timestamp(first_sale)).days / 30.44)

    cs, its = score(cat), score(item)
    check_against_db(cs, cstored, "Categories")
    check_against_db(its, istored, "Items")

    cs = cs.join(cstored[["naive_mae", "beats_naive_mae"]].rename(columns={"beats_naive_mae": "beats_naive"}))
    cs = cs.join(cnext.next_30)
    its = its.join(istored[["naive_mae", "beats_naive_mae"]].rename(columns={"beats_naive_mae": "beats_naive"}))
    its = its.join(inext.next_30).join(prod).join(last)
    its["category"] = its.forecast_category.fillna("Uncategorised")
    snapshot = cnext.snap.max()
    its["discontinued"] = (pd.Timestamp(snapshot) - pd.to_datetime(its.last_sale)).dt.days > DISCONTINUED_DAYS

    wins = sorted(cat.window_start.unique())
    first_win, last_win = pd.Timestamp(wins[0]).date(), (pd.Timestamp(wins[-1]) + pd.Timedelta(days=29)).date()
    n_closed = con.execute("SELECT COUNT(*) FROM Dim_Date WHERE is_store_closed = 1 AND calendar_date BETWEEN ? AND ?",
                           (str(first_win), str(last_win))).fetchone()[0]
    example = worked_example(con, its, cat_all, item_all)
    con.close()

    # Items that stopped selling are forecast ~0 and sell 0, so they score a near-perfect MASE;
    # the still-selling row shows the items without that flattery.
    live = its.index[~its.discontinued]
    summ = [summary_row("Categories", CAT_MODEL, cs, cat),
            summary_row(f"Items, all {len(its)}", ITEM_MODEL, its, item),
            summary_row(f"Items still selling ({len(live)})", ITEM_MODEL, its.loc[live], item[item.key.isin(live)])]
    wi = what_if.set_index(["cohort", "scenario"])
    bulk = [wi.loc[("Fast items", sc), "wmape_pct"] for sc in dict.fromkeys(what_if.scenario)
            if sc.startswith("Bulk orders known")]
    facts = dict(snapshot=snapshot, fc_start=cnext.s.min(), fc_end=cnext.e.max(), first_win=first_win,
                 last_win=last_win, cat=summ[0], item=summ[1], live=summ[2], n_other=n_products - len(its),
                 bulk_lo=min(bulk), bulk_hi=max(bulk), ex_item=example["item"],
                 item_base=wmape(item_all[item_all.version == "base"]),
                 cat_base=wmape(cat_all[cat_all.version == "base"]))
    # Whether the calendar rule helps on these windows decides how the Read Me words its caveat.
    helps = [facts[b] > s["wmape"] for b, s in (("item_base", summ[1]), ("cat_base", summ[0]))]
    facts["cal_effect"] = ("so these scores are slightly flattering" if all(helps) else
                           "yet on these windows the rule makes the forecasts worse, not better" if not any(helps) else
                           "and on these windows it helps one level and hurts the other")
    # The fairer figure: the same test at 30 window alignments (Calendar Check tab).
    facts["starts"] = {lvl: dict(n=len(g), base=g.base_wmape.mean(), capped=g.capped_wmape.mean(),
                                 helped=int(g.helped.sum()))
                       for lvl, g in starts.groupby("level")}

    cat = cat.sort_values(["key", "fold"]).assign(lead=lambda d: [(k,) for k in d.key])
    item = item.assign(avg=item.key.map(its.avg_actual)).sort_values(["avg", "key", "fold"], ascending=[False, True, True])
    item["lead"] = [(its.item_name[k], its.category[k]) for k in item.key]

    wb = Workbook()
    sheet_readme(wb, facts)
    sheet_how_it_works(wb, example, summ[0], summ[1])
    sheet_summary(wb, summ, len(its) - len(live))
    sheet_categories(wb, cs)
    sheet_items(wb, its)
    sheet_by_month(wb, "Category by Month", "Category Forecast vs Actual, Each Past 30-Day Window",
                   cat, ["Category"], [22], "category")
    sheet_by_month(wb, "Item by Month", "Item Forecast vs Actual, Each Past 30-Day Window",
                   item, ["Item", "Category"], [40, 20], "item")
    sheet_cics_summary(wb, cics)
    sheet_cics_items(wb, cics)
    sheet_cics_mapping(wb, cics)
    sheet_history(wb, hist_summ, hist_per, span_months)
    sheet_what_if(wb, what_if, n_closed)
    sheet_calendar_starts(wb, starts)
    sheet_grouping_search(wb, gs_cand, gs_path)
    wb.save(OUT)

    for s in summ:
        print("  %-24s WMAPE %.1f%%  bias %+.1f%%  beats naive %s  MAPE<=%.0f%%: %s  average MASE %.3f "
              "(median %.3f, below 1: %s)" % (s["level"], s["wmape"], s["bias"], s["beats"], MAPE_TARGET,
                                             s["target"], s["mase_avg"], s["mase_med"], s["mase_below"]))
    print("  without the calendar rule: items WMAPE %.1f%%, categories %.1f%%" % (facts["item_base"], facts["cat_base"]))
    for lvl, s in facts["starts"].items():
        print("  calendar rule over %d start dates (%s): %.1f%% -> %.1f%%, helped at %d"
              % (s["n"], lvl, s["base"], s["capped"], s["helped"]))
    print("  worked example: %s, forecast %.1f vs actual %.0f" % (example["item"], example["forecast"], example["actual"]))
    print("Wrote", os.path.relpath(OUT, ROOT))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
