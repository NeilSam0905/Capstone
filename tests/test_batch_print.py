"""
The Batch Sales Report PDF is the .xlsx as Excel prints it.

batch_pdf.py draws the PDF from the same layout functions batch_export.py
uses for the sheet's print setup, so the downloaded PDF, the Print Preview
and a print from Excel come out as one document. These tests pin the shared
parts: the sheet's page setup matches the layout, both files break pages at
the same rows, a supplier block that fits a page is never split, and the PDF
still renders without the Windows fonts.
"""
import io
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

import batch_export  # noqa: E402
import batch_pdf  # noqa: E402

DAYS = [f"2026-07-{d:02d}" for d in range(1, 32)]


def report(suppliers=6, items=12):
    out, pid = [], 0
    for s in range(suppliers):
        rows = []
        for i in range(items):
            pid += 1
            rows.append({"product_id": pid, "item_name": f"Item {s}-{i} with a long name (XL)",
                         "quantity": 3, "unit_price_php": None if i == 0 else 100.0,
                         "line_total": None if i == 0 else 300.0})
        out.append({"supplier": f"SUPPLIER {s}", "total_units": 3 * items,
                    "subtotal": 300.0 * (items - 1), "items": rows})
    return out


def daily(rep):
    q = {it["product_id"]: {"2026-07-02": 1, "2026-07-09": 2} for e in rep for it in e["items"]}
    return {"dates": DAYS, "marks": {"2026-07-04": "closed", "2026-07-20": "exam"},
            "quantities": q, "selling_days": 2}


def sheet(rep):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(batch_export.render_xlsx(rep, "2026-07", daily(rep)))).active


def pdf_pages(data):
    return len(re.findall(rb"/Type\s*/Page\b", data))


def test_the_sheet_carries_the_print_setup_the_pdf_draws():
    rep = report()
    ws = sheet(rep)
    assert ws.page_setup.orientation == "landscape"
    assert int(ws.page_setup.paperSize) == int(ws.PAPERSIZE_A4)
    assert ws.page_setup.scale == round(batch_export.print_scale(len(DAYS)) * 100)
    assert ws.print_title_rows.replace("$", "") == "1:1"   # the date row on every page
    assert len(ws.print_area.split(",")) == 2              # the grid, then the legend


def test_both_files_break_pages_at_the_same_rows():
    rep = report()
    rows = batch_export.grid_rows(rep)
    starts = batch_export.page_starts(rows, batch_export.print_scale(len(DAYS)))
    assert starts                                           # this report spans pages
    ws = sheet(rep)
    assert [b.id + 1 for b in ws.row_breaks.brk] == starts  # a break above each start
    pdf = batch_pdf.render(rep, "2026-07", daily(rep))
    assert pdf_pages(pdf) == len(starts) + 2                # + first page + legend page


def test_a_supplier_block_that_fits_a_page_is_never_split():
    rep = report(suppliers=8, items=10)
    rows = batch_export.grid_rows(rep)
    for start in batch_export.page_starts(rows, batch_export.print_scale(len(DAYS))):
        assert rows[start - 2][0] == "supplier"


def test_the_pdf_renders_without_the_windows_fonts(monkeypatch, tmp_path):
    monkeypatch.setattr(batch_pdf, "_font_dir", lambda: tmp_path)   # no fonts there
    rep = report(suppliers=1, items=2)
    rep[0]["supplier"] = "Café “Quoted” ₱ Supplier"
    assert batch_pdf.render(rep, "2026-07", daily(rep)).startswith(b"%PDF")


def test_an_empty_month_renders_in_both():
    ws = sheet([])
    assert "No sales were recorded" in str(ws["A2"].value)
    assert batch_pdf.render([], "2026-07", daily([])).startswith(b"%PDF")
