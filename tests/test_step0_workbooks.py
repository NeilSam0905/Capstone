"""
tests/test_step0_workbooks.py
------------------------------------------------------------------
step0 reads every tally workbook in rawdata/ (it used to read a fixed list
of four), so next academic year's workbook needs no code change. Checked:

  - the Tally Interface's import_* archives and Excel's ~$ lock files are
    not read (the archives are already in Fact_Sales);
  - no workbook at all raises FileNotFoundError - the pipeline then skips
    step0 and keeps its CSV - instead of writing an empty one;
  - a month in two workbooks is read once, from the newer workbook;
  - the ITEM PRICE column is captured during conversion, so step1 opens no
    workbook, and an unchanged set of workbooks is not converted again;
  - a sheet that numbers its rows down column A is read from column B;
  - the newest month's blank days at the end (not tallied yet) are left
    out rather than read as zero sales; an older month's are kept.
------------------------------------------------------------------
"""
import datetime
import os
import sys

import openpyxl
import pytest

import step0_convert_sales_with_zeros as step0  # conftest.py puts scripts/ on sys.path


def tally_workbook(path, month_sheets, mtime=None):
    """{sheet name: (first day, units sold that day)} -> a minimal TBS workbook."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, (day, units) in month_sheets.items():
        ws = wb.create_sheet(sheet)
        ws.append(["ITEMS", day, day + datetime.timedelta(days=1), "TOTAL QUANTITY"])
        ws.append(["TEST SUPPLIER (CONSIGNMENT)"])
        ws.append(["Test Tote", units, 0, units])
        ws.append(["TOTAL", units, 0, units])
    wb.save(path)
    if mtime is not None:
        os.utime(path, (mtime, mtime))
    return str(path)


def test_skips_app_imports_and_lock_files(tmp_path):
    jan = {"JANUARY 2027 - TBS": (datetime.datetime(2027, 1, 4), 3)}
    tally_workbook(tmp_path / "USTore TBS 2027.xlsx", jan)
    tally_workbook(tmp_path / "import_20270105-101010_tally.xlsx", jan)
    tally_workbook(tmp_path / "~$USTore TBS 2027.xlsx", jan)
    (tmp_path / "notes.csv").write_text("x")
    files = step0.find_workbooks(str(tmp_path))
    assert [os.path.basename(f) for f in files] == ["USTore TBS 2027.xlsx"]


def test_no_workbooks_is_file_not_found(tmp_path):
    with pytest.raises(FileNotFoundError):
        step0.find_workbooks(str(tmp_path))
    with pytest.raises(FileNotFoundError):
        step0.find_workbooks(str(tmp_path / "missing"))


def test_workbook_without_tally_sheets_is_file_not_found(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "OPEX 2027"
    wb.save(tmp_path / "OPEX.xlsx")
    with pytest.raises(FileNotFoundError):
        step0.plan_sheets(step0.find_workbooks(str(tmp_path)))


def test_a_new_workbook_is_read_without_a_code_change(tmp_path):
    old = tally_workbook(tmp_path / "USTore TBS OCTOBER A.Y. 2025-2026.xlsx",
                         {"JULY 2026 - TBS": (datetime.datetime(2026, 7, 1), 5)}, mtime=1_000_000)
    new = tally_workbook(tmp_path / "USTore TBS A.Y. 2026-2027.xlsx",
                         {"AUGUST 2026 - TBS": (datetime.datetime(2026, 8, 3), 7)}, mtime=2_000_000)
    plan, notes = step0.plan_sheets(step0.find_workbooks(str(tmp_path)))
    assert plan == [(old, ["JULY 2026 - TBS"]), (new, ["AUGUST 2026 - TBS"])]
    assert notes == []


def priced_workbook(path, units, price):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AUGUST 2026 - TBS"
    ws.append(["ITEMS", datetime.datetime(2026, 8, 3), "ITEM PRICE", "TOTAL QUANTITY"])
    ws.append(["TEST SUPPLIER"])
    ws.append(["Test Tote", units, price, units])
    ws.append(["TOTAL", units, None, units])
    wb.save(path)


def test_conversion_also_captures_the_price_column_for_step1(tmp_path):
    priced_workbook(tmp_path / "TBS.xlsx", 2, 150)
    plan, _ = step0.plan_sheets(step0.find_workbooks(str(tmp_path)))
    prices = {}
    step0.convert(plan, str(tmp_path / "sales.csv"), prices)
    assert prices["tbs"] == [("Test Tote", "2026-08", 150.0)]


def test_an_unchanged_run_is_skipped_and_a_changed_one_is_not(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "rawdata").mkdir()
    (tmp_path / "data").mkdir()
    monkeypatch.setattr(sys, "argv", ["step0"])
    priced_workbook(tmp_path / "rawdata" / "TBS.xlsx", 2, 150)

    step0.main()
    assert "Wrote" in capsys.readouterr().out
    step0.main()
    assert "have not changed" in capsys.readouterr().out

    priced_workbook(tmp_path / "rawdata" / "TBS.xlsx", 9, 150)        # the store updated the sheet
    step0.main()
    assert "Wrote" in capsys.readouterr().out
    assert "2026-08-03,Test Tote,9" in (tmp_path / step0.OUT_PATH).read_text()

    monkeypatch.setattr(sys, "argv", ["step0", "--force"])
    step0.main()
    assert "Wrote" in capsys.readouterr().out


def test_a_month_in_two_workbooks_is_read_once_from_the_newer(tmp_path):
    july_8 = {"JULY 2026 - TBS": (datetime.datetime(2026, 7, 1), 5)}
    july_31 = {"JULY 2026 - TBS": (datetime.datetime(2026, 7, 1), 9)}
    tally_workbook(tmp_path / "TBS copy 1.xlsx", july_8, mtime=1_000_000)
    newer = tally_workbook(tmp_path / "TBS copy 2.xlsx", july_31, mtime=2_000_000)
    plan, notes = step0.plan_sheets(step0.find_workbooks(str(tmp_path)))
    assert plan == [(newer, ["JULY 2026 - TBS"])]
    assert len(notes) == 1 and "2026-07" in notes[0] and "TBS copy 1.xlsx" in notes[0]

    out = tmp_path / "sales.csv"
    step0.convert(plan, str(out))
    rows = out.read_text().splitlines()[1:]
    assert rows == ["2026-07-01,Test Tote,9,TEST SUPPLIER (CONSIGNMENT)",
                    "2026-07-02,Test Tote,0,TEST SUPPLIER (CONSIGNMENT)"]


def month_sheet(wb, title, first_day, days, sold, numbered=False):
    """A dense month: `days` date columns, the item sold `sold` (one value
    per day, None for a blank cell) - optionally with a row number in A."""
    ws = wb.create_sheet(title)
    dates = [first_day + datetime.timedelta(days=k) for k in range(days)]
    lead = [None] if numbered else []
    ws.append(lead + [first_day, *dates, "TOTAL QUANTITY"])
    ws.append(lead + ["TEST SUPPLIER"])
    ws.append(([1] if numbered else []) + ["Test Tote", *sold, sum(q or 0 for q in sold)])
    ws.append(([2] if numbered else []) + ["Test Cap", *sold, sum(q or 0 for q in sold)])
    ws.append(lead + ["TOTAL"])


def converted(tmp_path, wb):
    path = tmp_path / "TBS.xlsx"
    wb.save(path)
    out = tmp_path / "sales.csv"
    step0.convert(step0.plan_sheets([str(path)])[0], str(out))
    return [line.split(",") for line in out.read_text().splitlines()[1:]]


def test_row_numbers_in_column_a_are_not_item_names(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    month_sheet(wb, "AUGUST 2026 - TBS", datetime.datetime(2026, 8, 1), 3, [2, None, 4],
                numbered=True)
    rows = converted(tmp_path, wb)
    assert {r[1] for r in rows} == {"Test Tote", "Test Cap"}
    assert {r[3] for r in rows} == {"TEST SUPPLIER"}
    assert sum(float(r[2]) for r in rows if r[1] == "Test Tote") == 6


def test_the_newest_months_blank_tail_is_not_read_as_zero_sales(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    # Every day has a column, so both sheets are dense (blank = sold zero).
    # September is over: its blank last two days are the store shut, real zeros.
    month_sheet(wb, "SEPTEMBER 2026 - TBS", datetime.datetime(2026, 9, 1), 30,
                [3] * 28 + [None, None])
    # October is still being filled in: a 0 typed on the 11th is a tally, the
    # blank 12th to 31st are days nobody has reached yet.
    month_sheet(wb, "OCTOBER 2026 - TBS", datetime.datetime(2026, 10, 1), 31,
                [2] * 9 + [None, 0] + [None] * 20)
    rows = converted(tmp_path, wb)
    september = sorted({r[0] for r in rows if r[0].startswith("2026-09")})
    october = sorted({r[0] for r in rows if r[0].startswith("2026-10")})
    assert (september[0], september[-1], len(september)) == ("2026-09-01", "2026-09-30", 30)
    assert (october[0], october[-1], len(october)) == ("2026-10-01", "2026-10-11", 11)
    assert [r[2] for r in rows if r[0] == "2026-10-10" and r[1] == "Test Tote"] == ["0"]
