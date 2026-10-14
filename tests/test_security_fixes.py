"""
tests/test_security_fixes.py
------------------------------------------------------------------
Three fixes from the backend security review:

  - dates are checked as real calendar days, and /api/events refuses one
    that is not (a stored "not-a-date" used to fail every pipeline run);
    the calendar's log readers skip such a row instead of crashing;
  - the batch report's CSV and XLSX keep item and supplier names as text,
    so a name starting with "=" is never run as a formula in Excel, while
    the sheet's own =SUM formulas stay formulas;
  - defusedxml is pinned, so openpyxl's XML-bomb guard comes with a fresh
    install.
------------------------------------------------------------------
"""
import io
import sqlite3
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "backend"))

import batch_export  # noqa: E402
import calendar_extent  # noqa: E402  (scripts/, via conftest.py)
import validation  # noqa: E402

# --------------------------------------------------------------- dates


def test_is_iso_date_wants_a_real_day():
    assert validation.is_iso_date("2026-10-14")
    assert validation.is_iso_date("2028-02-29")              # leap day
    for bad in ("2026-13-45", "2027-02-29", "not-a-date", "2026-1-5", "", None, 20261014):
        assert not validation.is_iso_date(bad), bad


def test_add_event_refuses_a_date_that_is_not_one():
    import app as backend
    # Skip the first-request calendar extension: this test must not write to
    # the real ustore.db, and the refusal happens before any database use.
    backend._calendar_checked = True
    client = backend.app.test_client()
    with client.session_transaction() as s:
        s["user"] = "staff"
    for bad in ("not-a-date", "2026-13-45"):
        r = client.post("/api/events", json={"calendar_date": bad, "event_name": "x"})
        assert r.status_code == 400
        assert r.get_json()["errors"]["calendar_date"] == "Date must be YYYY-MM-DD."


def test_log_readers_skip_an_unreadable_date_instead_of_crashing():
    c = sqlite3.connect(":memory:")
    c.executescript("""
        CREATE TABLE Closure_Log (closure_id INTEGER PRIMARY KEY AUTOINCREMENT,
                                  closure_date TEXT, is_closed INTEGER);
        CREATE TABLE Event_Log (event_id INTEGER PRIMARY KEY AUTOINCREMENT, event_date TEXT);
        INSERT INTO Closure_Log (closure_date, is_closed) VALUES ('2026-12-24', 1), ('garbage', 1);
        INSERT INTO Event_Log (event_date) VALUES ('2026-12-18'), ('not-a-date');
    """)
    from datetime import date
    assert calendar_extent.load_closure_log(c) == {date(2026, 12, 24): 1}
    assert calendar_extent.load_event_dates(c) == {date(2026, 12, 18)}

# ------------------------------------------------------------- exports


REPORT = [{
    "supplier": "+SUPPLIER",
    "total_units": 7,
    "items": [
        {"product_id": 1, "item_name": "=HYPERLINK(\"http://x\",\"click\")",
         "quantity": 5, "unit_price_php": 100.0, "line_total": 500.0},
        {"product_id": 2, "item_name": "Lanyard @180", "quantity": 2,
         "unit_price_php": None, "line_total": None},
    ],
}]


def test_csv_neutralises_formula_names_and_leaves_numbers_alone():
    text = batch_export.render_csv(REPORT, "2026-07").decode("utf-8-sig")
    assert "'=HYPERLINK" in text
    assert "'+SUPPLIER" in text
    assert "Lanyard @180" in text                            # "@" mid-name is fine
    assert "'5" not in text and ",5," in text                # quantities stay numbers


def test_xlsx_keeps_names_as_text_and_its_own_sums_as_formulas():
    from openpyxl import load_workbook
    daily = {"dates": ["2026-07-01", "2026-07-02"], "selling_days": 2,
             "quantities": {1: {"2026-07-01": 5}, 2: {"2026-07-02": 2}}}
    ws = load_workbook(io.BytesIO(batch_export.render_xlsx(REPORT, "2026-07", daily))).active
    cells = {c.value: c for row in ws.iter_rows() for c in row if c.value is not None}
    name = cells["=HYPERLINK(\"http://x\",\"click\")"]
    assert name.data_type == "s"                             # text, not a formula
    assert cells["+SUPPLIER"].data_type == "s"
    sums = [c for c in cells.values() if isinstance(c.value, str) and c.value.startswith("=SUM(")]
    assert sums and all(c.data_type == "f" for c in sums)   # the sheet's own totals


# -------------------------------------------------------- dependencies


def test_defusedxml_is_pinned_wherever_openpyxl_is():
    for req in (ROOT / "backend" / "requirements.txt", ROOT / "requirements" / "requirements.txt"):
        lines = [ln.split("#")[0].strip() for ln in req.read_text(encoding="utf-8").splitlines()]
        assert any(ln.startswith("openpyxl==") for ln in lines), req
        assert any(ln.startswith("defusedxml==") for ln in lines), req
