"""
app.py — USTore backend (Flask + SQLite), Phase 3 of docs/PROMPT_3_BACKEND.md.

Reads and writes ../ustore.db directly (see db.py for why: it's the same
database the ETL pipeline already builds and populates - this app is not a
second seed path). Endpoint shapes match what
`UST Prototype Design/src/services/dataService.js` already consumes, per
`UST Prototype Design/BACKEND_TODO.md` - so swapping the frontend from
mocks to this API is a same-shape change, not a rewrite.

BIR compliance (see docs/PROMPT_1_FRONTEND.md / docs/PROMPT_3_BACKEND.md §1): this is
an internal inventory counting service only. No checkout, payment, customer
total, or receipt endpoint exists anywhere in this file, and none should be
added.

Run:
    pip install -r requirements.txt
    python app.py
Serves on http://127.0.0.1:5000. The frontend dev server proxies /api to
this port (see vite.config.js) so no CORS is needed in normal dev use;
flask-cors is enabled as a fallback, limited to the dashboard's own
addresses (auth.allowed_origins).

Every /api route except /api/auth/* needs a signed-in session (auth.py).
"""
import csv
import io
import math
import re
import statistics
from datetime import date, datetime, timedelta
from pathlib import Path

from flask import Flask, Response, g, jsonify, request
from flask_cors import CORS

import auth
import batch_export
import batch_pdf
import catalog
import db as dbmod
import decision_rules
import pipeline
import validation
import names                  # after pipeline: it needs scripts/ on sys.path
import order_types            # scripts/order_types.py
import calendar_extent        # scripts/calendar_extent.py
import tally_corrections       # scripts/tally_corrections.py

app = Flask(__name__)
auth.configure(app)
# Only the dashboard may call this API from a browser, and it sends the
# session cookie, hence supports_credentials.
CORS(app, origins=auth.allowed_origins(), supports_credentials=True)

# The plain data files are gitignored; after a fresh clone or a pull they come
# from the encrypted vault/ (scripts/vault.py). Never overwrites local work.
_vault_note = pipeline.vault.ensure_unlocked()
if _vault_note:
    print(_vault_note, flush=True)

# Uploaded workbooks are archived here, so a spreadsheet imported through the
# interface lands in the same folder step0 reads its source workbooks from.
# Gitignored (real client data), and created on demand rather than assumed -
# a clean clone has no rawdata/.
RAWDATA_DIR = Path(__file__).resolve().parent.parent / "rawdata"

# An import is a whole workbook, so the JSON-sized default is far too small,
# but an unbounded upload is a trivial way to fill the disk. The largest file
# in rawdata/ today is ~2 MB.
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024


def _safe_name(name):
    """Filenames are built from user-supplied months/dates, so strip anything
    that could climb out of RAWDATA_DIR before it reaches the filesystem."""
    return re.sub(r"[^A-Za-z0-9._-]", "_", str(name))


def _write_rawdata(filename, payload):
    """Archive an uploaded file into rawdata/ and return the path, or None if
    it could not be written. A failed archive must never fail the import: the
    rows the user wanted loaded are the point, the copy is a side effect."""
    try:
        RAWDATA_DIR.mkdir(parents=True, exist_ok=True)
        path = RAWDATA_DIR / _safe_name(filename)
        path.write_bytes(payload)
        return str(path)
    except OSError as exc:
        app.logger.warning("could not write %s to rawdata/: %s", filename, exc)
        return None


def _norm_header(s):
    """'Units On Hand' / 'units_on_hand' / ' UNITS ON HAND ' -> 'unitsonhand'.
    The store's workbooks are hand-kept and their headers vary in case,
    spacing and punctuation; matching on letters and digits alone is what
    lets one importer read all of them."""
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def _read_table(storage):
    """Read an uploaded .csv/.xlsx into (headers, rows-as-dicts).

    Raises ValueError with a message meant for the user - the interface shows
    it verbatim, so it has to say what to fix, not what went wrong internally.
    """
    name = (storage.filename or "").lower()
    raw = storage.read()
    if not raw:
        raise ValueError("The file is empty.")

    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            # read_only + data_only: we want values, not formulas, and these
            # workbooks can carry thousands of rows of embedded formatting.
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:                       # openpyxl raises a zoo
            raise ValueError(f"Could not read that .xlsx file ({exc}).") from exc
        ws = wb.active
        table = [list(r) for r in ws.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        for encoding in ("utf-8-sig", "utf-8", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("Could not decode that CSV - save it as UTF-8 and retry.")
        table = [row for row in csv.reader(io.StringIO(text))]
    else:
        raise ValueError("Upload a .csv or .xlsx file.")

    table = [r for r in table if any(str(c).strip() for c in r if c is not None)]
    if not table:
        raise ValueError("The file has no rows.")

    headers = [_norm_header(h) for h in table[0]]
    rows = []
    for values in table[1:]:
        row = {}
        for i, h in enumerate(headers):
            if h:
                row[h] = values[i] if i < len(values) else None
        rows.append(row)
    return headers, rows


def _pick(row, *names):
    """First non-empty value among several accepted column spellings."""
    for n in names:
        v = row.get(_norm_header(n))
        if v is not None and str(v).strip() != "":
            return v
    return None


def _as_int(value):
    """Spreadsheet numbers arrive as 5, 5.0 or '5' depending on the cell.
    Returns None when the value is not a whole number."""
    try:
        f = float(str(value).strip().replace(",", ""))
    except (TypeError, ValueError):
        return None
    return int(f) if f == int(f) else None


def _as_iso_date(value):
    """Accept a real date cell, an ISO string, or nothing. Deliberately does
    NOT accept DD/MM or MM/DD: the whole pipeline refuses ambiguous dates for
    the reason set out in the README, and an importer is exactly where a
    silently-misread date would do the most damage."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value or "").strip()[:10]
    return s if validation.is_iso_date(s) else None


def _product_index(c):
    """name_key -> product_id, for matching a spreadsheet's item column: each
    item's own name and every sheet name the vocabulary maps onto it."""
    return names.name_index(c, VOCAB_CSV)


def _lookup(index, item):
    return index.get(names.name_matcher.name_key(item))


def _accepted_upload(storage):
    """Reject anything that is not a spreadsheet BEFORE it is archived.

    rawdata/ is a pipeline input folder - step0 globs it - so letting an
    unparseable upload land there would put junk in front of the ETL. The
    extension check therefore gates the archive, not just the parse."""
    name = (storage.filename or "").lower()
    if not (name.endswith(".csv") or name.endswith(".xlsx")):
        return "Upload a .csv or .xlsx file."
    return None


def _archive_upload(storage):
    """Keep the file exactly as uploaded, under a timestamped name so two
    imports of the same workbook do not overwrite each other."""
    storage.stream.seek(0)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return _write_rawdata(f"import_{stamp}_{storage.filename or 'upload'}",
                          storage.stream.read())


def con():
    if "con" not in g:
        g.con = dbmod.get_db()
    return g.con


def _extend_calendar(c):
    """Make sure Dim_Date runs to the end of next year (scripts/
    calendar_extent.py), so a date staff enter is never refused just because
    the calendar file stops. One MAX() when it already does. Best-effort: if
    the pipeline holds the write lock right now, the date checks downstream
    still say what is wrong, and the next request tries again."""
    try:
        if calendar_extent.extend_dim_date(c):
            c.commit()
    except Exception:                       # noqa: BLE001 - never fail a request over this
        app.logger.exception("could not extend Dim_Date")


_calendar_checked = False


@app.before_request
def _extend_calendar_once():
    """Once per server process, on its first request: covers a server that
    has been running since last year, before anything reads the calendar.
    Not at import time - the test suite imports this module."""
    global _calendar_checked
    if not _calendar_checked and request.path.startswith("/api/"):
        _calendar_checked = True
        _extend_calendar(con())


@app.teardown_appcontext
def _close_con(_exc):
    c = g.pop("con", None)
    if c is not None:
        c.close()


# ------------------------------------------------------------------- auth

# Reachable signed out: the login form needs these three.
PUBLIC_PATHS = {"/api/auth/login", "/api/auth/logout", "/api/auth/session"}


@app.before_request
def _require_login():
    """Every /api route needs a signed-in user, and every write must come from
    the dashboard's own address (see auth.origin_allowed for why CORS alone
    does not cover writes). CORS preflights pass through: they carry no
    cookie and change nothing."""
    if request.method == "OPTIONS" or not request.path.startswith("/api/"):
        return None
    if request.method not in auth.SAFE_METHODS and not auth.origin_allowed(request):
        return jsonify({"ok": False, "error": "This request did not come from the dashboard."}), 403
    if request.path in PUBLIC_PATHS:
        return None
    if not auth.current_user():
        return jsonify({"ok": False, "auth_required": True,
                        "error": "Your session has ended. Sign in again."}), 401
    return None


@app.get("/api/auth/session")
def get_session():
    user = auth.current_user()
    return jsonify({"authenticated": user is not None, "user": user})


@app.post("/api/auth/login")
def login():
    payload = request.get_json(silent=True) or {}
    addr = request.remote_addr or "?"
    wait = auth.locked_for(addr)
    if wait:
        return jsonify({"ok": False, "error": f"Too many wrong attempts. Try again in {wait} seconds."}), 429
    if not str(payload.get("username") or "").strip() or not payload.get("password"):
        return jsonify({"ok": False, "error": "Enter your username and password."}), 400
    user = auth.check_login(con(), payload.get("username"), payload.get("password"), addr)
    if user is None:
        return jsonify({"ok": False, "error": "Wrong username or password."}), 401
    auth.sign_in(user)
    return jsonify({"ok": True, "user": user})


@app.post("/api/auth/logout")
def logout():
    auth.sign_out()
    return jsonify({"ok": True})


def _needs_admin(what):
    """None when this edit may go ahead; otherwise the response refusing it.

    Changing or removing a record that already exists - a stock count, a
    closure, an event, a restock, an upcoming order, an item's name or status -
    needs the admin. The admin signed in does it directly; anyone else sends
    the admin's password in the X-Admin-Password header, which the interface
    asks for in a confirmation dialog when a write comes back `admin_required`
    (services/dataService.js). Adding a new record (a tally, a first count, a
    new closure or event) does not. Checked here, on the server, so the rule
    holds however the request is made. Wrong passwords lock the address out
    like wrong logins do (auth.check_admin_password)."""
    if auth.is_admin(auth.current_user()):
        return None
    addr = request.remote_addr or "?"
    wait = auth.locked_for(f"admin:{addr}")
    if wait:
        return jsonify({"ok": False, "error": f"Too many wrong admin passwords. Try again in {wait} seconds."}), 429
    password = request.headers.get(auth.ADMIN_PASSWORD_HEADER)
    if not password:
        return jsonify({"ok": False, "admin_required": True, "action": what,
                        "error": f"{what} needs the admin's password."}), 403
    if not auth.check_admin_password(con(), password, addr):
        return jsonify({"ok": False, "admin_required": True, "action": what,
                        "error": "That is not the admin password."}), 403
    return None


def _bool_flag(v):
    return request.args.get(v) in ("1", "true", "True")


def _us_date(iso):
    """'2026-10-09' -> '10/09/2026', for dates baked into prose the frontend
    renders as-is. Structured date fields stay ISO - the frontend formats
    those itself (lib/format.js usDate) and anything feeding an API call or an
    <input type="date"> must remain ISO. This is only for the advisory
    `description` strings, which are sentences, not data."""
    try:
        y, m, d = str(iso).split("-")
        return f"{m}/{d}/{y}" if len(y) == 4 else str(iso)
    except (ValueError, AttributeError):
        return str(iso)


# --------------------------------------------------------------- metadata

@app.get("/api/meta")
def get_meta():
    c = con()
    products = dbmod.rows(c, "SELECT product_id, unit_price_php FROM Dim_Product")
    stats, _ = catalog.compute_stats(c)
    dim_date_span = dbmod.one(c, "SELECT MIN(calendar_date) a, MAX(calendar_date) b FROM Dim_Date")
    # The last day the published school calendar (calendar_ranges.csv) covers.
    # Dim_Date runs further (scripts/calendar_extent.py), but past this day it
    # knows no enrollment, exam week or break - staff have to log closures
    # and events themselves, which is what the Tally Interface warns about.
    academic_end = dbmod.one(
        c, "SELECT MAX(calendar_date) d FROM Dim_Date WHERE semester_id IS NOT NULL")["d"]
    sales_span = dbmod.one(c, f"""
        SELECT MIN(d.calendar_date) a, MAX(d.calendar_date) b
        FROM Fact_Sales f JOIN Dim_Date d ON d.date_id = f.date_id
        WHERE {dbmod.SALE_ONLY}
    """)
    fact_sales_rows = c.execute("SELECT COUNT(*) FROM Fact_Sales").fetchone()[0]
    total_units = c.execute(f"SELECT SUM(f.quantity_sold) FROM Fact_Sales f WHERE {dbmod.SALE_ONLY}").fetchone()[0]
    n_params = c.execute("SELECT COUNT(*) FROM Dim_Parameters").fetchone()[0]
    n_reorder = c.execute("SELECT COUNT(*) FROM Result_Prescriptive").fetchone()[0]
    tables = {r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    has_forecast = "Result_Forecast" in tables and c.execute(
        "SELECT COUNT(*) FROM Result_Forecast "
        "WHERE model_type IS NOT 'policy_rate'").fetchone()[0] > 0

    priced = [p for p in products if p["unit_price_php"] is not None]
    unpriced_units = sum(
        stats[p["product_id"]]["total_units"]
        for p in products
        if p["unit_price_php"] is None and p["product_id"] in stats
    )

    return jsonify({
        "generated_at": date.today().isoformat(),
        "source": "ustore.db (live)",
        "sales_span": [sales_span["a"], sales_span["b"]],
        "calendar_span": [dim_date_span["a"], dim_date_span["b"]],
        "academic_calendar_end": academic_end,
        "fact_sales_rows": fact_sales_rows,
        "total_units": total_units,
        "products": len(products),
        "products_with_sales": len(stats),
        "products_with_price": len(priced),
        "unpriced_units": unpriced_units,
        "products_with_stock": sum(1 for s in stats.values() if s["current_stock"] is not None),
        "recent_window_days": 60,
        "available": {
            "sales": True,
            "fsn": True,
            "batch_report": True,
            "forecast": bool(has_forecast),
            "reorder": n_params > 0 and n_reorder > 0,
        },
        # Only populated when the thing is actually pending. This used to be a
        # fixed string blaming a missing cmdstan build, rendered even on
        # databases that had forecasts in them. Both halves were wrong: recent
        # prophet releases ship their own cmdstan, so a missing forecast means
        # step4 has not been run (it is the long, opt-out step in the Tally
        # Interface's pipeline runner), not that it cannot be.
        "pending_reason": {
            "forecast": None if has_forecast else FORECAST_PENDING_REASON,
            "reorder": "Dim_Parameters is empty - no lead times, ordering or holding costs "
                       "have been collected yet (Block 5, the USTore site visit)."
                       if not (n_params > 0 and n_reorder > 0) else None,
        },
    })


@app.get("/api/suppliers")
def get_suppliers():
    c = con()

    # ?year=YYYY, ?month=YYYY-MM, or ?start=&end=YYYY-MM-DD — only suppliers
    # with a sale in that period, for the batch report's filter. Taken from
    # build_batch_report()'s own output rather than a query of its own, so the
    # dropdown cannot offer a supplier whose section that report would then
    # come back empty for. An unparseable period is ignored rather than
    # rejected: a filter list is not worth a failed screen.
    if request.args.get("year") or request.args.get("month") or request.args.get("start"):
        period, error = _batch_period()
        if not error:
            sold = sorted({e["supplier"] for e in build_batch_report(c, period)})
            return jsonify([catalog.ALL_SUPPLIERS, *sold])

    cat = catalog.compute_catalog(c)
    # ?forecastable=1 — only suppliers with at least one SKU the model actually
    # forecast. The Demand Forecast screen lists nothing but forecastable SKUs,
    # so offering it a supplier whose items are all unforecast leads straight
    # to an empty page.
    if _bool_flag("forecastable"):
        ids = _forecastable_ids(c)
        # An empty set means the pipeline has not produced any forecast at all.
        # That screen then falls back to every SKU with sales history, so the
        # supplier list must not narrow to nothing underneath it.
        if ids:
            cat = [p for p in cat if p["product_id"] in ids]
    return jsonify([catalog.ALL_SUPPLIERS, *sorted({p["supplier_name"] for p in cat})])


@app.get("/api/categories")
def get_categories():
    cat = catalog.compute_catalog(con())
    return jsonify([catalog.ALL_CATEGORIES, *sorted({p["category"] for p in cat})])


@app.get("/api/months")
def get_months():
    return jsonify(catalog.months_seen(con()))


# ---------------------------------------------------------------- catalog

@app.get("/api/products")
def get_products():
    c = con()
    supplier = request.args.get("supplier") or None
    category = request.args.get("category") or None
    # The topbar's date range belongs here, not only on /api/sales/monthly.
    # Every product-derived panel (total units, top sellers, category mix)
    # reads this endpoint, so without it "Last 3 Months" moved the trend line
    # and left every other figure on the page at its all-time value.
    date_range = request.args.get("dateRange") or None
    cat = catalog.compute_catalog(c, date_range,
                                  request.args.get("start"), request.args.get("end"))
    cat = [p for p in cat if catalog.matches(p, supplier, category)]
    if _bool_flag("has_history"):
        cat = [p for p in cat if p["is_active"] and p["total_units"] >= 0 and p["has_history"]]
        cat.sort(key=lambda p: p["item_name"])
    return jsonify(cat)


DATA_DIR = Path(__file__).resolve().parent.parent / "data"
VOCAB_CSV = DATA_DIR / "vocab_mapping_FINAL_v5.csv"
INVENTORY_SOURCE_CSV = DATA_DIR / "USTore_inventory_excel_long.csv"


@app.post("/api/products")
def add_product():
    """Create one new item, permanently.

    Writes to three places, because anything less does not survive:

      - `Dim_Product`, so the item is usable in this session immediately;
      - `data/vocab_mapping_FINAL_v5.csv`, the controlled vocabulary;
      - `data/USTore_inventory_excel_long.csv`, the inventory source, which is
        what `step1_apply_mapping.py` actually derives the product roster from.

    Writing to the vocabulary is a deliberate, user-directed exception to the
    rule that it is a hand-maintained file. The two CSV appends happen BEFORE
    the database insert so that a filesystem failure leaves nothing behind: a
    row in Dim_Product that no CSV backs would vanish at the next pipeline run
    and look like data loss.
    """
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("item_name") or "").strip()
    category = str(payload.get("category") or "").strip() or "Uncategorised"
    supplier = str(payload.get("supplier_name") or "").strip() or None

    errors = {}
    if not name:
        errors["item_name"] = "Enter an item name."
    elif len(name) > 200:
        errors["item_name"] = "Item name is too long."

    c = con()
    if not errors and c.execute(
        "SELECT 1 FROM Dim_Product WHERE LOWER(TRIM(item_name)) = ?", (name.lower(),)
    ).fetchone():
        errors["item_name"] = "An item with that name already exists."
    if errors:
        return jsonify({"ok": False, "errors": errors}), 400

    # CSVs first: see the docstring. If these fail the item is never created,
    # rather than created in a form that the next pipeline run would erase.
    try:
        names.register_new_item(VOCAB_CSV, INVENTORY_SOURCE_CSV, name, category,
                                payload.get("unit_price_php"))
    except (OSError, StopIteration, csv.Error) as exc:
        app.logger.exception("could not register %r in the source CSVs", name)
        return jsonify({"ok": False, "errors": {
            "item_name": f"Could not write to the data files ({exc}). Item not added."
        }}), 500

    # fsn_class and lead_time_days are left NULL on purpose: step3 classifies
    # and step5a sets lead times. Guessing either here would put a number in
    # front of the prescriptive math that no rule produced.
    cur = c.execute("""
        INSERT INTO Dim_Product (item_name, category, supplier_name, entry_date, is_active, is_hvl)
        VALUES (?, ?, ?, ?, 1, 0)
    """, (name, category, supplier, date.today().isoformat()))
    c.commit()

    return jsonify({"ok": True, "product": {
        "product_id": cur.lastrowid,
        "item_name": name,
        "category": category,
        "supplier_name": supplier or catalog.UNATTRIBUTED,
    }})


@app.post("/api/products/<int:product_id>/rename")
def rename_product(product_id):
    """The tally sheet now calls this item something else ({"new_name"}).

    Adds the new name to the vocabulary as another name for the item, so the
    next pipeline run reads rows under it as this item: history kept, name on
    the dashboard unchanged (see names.rename)."""
    gate = _needs_admin("Renaming an item")
    if gate:
        return gate
    payload = request.get_json(silent=True) or {}
    try:
        result = names.rename(con(), VOCAB_CSV, product_id, payload.get("new_name"))
    except names.ReviewError as exc:
        return jsonify({"ok": False, "error": str(exc)}), exc.status
    except (OSError, StopIteration, csv.Error) as exc:
        app.logger.exception("could not write the vocabulary for a rename")
        return jsonify({"ok": False, "error": f"Could not write to the vocabulary ({exc})."}), 500
    return jsonify({"ok": True, **result})


@app.get("/api/names/review")
def get_name_review():
    """Item names the vocabulary does not know yet, from the tally sheets
    (provisional items, step1) and from imported files (held rows), with a
    suggestion where name_matcher has one. See backend/names.py."""
    try:
        return jsonify(names.review_list(con(), VOCAB_CSV))
    except OSError as exc:
        return jsonify({"ok": False, "error": f"Could not read the vocabulary ({exc})."}), 500


@app.post("/api/names/review")
def settle_name():
    """Settle one name: {"raw_name", "action": "same" | "new" | "discard",
    "product_id" (same), "category"/"supplier_name" (new, import-only)}.
    Appends one row to the vocabulary; see names.settle."""
    payload = request.get_json(silent=True) or {}
    try:
        result = names.settle(con(), VOCAB_CSV, INVENTORY_SOURCE_CSV,
                              payload.get("raw_name"), payload.get("action"),
                              product_id=payload.get("product_id"),
                              category=payload.get("category"),
                              supplier_name=payload.get("supplier_name"))
    except names.ReviewError as exc:
        return jsonify({"ok": False, "error": str(exc)}), exc.status
    except (OSError, StopIteration, csv.Error) as exc:
        app.logger.exception("could not write the vocabulary for a name review")
        return jsonify({"ok": False, "error": f"Could not write to the data files ({exc})."}), 500
    return jsonify({"ok": True, **result})


@app.put("/api/products/<int:product_id>/status")
def set_product_status(product_id):
    """Mark an item discontinued, or back in stock ({"discontinued": bool}).

    Stored in Product_Status by item name (see backend/db.py for why), so it
    survives a pipeline run. A discontinued item leaves the tally item
    pickers; its history, counts and stock figures are untouched."""
    gate = _needs_admin("Changing an item's status")
    if gate:
        return gate
    payload = request.get_json(silent=True) or {}
    if not isinstance(payload.get("discontinued"), bool):
        return jsonify({"ok": False, "error": "Send discontinued: true or false."}), 400
    c = con()
    product = dbmod.one(c, "SELECT item_name FROM Dim_Product WHERE product_id = ?", (product_id,))
    if not product:
        return jsonify({"ok": False, "error": "No such item."}), 404
    c.execute(dbmod.PRODUCT_STATUS_DDL)   # in case startup could not create it
    c.execute("""
        INSERT INTO Product_Status (item_name, discontinued, changed_at) VALUES (?, ?, ?)
        ON CONFLICT (item_name) DO UPDATE SET
            discontinued = excluded.discontinued, changed_at = excluded.changed_at
    """, (product["item_name"], int(payload["discontinued"]),
          datetime.now().isoformat(timespec="seconds")))
    c.commit()
    return jsonify({"ok": True, "product_id": product_id, "item_name": product["item_name"],
                    "discontinued": payload["discontinued"]})


@app.get("/api/products/<int:product_id>/history")
def get_product_history(product_id):
    c = con()
    rows = dbmod.rows(c, f"""
        SELECT substr(d.calendar_date, 1, 7) AS month,
               SUM(f.quantity_sold) AS units, COUNT(DISTINCT f.date_id) AS tally_days
        FROM Fact_Sales f JOIN Dim_Date d ON d.date_id = f.date_id
        WHERE f.product_id = ? AND {dbmod.SALE_ONLY}
        GROUP BY 1 ORDER BY 1
    """, (product_id,))
    return jsonify(rows)


# ------------------------------------------------------------------ sales

@app.get("/api/sales/monthly")
def get_monthly_units():
    c = con()
    cat = catalog.compute_catalog(c)
    supplier = request.args.get("supplier") or None
    category = request.args.get("category") or None
    date_range = request.args.get("dateRange") or None
    start, end = request.args.get("start"), request.args.get("end")
    keep = {p["product_id"] for p in cat if catalog.matches(p, supplier, category)}
    price_by_id = {p["product_id"]: p["unit_price_php"] for p in cat}
    months = catalog.months_seen(c)

    monthly = dbmod.rows(c, f"""
        SELECT f.product_id, substr(d.calendar_date, 1, 7) AS month, SUM(f.quantity_sold) AS units
        FROM Fact_Sales f JOIN Dim_Date d ON d.date_id = f.date_id
        WHERE {dbmod.SALE_ONLY}
        GROUP BY 1, 2
    """)
    by_month = {}
    for row in monthly:
        if row["product_id"] not in keep or not catalog.in_range(row["month"], date_range, months, start, end):
            continue
        acc = by_month.setdefault(row["month"], {"month": row["month"], "units": 0, "revenue": 0.0, "priced_units": 0})
        acc["units"] += row["units"]
        price = price_by_id.get(row["product_id"])
        if price is not None:
            acc["revenue"] += row["units"] * price
            acc["priced_units"] += row["units"]
    return jsonify(sorted(by_month.values(), key=lambda a: a["month"]))


# How many trailing days of daily sales the Overview's performance line shows,
# and the window its top earners are ranked over.
PERFORMANCE_DAILY_DAYS = 90
TOP_EARNER_WINDOW_DAYS = 365
TOP_EARNER_SHARE = 0.80          # list items down to this share of sales
SMALL_SELLER_UNITS = 10          # "sold this many or fewer" = negligible
STOPPED_MIN_UNITS = 5            # sold this many last month, none this month
CATEGORY_DRIVERS = 2             # items named behind each category's change


def _month_before(month, years=0, months=1):
    y, m = int(month[:4]) - years, int(month[5:7]) - months
    while m < 1:
        y, m = y - 1, m + 12
    return f"{y:04d}-{m:02d}"


@app.get("/api/overview/performance")
def get_overview_performance():
    """The Overview's story in pesos: how the store is doing, and which
    products drive the sales.

    Summed here from Fact_Sales x Dim_Product's unit price (unpriced items left
    out; `priced_unit_share` says how much that is).

    Every date is anchored on the last day any sale was tallied, store-wide,
    not on today: the latest month is the latest TALLIED month, and a filter
    cannot move it. Supplier / category narrow the sums.

    `by_category` splits the latest and previous months by category, so the
    Overview can say where sales moved rather than only that they did, and
    names the items that moved each one most. `stopped_selling` lists the
    items that sold last month and not once this month: one that suddenly
    stops has often run out, so they are the ones to check first."""
    c = con()
    cat = catalog.compute_catalog(c)
    supplier = request.args.get("supplier") or None
    category = request.args.get("category") or None
    by_id = {p["product_id"]: p for p in cat}
    keep = {p["product_id"] for p in cat if catalog.matches(p, supplier, category)}

    rows = dbmod.rows(c, f"""
        SELECT f.product_id, d.calendar_date AS day, SUM(f.quantity_sold) AS units
        FROM Fact_Sales f JOIN Dim_Date d ON d.date_id = f.date_id
        WHERE {dbmod.SALE_ONLY} AND f.quantity_sold > 0
        GROUP BY 1, 2
    """)
    if not rows:
        return jsonify({"available": False, "reason": "No sales have been tallied yet."})
    history_end = max(r["day"] for r in rows)
    end = date.fromisoformat(history_end)
    window_start = (end - timedelta(days=TOP_EARNER_WINDOW_DAYS - 1)).isoformat()
    daily_start = (end - timedelta(days=PERFORMANCE_DAILY_DAYS - 1)).isoformat()
    latest = history_end[:7]
    previous = _month_before(latest)

    daily, monthly, items, by_category, by_item = {}, {}, {}, {}, {}
    units_all = units_priced = 0
    for r in rows:
        pid = r["product_id"]
        if pid not in keep:
            continue
        price = (by_id.get(pid) or {}).get("unit_price_php")
        month = r["day"][:7]
        m = monthly.setdefault(month, {"month": month, "php": 0.0, "units": 0})
        m["units"] += r["units"]
        if month in (latest, previous):
            bi = by_item.setdefault(pid, {"units": 0, "previous_units": 0, "php": 0.0, "previous_php": 0.0})
            bi["units" if month == latest else "previous_units"] += r["units"]
        if r["day"] >= window_start:
            units_all += r["units"]
        if price is None:
            continue
        php = r["units"] * price
        m["php"] += php
        if month in (latest, previous):
            name = (by_id.get(pid) or {}).get("category") or "Uncategorised"
            bc = by_category.setdefault(name, {"category": name, "php": 0.0, "previous_php": 0.0})
            bc["php" if month == latest else "previous_php"] += php
            by_item[pid]["php" if month == latest else "previous_php"] += php
        if r["day"] >= daily_start:
            daily[r["day"]] = daily.get(r["day"], 0.0) + php
        if r["day"] >= window_start:
            units_priced += r["units"]
            it = items.setdefault(pid, {"php": 0.0, "units": 0})
            it["php"] += php
            it["units"] += r["units"]

    def month_figure(mo):
        m = monthly.get(mo)
        return None if not m else {"month": mo, "php": round(m["php"], 2), "units": m["units"]}

    # The latest month is complete only if the tallies run to its last day.
    month_end = (date(int(latest[:4]) + (latest[5:] == "12"), int(latest[5:]) % 12 + 1, 1)
                 - timedelta(days=1))
    days = [(end - timedelta(days=i)).isoformat() for i in range(PERFORMANCE_DAILY_DAYS - 1, -1, -1)]

    ranked = sorted(items.items(), key=lambda kv: -kv[1]["php"])
    total = sum(v["php"] for _, v in ranked) or 0.0
    top, cum = [], 0.0
    for pid, v in ranked:
        cum += v["php"]
        p = by_id.get(pid, {})
        top.append({
            "product_id": pid, "item_name": p.get("item_name"), "category": p.get("category"),
            "supplier_name": p.get("supplier_name"), "fsn_class": p.get("fsn_class"),
            "units": v["units"], "php": round(v["php"], 2),
            "share": v["php"] / total if total else 0.0, "cum_share": cum / total if total else 0.0,
        })
    n_for_share = next((i + 1 for i, t in enumerate(top) if t["cum_share"] >= TOP_EARNER_SHARE), len(top))
    small = [t for t in top if t["units"] <= SMALL_SELLER_UNITS]

    def category_of(pid):
        return (by_id.get(pid) or {}).get("category") or "Uncategorised"

    def drivers(category, diff):
        """The items that moved `category` furthest the way it moved."""
        moved = [(pid, v["php"] - v["previous_php"]) for pid, v in by_item.items()
                 if category_of(pid) == category]
        moved = sorted((t for t in moved if t[1] * diff > 0), key=lambda t: -abs(t[1]))
        return [{"item_name": by_id[pid]["item_name"], "diff": round(d, 2)}
                for pid, d in moved[:CATEGORY_DRIVERS]]

    stopped = sorted(
        ({"product_id": pid, "item_name": by_id[pid]["item_name"], "category": category_of(pid),
          "previous_units": v["previous_units"], "previous_php": round(v["previous_php"], 2)}
         for pid, v in by_item.items()
         if v["previous_units"] >= STOPPED_MIN_UNITS and v["units"] == 0),
        key=lambda s: (-s["previous_php"], -s["previous_units"]))

    return jsonify({
        "available": True,
        "history_end": history_end,
        "latest_month": month_figure(latest),
        "latest_month_complete": end == month_end,
        "previous_month": month_figure(previous),
        "same_month_last_year": month_figure(_month_before(latest, years=1, months=0)),
        "by_category": sorted(
            ({"category": v["category"], "php": round(v["php"], 2),
              "previous_php": round(v["previous_php"], 2),
              "drivers": drivers(v["category"], v["php"] - v["previous_php"])}
             for v in by_category.values()),
            key=lambda v: -v["php"]),
        "stopped_selling": {"min_units": STOPPED_MIN_UNITS, "items": stopped},
        "daily": [{"date": d, "php": round(daily.get(d, 0.0), 2)} for d in days],
        "priced_unit_share": units_priced / units_all if units_all else None,
        "top": {
            "window_start": window_start, "window_end": history_end,
            "total_php": round(total, 2), "n_selling": len(top),
            "n_for_share": n_for_share, "share": TOP_EARNER_SHARE,
            "top10_share": top[min(10, len(top)) - 1]["cum_share"] if top else 0.0,
            "items": top[:max(n_for_share, 10)],
        },
        "small_sellers": {
            "max_units": SMALL_SELLER_UNITS, "count": len(small),
            "share": sum(t["php"] for t in small) / total if total else 0.0,
        },
    })


# A batch report covers a year, a month, one day or one week. A span of days
# stops at a month: its grid has a column per day, and past a month it no
# longer fits a page at a readable size. A year's grid has a column per month.
BATCH_MAX_DAYS = 31
YEAR_RE = re.compile(r"^\d{4}$")


def _batch_period():
    """The report period the request asks for, as (period, error).

    ?year=YYYY is a year; ?month=YYYY-MM a month; ?start=&end=YYYY-MM-DD a
    span of days (one day is start == end, a week Monday to Sunday). The
    period is what batch_pdf.period_span() and the renderers take: the year
    or month string, or the (start, end) pair."""
    start, end = request.args.get("start"), request.args.get("end")
    if start or end:
        try:
            if not (validation.ISO_DATE_RE.match(start or "") and validation.ISO_DATE_RE.match(end or "")):
                raise ValueError
            first, last = date.fromisoformat(start), date.fromisoformat(end)
        except ValueError:
            return None, "start= and end= must both be dates, YYYY-MM-DD."
        if last < first:
            return None, "end= is before start=."
        if (last - first).days >= BATCH_MAX_DAYS:
            return None, f"A batch report covers at most {BATCH_MAX_DAYS} days; use month= for a month."
        return (start, end), None
    year = request.args.get("year")
    if year is not None:
        if not YEAR_RE.match(year):
            return None, "year= must be a year, YYYY."
        return year, None
    month = request.args.get("month")
    if not month or not validation.ISO_MONTH_RE.match(month):
        return None, "year=YYYY, month=YYYY-MM, or start= and end=YYYY-MM-DD, is required."
    return month, None


def _period_error(message):
    return jsonify({"ok": False, "errors": {"period": message}}), 400


def build_batch_report(c, period, only_supplier=None):
    """The period's sales grouped by supplier — the shared body of both
    /api/reports/batch (JSON, for the screen) and /api/reports/batch.pdf
    (the document Purchasing and Finance actually receive). Neither may
    compute this differently from the other. `period` is a month
    ('YYYY-MM'), a year ('YYYY') or a (start, end) span of days; see
    _batch_period().

    `only_supplier` narrows the report to that supplier's section. It is applied
    here, in the shared body, rather than on the screen: a filter that trimmed
    the table but not the PDF/CSV/XLSX beside it would hand someone an export
    that disagrees with what they were looking at when they clicked."""
    cat = catalog.compute_catalog(c)
    product_by_id = {p["product_id"]: p for p in cat}

    totals = dbmod.rows(c, f"""
        SELECT f.product_id, SUM(f.quantity_sold) AS units
        FROM Fact_Sales f JOIN Dim_Date d ON d.date_id = f.date_id
        WHERE d.calendar_date BETWEEN ? AND ? AND {dbmod.SALE_ONLY}
        GROUP BY 1
    """, batch_pdf.period_span(period))

    by_supplier = {}
    for row in totals:
        if row["units"] <= 0:
            continue
        p = product_by_id.get(row["product_id"])
        if not p:
            continue
        supplier = p["supplier_name"] or catalog.UNATTRIBUTED
        entry = by_supplier.setdefault(supplier, {
            "supplier": supplier, "items": [], "total_units": 0, "subtotal": 0.0, "unpriced_units": 0,
        })
        line_total = row["units"] * p["unit_price_php"] if p["unit_price_php"] is not None else None
        entry["items"].append({
            # product_id travels with the line so the XLSX renderer can join
            # each row to its per-date quantities. Two suppliers can carry the
            # same item_name, so the name is not a key.
            "product_id": row["product_id"],
            "item_name": p["item_name"],
            "quantity": row["units"],
            "unit_price_php": p["unit_price_php"],
            "line_total": line_total,
        })
        entry["total_units"] += row["units"]
        if line_total is not None:
            entry["subtotal"] += line_total
        else:
            entry["unpriced_units"] += row["units"]

    out = []
    for entry in by_supplier.values():
        # Named `only_supplier`, not `supplier`: the loop above binds `supplier`
        # to each row's own supplier, so a parameter of that name is shadowed
        # by the time this runs and every section but the last one's is dropped.
        if (only_supplier and only_supplier != catalog.ALL_SUPPLIERS
                and entry["supplier"] != only_supplier):
            continue
        entry["items"].sort(key=lambda i: i["quantity"], reverse=True)
        out.append(entry)
    out.sort(key=lambda e: e["subtotal"], reverse=True)
    return out


def build_batch_daily(c, period):
    """The period laid out day by day, for the TBS-shaped .xlsx export.

    `build_batch_report` aggregates each item to a period total, which is
    what the screen, the PDF and the CSV show. The store's own TBS workbook
    is a grid instead - one column per calendar day - so this returns the
    same Fact_Sales rows at day resolution. The two cannot disagree: the
    row totals are SUM() formulas over these columns, and both read the
    same table.

    Dates come from Dim_Date rather than from Python's calendar so the
    column set is the pipeline's own calendar, including the closure flags
    the Tally Interface writes.
    """
    days = dbmod.rows(c, """
        SELECT calendar_date, is_store_closed, is_event_day,
               is_enrollment_period, is_exam_week, is_sem_break
        FROM Dim_Date
        WHERE calendar_date BETWEEN ? AND ?
        ORDER BY calendar_date
    """, batch_pdf.period_span(period))

    qty_rows = dbmod.rows(c, f"""
        SELECT f.product_id, d.calendar_date AS day, SUM(f.quantity_sold) AS units
        FROM Fact_Sales f JOIN Dim_Date d ON d.date_id = f.date_id
        WHERE d.calendar_date BETWEEN ? AND ? AND {dbmod.SALE_ONLY}
        GROUP BY 1, 2
        HAVING SUM(f.quantity_sold) > 0
    """, batch_pdf.period_span(period))

    by_product = {}
    for row in qty_rows:
        by_product.setdefault(row["product_id"], {})[row["day"]] = row["units"]

    # One mark per day, so the export can shade the date columns the way the
    # store colour-codes theirs. First match wins: a closed day is a closed
    # day whatever else the calendar says about it, and an event during
    # enrollment is read as the event. These are exactly the Dim_Date
    # regressors the forecast is fitted on (see step4_forecast_model.py), so
    # the sheet shows the same calendar the model sees.
    MARKS = (("is_store_closed", "closed"), ("is_event_day", "event"),
             ("is_enrollment_period", "enrollment"), ("is_exam_week", "exam"),
             ("is_sem_break", "break"))

    def mark(day):
        for column, name in MARKS:
            if day[column]:
                return name
        return None

    marks = {d["calendar_date"]: mark(d) for d in days}
    columns = [d["calendar_date"] for d in days]

    # A year is laid out a column per month, each the month's total. The
    # marks stay per day: no month column is shaded, and the legend still
    # counts the year's closed, event and exam days.
    if batch_pdf.is_year(period):
        columns = [f"{period}-{m:02d}" for m in range(1, 13)]
        for product, per_day in by_product.items():
            per_month = {}
            for day, units in per_day.items():
                per_month[day[:7]] = per_month.get(day[:7], 0) + units
            by_product[product] = per_month

    return {
        "dates": columns,
        "marks": {d: m for d, m in marks.items() if m},
        "quantities": by_product,
        # The store writes this in the legend block as "N DAYS". It counts
        # days that actually sold, not days the store was open - an open day
        # with no recorded sale is not a selling day on their sheet either.
        "selling_days": len({r["day"] for r in qty_rows}),
    }


@app.get("/api/reports/batch")
def get_batch_report():
    period, error = _batch_period()
    if error:
        return _period_error(error)
    return jsonify(build_batch_report(con(), period, request.args.get("supplier") or None))


@app.get("/api/reports/batch.pdf")
def get_batch_report_pdf():
    """The batch sales report as a real PDF file.

    Rendered server-side with **fpdf2**, chosen because it is pure Python with
    no system libraries behind it. This is what previously blocked the feature:
    weasyprint needs GTK/Pango/Cairo installed separately on Windows, and
    reportlab pulls a C extension - neither is something a store's machine
    should have to acquire to print a monthly report. fpdf2 installs from a
    plain wheel and was already present in this environment.

    `?inline=1` serves it for viewing in the browser (the Print Preview
    button); the default is a download (Export as PDF).
    """
    period, error = _batch_period()
    if error:
        return _period_error(error)

    c = con()
    report = build_batch_report(c, period, request.args.get("supplier") or None)
    pdf_bytes = batch_pdf.render(report, period, build_batch_daily(c, period))

    disposition = "inline" if _bool_flag("inline") else "attachment"
    filename = f"USTore_Batch_Sales_Report_{batch_pdf.period_slug(period)}.pdf"
    return Response(pdf_bytes, mimetype="application/pdf", headers={
        "Content-Disposition": f'{disposition}; filename="{filename}"',
        "Content-Length": str(len(pdf_bytes)),
    })


@app.get("/api/reports/batch.csv")
def get_batch_report_csv():
    """The batch sales report as CSV, in the same row shape as the store's
    own TBS workbooks (see batch_export.py). Same build_batch_report() body
    as the PDF and the screen - never a different number."""
    period, error = _batch_period()
    if error:
        return _period_error(error)

    report = build_batch_report(con(), period, request.args.get("supplier") or None)
    csv_bytes = batch_export.render_csv(report, period)
    filename = f"USTore_Batch_Sales_Report_{batch_pdf.period_slug(period)}.csv"
    return Response(csv_bytes, mimetype="text/csv", headers={
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Content-Length": str(len(csv_bytes)),
    })


@app.get("/api/reports/batch.xlsx")
def get_batch_report_xlsx():
    """The batch sales report as an .xlsx workbook, laid out as the store's
    own TBS sheet: supplier blocks, one column per calendar day, TOTAL
    QUANTITY / ITEM PRICE / FOR REMITTANCE. Same build_batch_report() body
    as the CSV, PDF and screen, plus build_batch_daily() for the grid."""
    period, error = _batch_period()
    if error:
        return _period_error(error)

    c = con()
    report = build_batch_report(c, period, request.args.get("supplier") or None)
    xlsx_bytes = batch_export.render_xlsx(report, period, build_batch_daily(c, period))
    filename = f"USTore_Batch_Sales_Report_{batch_pdf.period_slug(period)}.xlsx"
    return Response(xlsx_bytes,
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={
                        "Content-Disposition": f'attachment; filename="{filename}"',
                        "Content-Length": str(len(xlsx_bytes)),
                    })


# -------------------------------------------------------------------- FSN

@app.get("/api/fsn/sensitivity")
def get_fsn_sensitivity():
    return jsonify(catalog.fsn_sensitivity(con()))


@app.get("/api/fsn/decision-rules")
def get_fsn_decision_rules():
    """What to do about each SKU's FSN class (manuscript Sec. 3.3.1), with the
    per-semester velocity the "Review for Discontinuation" rule was judged on.
    ?threshold= (units per selling day) overrides the default."""
    threshold = decision_rules.parse_threshold(request.args.get("threshold"))
    return jsonify(decision_rules.evaluate(con(), threshold))


# --------------------------------------------------------------- reorder

# One tally/billing cycle. Matches the 30-day horizon the forecast targets
# (§3.3.2) and the monthly supplier-remittance cycle the store already runs on,
# so "enough to last until the next count" means the same thing in both places.
REVIEW_PERIOD_DAYS = 30

# The forecast -> prescriptive contract's three rate states, as the API sees
# them. Only the third needs a name here: the screen has to render "we do not
# know" differently from "you have enough", and before the contract those two
# were the same pixel. See docs/PRESCRIPTIVE_CONTRACT.md.
RATE_INSUFFICIENT = "insufficient_data"
POLICY_STATE_PRICED = "priced"
POLICY_STATE_INSUFFICIENT = "insufficient_data"
TIER_NOT_STOCKABLE = "not_stockable"
NOT_STOCKABLE_NOTE = (
    "Made-to-order candidate. This item's own history shows that every extra "
    "unit of stock returns less service than the store's efficiency floor, so "
    "the reorder point covers normal demand and buys no safety stock. Low "
    "availability here is a deliberate choice, not an oversight - if it must be "
    "available on demand, it needs a standing order, not a bigger buffer."
)
INSUFFICIENT_DATA_NOTE = (
    "Insufficient data - manual review. This item has no demand in the trailing "
    "365 days, so the pipeline cannot estimate a rate for it and deliberately "
    "emits no reorder point. A blank here means \"not known\", NOT \"stock is "
    "fine\": first-season designs, one-off drops and discontinued lines all land "
    "here, and they need a person to decide, not a number."
)

ORDER_QTY_NOTE = (
    "Suggested quantity brings stock up to the reorder point plus "
    f"{REVIEW_PERIOD_DAYS} days of demand at the observed rate. It is deliberately "
    "not EOQ: under the provisional ordering/holding costs (Block 5/B9, still "
    "estimates), EOQ works out larger than a whole year of demand for nearly "
    "every SKU, so ordering it would mean buying years of stock. EOQ stays in "
    "the table below as the theoretical figure, flagged where it exceeds annual "
    "demand."
)


@app.get("/api/reorder")
def get_reorder():
    items = _reorder_items(con())
    if items is None:
        return jsonify({
            "available": False,
            "reason": "Dim_Parameters is empty - no lead times, ordering or holding costs "
                      "have been collected yet (Block 5, the USTore site visit).",
            "data": None,
        })

    priced = [i for i in items if i["policy_state"] == POLICY_STATE_PRICED]
    flagged = [i for i in items if i["policy_state"] == POLICY_STATE_INSUFFICIENT]
    due = [i for i in items if i["needs_reorder"]]
    eoq_over = sum(
        1 for i in priced if i["scenarios"].get("low_admin_cost", {}).get("exceeds_annual_demand")
    )

    return jsonify({
        "available": True,
        "reason": None,
        "data": {
            "items": items,
            "summary": {
                "priced_skus": len(priced),
                # Reported, not hidden: these SKUs are IN the table and on the
                # screen, carrying a state instead of a number.
                "insufficient_data_skus": len(flagged),
                "not_stockable_skus": sum(
                    1 for i in priced if i.get("service_tier") == TIER_NOT_STOCKABLE),
                "insufficient_data_note": INSUFFICIENT_DATA_NOTE,
                "total_skus": len(items),
                "with_stock_count": sum(1 for i in items if i["current_stock"] is not None),
                "no_stock_count": sum(1 for i in items if i["current_stock"] is None),
                "reorder_now": len(due),
                "approaching_rop": sum(1 for i in items if i["approaching_rop"]),
                "suggested_units_total": sum(i["suggested_order_qty"] for i in due),
                "reorder_status_unknown": len(flagged),
                "suppliers_affected": len({i["supplier_name"] for i in due if i["supplier_name"]}),
                "review_period_days": REVIEW_PERIOD_DAYS,
                "order_qty_note": ORDER_QTY_NOTE,
                "eoq_exceeding_annual_demand": eoq_over,
            },
        },
    })


def _reorder_items(c):
    """Every SKU in Result_Prescriptive with its stock position and what to do
    about it, sorted by name - or None when the prescriptive step has nothing
    (Dim_Parameters empty). Shared by /api/reorder and the forecast guidance,
    so the two can never suggest different order quantities for one item."""
    n_params = c.execute("SELECT COUNT(*) FROM Dim_Parameters").fetchone()[0]
    reorder_rows = dbmod.rows(c, """
        SELECT rp.*, p.item_name, p.supplier_name
        FROM Result_Prescriptive rp
        JOIN Dim_Product p ON p.product_id = rp.product_id
        ORDER BY p.item_name, rp.ordering_cost_scenario
    """)
    if n_params == 0 or not reorder_rows:
        return None

    by_product = {}
    for r in reorder_rows:
        pid = r["product_id"]
        item = by_product.setdefault(pid, {
            "product_id": pid,
            "item_name": r["item_name"],
            "supplier_name": r["supplier_name"],
            "fsn_class": r["fsn_class"],
            "lead_time_days": r["lead_time_days"],
            "lead_time_category": r["lead_time_category"],
            "avg_daily_demand": r["avg_daily_demand"],
            "annual_demand": r["annual_demand"],
            "sigma_demand": r["sigma_demand"],
            "sigma_source": r["sigma_source"],
            "z_value": r["z_value"],
            "safety_stock": r["safety_stock"],
            "reorder_point": r["reorder_point"],
            "demand_method": r["demand_method"],
            # The forecast -> prescriptive contract. rate_source is what the
            # screen must branch on: 'insufficient_data' means this SKU has
            # no learnable demand rate and carries NO reorder point. See
            # docs/PRESCRIPTIVE_CONTRACT.md.
            "rate_source": r.get("rate_source"),
            "service_tier": r.get("service_tier"),
            "buffer_quantile": r.get("buffer_quantile"),
            "buffer_source": r.get("buffer_source"),
            "safety_stock_normal_legacy": r.get("safety_stock_normal_legacy"),
            "is_provisional": bool(r["is_provisional"]),
            "scenarios": {},
        })
        item["scenarios"][r["ordering_cost_scenario"]] = {
            "ordering_cost_php": r["ordering_cost_php"],
            "holding_cost_php_per_unit_year": r["holding_cost_php_per_unit_year"],
            "eoq": r["eoq"],
            "cost_at_eoq": r["cost_at_eoq"],
            "cost_at_half_eoq": r["cost_at_half_eoq"],
            "cost_at_double_eoq": r["cost_at_double_eoq"],
        }

    # ---- what to actually do about it -------------------------------
    # Joining stock here rather than leaving the frontend to do it: the
    # order quantity below has to be derived from on-hand, and a screen must
    # not invent an analytic the backend has not computed.
    stats, _ = catalog.compute_stats(c)

    for item in by_product.values():
        st = stats.get(item["product_id"], {})
        stock = st.get("current_stock")

        item["current_stock"] = stock
        item["stock_as_of"] = st.get("stock_as_of")
        item["stock_source"] = st.get("stock_source")

        # A SKU with no learnable demand rate has NO reorder point, and the
        # old `item["reorder_point"] or 0.0` turned that NULL into 0.0 - which
        # then rendered as needs_reorder=False, i.e. "you have enough stock".
        # That is a recommendation the pipeline has no evidence for, and it is
        # exactly the silent zero the contract exists to prevent. These SKUs
        # get an explicit state instead, and every downstream number stays
        # None rather than being computed off a fabricated zero.
        if item.get("rate_source") == RATE_INSUFFICIENT or item["reorder_point"] is None:
            item["rate_source"] = RATE_INSUFFICIENT
            item["policy_state"] = POLICY_STATE_INSUFFICIENT
            item["policy_note"] = INSUFFICIENT_DATA_NOTE
            item["days_cover_remaining"] = None
            item["needs_reorder"] = None
            item["approaching_rop"] = None
            item["order_up_to_level"] = None
            item["suggested_order_qty"] = None
            item["suggested_order_min"] = None
            for scen in item["scenarios"].values():
                scen["exceeds_annual_demand"] = None
            continue

        item["policy_state"] = POLICY_STATE_PRICED
        # A not_stockable SKU has a reorder point covering normal demand and no
        # safety stock at all. That is deliberate and the screen must say so,
        # otherwise it reads as an under-stocked item somebody should "fix".
        item["policy_note"] = (
            NOT_STOCKABLE_NOTE if item.get("service_tier") == TIER_NOT_STOCKABLE
            else None)
        add = item["avg_daily_demand"] or 0.0
        rop = item["reorder_point"]

        item["days_cover_remaining"] = round(stock / add, 1) if (stock is not None and add > 0) else None
        item["needs_reorder"] = stock is not None and stock <= rop
        item["approaching_rop"] = stock is not None and rop < stock <= rop * 1.2

        # Order-up-to quantity, NOT EOQ. See ORDER_QTY_NOTE: EOQ under the
        # current provisional cost inputs comes out larger than a year of
        # demand for essentially every SKU, so it cannot be handed to staff
        # as "order this many". This is the textbook (s, S) alternative -
        # bring stock up to the reorder point plus one review period of
        # demand - and it uses only the three inputs that are actually
        # measured here (on-hand, ADUS, lead time), not the two cost
        # estimates that are still guesses.
        target = rop + REVIEW_PERIOD_DAYS * add
        item["order_up_to_level"] = round(target, 1)
        item["suggested_order_qty"] = (
            max(0, int(round(target - stock))) if item["needs_reorder"] else 0
        )
        # The same order without the safety stock: what the expected sales
        # alone would use up. The pair is the range the Overview shows, so a
        # provisional quantity is not read as exact - the low end if delivery
        # and sales go as averaged, the high end with the full buffer.
        ss = item["safety_stock"] or 0.0
        item["suggested_order_min"] = (
            min(item["suggested_order_qty"], max(0, int(round(target - ss - stock))))
            if item["needs_reorder"] else 0
        )

        for scen in item["scenarios"].values():
            scen["exceeds_annual_demand"] = bool(
                item["annual_demand"] and scen["eoq"] > item["annual_demand"]
            )

    return sorted(by_product.values(), key=lambda i: i["item_name"])


# ---------------------------------------------------------------- stock

@app.get("/api/stock")
def get_stock():
    c = con()
    cat = catalog.compute_catalog(c)
    supplier = request.args.get("supplier") or None
    category = request.args.get("category") or None
    items = [p for p in cat if catalog.matches(p, supplier, category) and p["current_stock"] is not None]
    items.sort(key=lambda p: (p["days_of_supply"] is None, p["days_of_supply"] or 0))
    return jsonify({"items": items, "covered": len(items), "total": len(cat)})


# ------------------------------------------------------------ inventory

@app.get("/api/inventory")
def get_inventory_counts():
    """Staff-entered stock counts for one month (?month=YYYY-MM), newest
    first, joined to the product name so the table needs no second call.

    Also returns `workbook_month`: the last month the historical inventory
    workbook covers. The interface shows it so whoever is counting can see
    where the existing data stops and what they are extending."""
    month = request.args.get("month") or ""
    if not validation.ISO_MONTH_RE.match(month):
        return jsonify({"ok": False, "errors": {"count_month": "Month must be YYYY-MM."}}), 400

    counts = dbmod.rows(con(), """
        SELECT ic.count_id, ic.product_id, p.item_name, p.supplier_name, p.category,
               ic.count_month, ic.quantity, ic.note, ic.date_logged
          FROM Inventory_Count ic
          JOIN Dim_Product p ON p.product_id = ic.product_id
         WHERE ic.count_month = ?
         ORDER BY ic.count_id DESC
    """, (month,))
    for r in counts:
        r["supplier_name"] = r["supplier_name"] or catalog.UNATTRIBUTED

    workbook = catalog.load_csv_stock()
    return jsonify({
        "month": month,
        "counts": counts,
        "total_units": sum(r["quantity"] for r in counts),
        "workbook_month": max((v["month"] for v in workbook.values()), default=None),
    })


@app.post("/api/inventory/import")
def import_inventory_counts():
    """Bulk stock counts from a .csv/.xlsx, archived into rawdata/ as uploaded.

    Expected columns (header spelling is matched loosely - case, spacing and
    punctuation are ignored):

        Item            required, must match a Dim_Product item_name
        Units On Hand   required, whole number >= 0  (aliases: Quantity, Stock)
        Month           optional, YYYY-MM - falls back to the ?month= form field
        Note            optional

    Upserts exactly like the single-count form: one count per product per
    month, a re-import corrects rather than accumulates.

    Rows that cannot be used are REPORTED, never silently dropped - the same
    rule step2 follows with Exception_Log. Valid rows are still applied, so a
    500-row workbook is not rejected wholesale over two bad lines.

    Item may be an item's name or any tally-sheet name the vocabulary maps
    onto it. A valid row whose name is neither is HELD (Pending_Import_Row)
    and counted under "Names to review"; it is applied once the name is
    settled there.
    """
    storage = request.files.get("file")
    if storage is None or not storage.filename:
        return jsonify({"ok": False, "error": "Choose a .csv or .xlsx file to import."}), 400
    bad = _accepted_upload(storage)
    if bad:
        return jsonify({"ok": False, "error": bad}), 400

    default_month = (request.form.get("month") or "").strip()
    if default_month and not validation.ISO_MONTH_RE.match(default_month):
        return jsonify({"ok": False, "error": "Month must be YYYY-MM."}), 400

    saved = _archive_upload(storage)
    storage.stream.seek(0)
    try:
        _, rows = _read_table(storage)
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc), "saved_to": saved}), 400

    c = con()
    products = _product_index(c)
    today_month = date.today().strftime("%Y-%m")
    imported = updated = 0
    rejected = []
    held = {}                                      # name -> rows held for review
    now = datetime.now().isoformat(timespec="seconds")

    for i, row in enumerate(rows, start=2):        # start=2: row 1 is the header
        item = _pick(row, "Item", "Item Name", "Product", "Canonical Item Name")
        qty = _as_int(_pick(row, "Units On Hand", "Quantity", "Units", "Stock", "On Hand"))
        month = str(_pick(row, "Month", "Count Month") or default_month).strip()[:7]
        note = _pick(row, "Note", "Remarks")

        if not item:
            rejected.append({"row": i, "item": "", "reason": "No item name."}); continue
        item = str(item).strip()
        if qty is None or qty < 0:
            rejected.append({"row": i, "item": item,
                             "reason": "Units on hand must be a whole number >= 0."}); continue
        if not validation.ISO_MONTH_RE.match(month):
            rejected.append({"row": i, "item": item,
                             "reason": "No usable month (YYYY-MM)."}); continue
        if month > today_month:
            rejected.append({"row": i, "item": item,
                             "reason": "Month is in the future."}); continue
        pid = _lookup(products, item)
        if pid is None:
            names.hold_row(c, "inventory", item, qty, storage.filename, count_month=month,
                           note=(str(note).strip() if note else None),
                           entered_by=auth.current_user())
            held[item] = held.get(item, 0) + 1
            continue

        existed = dbmod.one(c, "SELECT 1 FROM Inventory_Count WHERE product_id = ? AND count_month = ?",
                            (pid, month))
        c.execute("""
            INSERT INTO Inventory_Count (product_id, count_month, quantity, note, counted_by, date_logged)
            VALUES (?, ?, ?, ?, 'import', ?)
            ON CONFLICT (product_id, count_month) DO UPDATE SET
                quantity = excluded.quantity, note = excluded.note, date_logged = excluded.date_logged
        """, (pid, month, qty, (str(note).strip() if note else None), now))
        if existed:
            updated += 1
        else:
            imported += 1

    c.commit()
    return jsonify({
        "ok": True, "imported": imported, "updated": updated,
        "rejected": rejected[:50], "rejected_total": len(rejected),
        **_held_summary(held),
        "rows_read": len(rows), "saved_to": saved,
    })


def _held_summary(held):
    return {"held": sum(held.values()),
            "held_names": [{"name": n, "rows": k} for n, k in sorted(held.items())][:50]}


@app.post("/api/inventory")
def add_inventory_count():
    """Record (or correct) one product's stock count for one month.

    Upsert, not append: `Inventory_Count` is UNIQUE on (product_id,
    count_month), and a recount replaces the earlier figure rather than
    adding to it. The response says which of the two happened so the
    interface can tell the user it overwrote something."""
    payload = request.get_json(silent=True) or {}
    c = con()
    errors = validation.validate_inventory_count(c, payload)
    if errors:
        return jsonify({"ok": False, "errors": errors}), 400

    product_id = int(payload["product_id"])
    count_month = payload["count_month"]
    quantity = int(float(payload["quantity"]))
    note = (payload.get("note") or "").strip() or None

    previous = dbmod.one(c, """
        SELECT quantity FROM Inventory_Count WHERE product_id = ? AND count_month = ?
    """, (product_id, count_month))
    # A first count is a new record; replacing a different figure is an edit.
    # A correction from Current Inventory ({"correction": true}) is an edit
    # even as the month's first count: it changes the on-hand figure the store
    # already has on record (an earlier count or the workbook), and the
    # interface sends it only when the figure differs from that one.
    if payload.get("correction") is True:
        gate = _needs_admin("Correcting the units on hand")
        if gate:
            return gate
    elif previous and previous["quantity"] != quantity:
        gate = _needs_admin("Changing a recorded stock count")
        if gate:
            return gate

    c.execute("""
        INSERT INTO Inventory_Count (product_id, count_month, quantity, note, counted_by, date_logged)
        VALUES (?, ?, ?, ?, ?, ?)
        ON CONFLICT (product_id, count_month) DO UPDATE SET
            quantity = excluded.quantity,
            note = excluded.note,
            counted_by = excluded.counted_by,
            date_logged = excluded.date_logged
    """, (product_id, count_month, quantity, note, auth.current_user(),
          datetime.now().isoformat(timespec="seconds")))
    c.commit()

    product = dbmod.one(c, "SELECT item_name FROM Dim_Product WHERE product_id = ?", (product_id,))
    return jsonify({
        "ok": True,
        "replaced": previous["quantity"] if previous else None,
        "count": {
            "product_id": product_id,
            "item_name": product["item_name"],
            "count_month": count_month,
            "quantity": quantity,
            "note": note,
        },
    })


@app.delete("/api/inventory/<int:count_id>")
def delete_inventory_count(count_id):
    gate = _needs_admin("Removing a stock count")
    if gate:
        return gate
    c = con()
    if not dbmod.one(c, "SELECT 1 FROM Inventory_Count WHERE count_id = ?", (count_id,)):
        return jsonify({"ok": False, "error": "No such inventory count."}), 404
    c.execute("DELETE FROM Inventory_Count WHERE count_id = ?", (count_id,))
    c.commit()
    return jsonify({"ok": True})


# ------------------------------------------------------------- restocks

def _restock_dates(payload, require_ordered=True):
    """(ordered_on, delivered_on, errors) from a restock form. Delivered may be
    blank (still on its way), never before the order, and neither in the future."""
    today_iso = date.today().isoformat()
    errors = {}
    ordered = str(payload.get("ordered_on") or "").strip()
    delivered = str(payload.get("delivered_on") or "").strip() or None
    if require_ordered:
        if not validation.is_iso_date(ordered):
            errors["ordered_on"] = "Pick the date it was ordered."
        elif ordered > today_iso:
            errors["ordered_on"] = "The order date cannot be in the future."
    if delivered is not None:
        if not validation.is_iso_date(delivered):
            errors["delivered_on"] = "Delivered date must be YYYY-MM-DD."
        elif delivered > today_iso:
            errors["delivered_on"] = "The delivery date cannot be in the future."
        elif ordered and validation.is_iso_date(ordered) and delivered < ordered:
            errors["delivered_on"] = "It cannot arrive before it was ordered."
    return ordered, delivered, errors


@app.get("/api/restocks")
def get_restocks():
    """Recorded restocks, newest first, and per supplier what they add up to:
    deliveries so far, the median days from order to delivery, and whether
    that median is used yet (step5a: from MIN_DELIVERIES deliveries on, at the
    next pipeline run). `lead_time_now` is the lead time the reorder points
    currently use for that supplier's items (the most common one)."""
    import step5a_set_lead_times as step5a
    c = con()
    if "Restock_Log" not in _table_names(c):
        return jsonify({"restocks": [], "suppliers": [], "min_deliveries": step5a.MIN_DELIVERIES})
    restocks = dbmod.rows(c, """
        SELECT restock_id, supplier_name, ordered_on, delivered_on, note, entered_by,
               CASE WHEN delivered_on IS NOT NULL
                    THEN CAST(ROUND(julianday(delivered_on) - julianday(ordered_on)) AS INTEGER) END AS days
          FROM Restock_Log ORDER BY ordered_on DESC, restock_id DESC LIMIT 200
    """)
    now = {r["supplier_name"]: r["lead_time_days"] for r in dbmod.rows(c, """
        SELECT supplier_name, lead_time_days FROM (
            SELECT supplier_name, lead_time_days, COUNT(*) n,
                   ROW_NUMBER() OVER (PARTITION BY supplier_name ORDER BY COUNT(*) DESC) rk
              FROM Dim_Product WHERE supplier_name IS NOT NULL GROUP BY 1, 2)
         WHERE rk = 1""")}
    days = step5a.delivery_days(c)
    open_orders = {r["supplier_name"]: r["n"] for r in dbmod.rows(c, """
        SELECT supplier_name, COUNT(*) n FROM Restock_Log WHERE delivered_on IS NULL GROUP BY 1""")}
    suppliers = []
    for s in sorted(set(days) | set(open_orders)):
        d = days.get(s, [])
        suppliers.append({
            "supplier_name": s, "deliveries": len(d), "open": open_orders.get(s, 0),
            "median_days": round(statistics.median(d), 1) if d else None,
            "min_days": min(d) if d else None, "max_days": max(d) if d else None,
            "in_use": len(d) >= step5a.MIN_DELIVERIES, "lead_time_now": now.get(s),
        })
    return jsonify({"restocks": restocks, "suppliers": suppliers, "min_deliveries": step5a.MIN_DELIVERIES})


@app.post("/api/restocks")
def add_restock():
    """Record a restock order: {supplier_name, ordered_on, delivered_on?, note?}."""
    payload = request.get_json(silent=True) or {}
    supplier = str(payload.get("supplier_name") or "").strip()
    ordered, delivered, errors = _restock_dates(payload)
    if not supplier:
        errors["supplier_name"] = "Choose the supplier."
    if errors:
        return jsonify({"ok": False, "errors": errors}), 400
    c = con()
    c.execute(dbmod.RESTOCK_LOG_DDL)               # in case startup could not create it
    cur = c.execute("""
        INSERT INTO Restock_Log (supplier_name, ordered_on, delivered_on, note, entered_by, date_logged)
        VALUES (?, ?, ?, ?, ?, ?)
    """, (supplier, ordered, delivered, (str(payload.get("note") or "").strip() or None),
          auth.current_user(), datetime.now().isoformat(timespec="seconds")))
    c.commit()
    return jsonify({"ok": True, "restock_id": cur.lastrowid})


@app.put("/api/restocks/<int:restock_id>")
def mark_restock_delivered(restock_id):
    """Set (or correct) the day a restock arrived: {delivered_on}."""
    payload = request.get_json(silent=True) or {}
    c = con()
    row = dbmod.one(c, "SELECT ordered_on, delivered_on FROM Restock_Log WHERE restock_id = ?",
                    (restock_id,)) if "Restock_Log" in _table_names(c) else None
    if row is None:
        return jsonify({"ok": False, "error": "No such restock."}), 404
    # Marking an order delivered is a new fact; changing the date is an edit.
    if row["delivered_on"]:
        gate = _needs_admin("Changing a delivery date")
        if gate:
            return gate
    _o, delivered, errors = _restock_dates({**payload, "ordered_on": row["ordered_on"]})
    if delivered is None:
        errors["delivered_on"] = "Pick the date it arrived."
    if errors:
        return jsonify({"ok": False, "errors": errors}), 400
    c.execute("UPDATE Restock_Log SET delivered_on = ? WHERE restock_id = ?", (delivered, restock_id))
    c.commit()
    return jsonify({"ok": True})


@app.delete("/api/restocks/<int:restock_id>")
def delete_restock(restock_id):
    gate = _needs_admin("Removing a restock")
    if gate:
        return gate
    c = con()
    if "Restock_Log" not in _table_names(c) or not dbmod.one(
            c, "SELECT 1 FROM Restock_Log WHERE restock_id = ?", (restock_id,)):
        return jsonify({"ok": False, "error": "No such restock."}), 404
    c.execute("DELETE FROM Restock_Log WHERE restock_id = ?", (restock_id,))
    c.commit()
    return jsonify({"ok": True})


# -------------------------------------------------------------- calendar

@app.get("/api/calendar")
def get_calendar():
    return jsonify(dbmod.rows(con(), "SELECT * FROM Dim_Date ORDER BY calendar_date"))


@app.get("/api/calendar/closed")
def get_closed_dates():
    """Every date the store was shut, from BOTH records of it.

    There are two, and they do not agree:

      Dim_Date.is_store_closed   the published academic calendar
                                 (calendar_ranges.csv) plus anything staff
                                 flagged through this interface.
      Dim_Day_Status             read out of the TBS workbooks themselves,
                                 where each date's header cell is coloured
                                 and a per-sheet legend gives the meaning
                                 ("NO OPERATION", "HOLIDAY (DAY OF VALOR)").

    The workbook is the store's own record of what it actually did, and it
    knows things the academic calendar cannot: 48 of the dates it marks
    "NO OPERATION" are Sundays, which no term calendar has an opinion
    about. Showing only Dim_Date therefore under-reports closures badly.

    Rows are merged by date and carry their provenance, so the interface
    can show WHY a date is flagged and where that came from rather than
    presenting two different sets of truth as one list. Returned most
    recent first: the list spans 2023-2026 and the recent end is the part
    anyone is checking.
    """
    c = con()
    rows = dbmod.rows(c, """
        WITH merged AS (
            SELECT calendar_date,
                   1 AS from_calendar, 0 AS from_workbook,
                   NULL AS label
            FROM Dim_Date WHERE is_store_closed = 1

            UNION ALL

            SELECT calendar_date,
                   0 AS from_calendar, 1 AS from_workbook,
                   legend_label AS label
            FROM Dim_Day_Status WHERE is_store_closed = 1
        )
        SELECT calendar_date,
               MAX(from_calendar) AS from_calendar,
               MAX(from_workbook) AS from_workbook,
               MAX(label)         AS label
        FROM merged
        GROUP BY calendar_date
        ORDER BY calendar_date DESC
    """) if "Dim_Day_Status" in _table_names(c) else dbmod.rows(c, """
        SELECT calendar_date, 1 AS from_calendar, 0 AS from_workbook,
               NULL AS label
        FROM Dim_Date WHERE is_store_closed = 1
        ORDER BY calendar_date DESC
    """)

    # Closure_Log holds the reason a staff member typed. It is the better
    # label when there is one, because it describes this specific closure
    # rather than the legend category the workbook filed it under.
    reasons = {}
    if "Closure_Log" in _table_names(c):
        # Oldest first, so a later entry's reason (an edit) overwrites an
        # earlier one for the same day.
        for r in dbmod.rows(c, "SELECT closure_date, reason FROM Closure_Log "
                               "WHERE is_closed = 1 AND reason IS NOT NULL ORDER BY closure_id"):
            if r["reason"]:
                reasons[r["closure_date"]] = r["reason"]

    out = []
    for r in rows:
        d = r["calendar_date"]
        src = ("both" if r["from_calendar"] and r["from_workbook"]
               else "workbook" if r["from_workbook"] else "calendar")
        out.append({
            "calendar_date": d,
            "source": src,
            "reason": reasons.get(d) or r["label"] or None,
        })
    return jsonify(out)


def _table_names(c):
    return {r[0] for r in c.execute(
        "SELECT name FROM sqlite_master WHERE type='table'")}


@app.get("/api/calendar/<iso_date>")
def get_date_flags(iso_date):
    row = dbmod.one(con(), "SELECT * FROM Dim_Date WHERE calendar_date = ?", (iso_date,))
    return jsonify(row)


@app.put("/api/calendar/<iso_date>/closure")
def set_closure(iso_date):
    if not validation.is_iso_date(iso_date):
        return jsonify({"ok": False, "errors": {"calendar_date": "Date must be YYYY-MM-DD."}}), 400
    payload = request.get_json(silent=True) or {}
    closed = 1 if payload.get("closed") else 0
    reason = (payload.get("reason") or "").strip() or None

    c = con()
    # So the UPDATE below has a row to land on. A date past even the extended
    # calendar is still logged, and populate_dim_date.py applies it once the
    # calendar reaches it.
    _extend_calendar(c)
    # Marking a day closed adds a closure; reopening a closed day removes one.
    if not closed:
        was = dbmod.one(c, "SELECT is_store_closed FROM Dim_Date WHERE calendar_date = ?", (iso_date,))
        if was and was["is_store_closed"]:
            gate = _needs_admin("Reopening a closed day")
            if gate:
                return gate
    # Remediation D3. Log the toggle in Closure_Log (durable across a
    # populate_dim_date.py rebuild, which reads it back - see that
    # script) AND update Dim_Date directly (so the change is visible
    # immediately, same dual-write shape as add_event() below).
    c.execute("""
        INSERT INTO Closure_Log (closure_date, is_closed, reason, created_by, date_logged)
        VALUES (?, ?, ?, ?, ?)
    """, (iso_date, closed, reason, auth.current_user(), datetime.now().isoformat(timespec="seconds")))
    c.execute("UPDATE Dim_Date SET is_store_closed = ? WHERE calendar_date = ?", (closed, iso_date))
    c.commit()
    return jsonify({"ok": True})


@app.put("/api/calendar/closures/<iso_date>")
def edit_closure(iso_date):
    """Move a closure to another day, or change its reason - upcoming ones
    when the school calendar shifts, past ones to correct the record:
    {calendar_date, reason?}. A closure is a date in Dim_Date, not a row of
    its own, so a move is logged as the two toggles it amounts to - the old
    day reopened, the new day closed - in Closure_Log, which every
    populate_dim_date.py rebuild replays (latest entry per day wins). That
    works for a closure from the published calendar as well as one flagged
    here: the log overrides the calendar file without editing it. A past
    change reaches the forecasts at the next pipeline run (the Tally page's
    staleness banner says so)."""
    gate = _needs_admin("Changing a closure")
    if gate:
        return gate
    payload = request.get_json(silent=True) or {}
    new_date = str(payload.get("calendar_date") or "").strip()
    reason = (payload.get("reason") or "").strip() or None

    if not validation.is_iso_date(iso_date) or not validation.is_iso_date(new_date):
        return jsonify({"ok": False, "errors": {"calendar_date": "Date must be YYYY-MM-DD."}}), 400

    c = con()
    _extend_calendar(c)
    gap = _calendar_gap(c, new_date)
    if gap:
        return jsonify({"ok": False, "errors": {"calendar_date": gap}}), 400
    row = dbmod.one(c, "SELECT is_store_closed FROM Dim_Date WHERE calendar_date = ?", (iso_date,))
    if not row or not row["is_store_closed"]:
        return jsonify({"ok": False, "error": f"{iso_date} is not marked closed."}), 404

    user, now = auth.current_user(), datetime.now().isoformat(timespec="seconds")
    log = "INSERT INTO Closure_Log (closure_date, is_closed, reason, created_by, date_logged) VALUES (?, ?, ?, ?, ?)"
    if new_date != iso_date:
        c.execute(log, (iso_date, 0, f"Moved to {new_date}", user, now))
        c.execute("UPDATE Dim_Date SET is_store_closed = 0 WHERE calendar_date = ?", (iso_date,))
    c.execute(log, (new_date, 1, reason, user, now))
    c.execute("UPDATE Dim_Date SET is_store_closed = 1 WHERE calendar_date = ?", (new_date,))
    c.commit()
    return jsonify({"ok": True, "closure": {"calendar_date": new_date, "reason": reason,
                                            "previous_date": iso_date}})


# ----------------------------------------------------------------- tally

def _order_type(transaction_type, raw):
    """Fact_Sales.order_type for a new row: walk_in / bulk / pre_order for a
    sale (walk-in when not given), None for a damaged / promo / transfer
    removal, which is not a sale of any kind, and False when `raw` is not an
    order type at all. The forecasts train on walk-in sales only
    (scripts/order_types.py)."""
    if str(transaction_type).upper() != "SALE":
        return None
    order_type = order_types.normalize(raw)
    return False if order_type is None else order_type


@app.get("/api/orders/upcoming")
def get_upcoming_orders():
    """Known bulk / organisation orders and pre-orders still to come (and any
    from the last 14 days, so a just-collected one can still be seen). The
    forecasts add each on its expected date (scripts/order_types.py)."""
    c = con()
    if "Upcoming_Order" not in _table_names(c):
        return jsonify([])
    since = (date.today() - timedelta(days=14)).isoformat()
    return jsonify(dbmod.rows(c, """
        SELECT o.order_id, o.product_id, p.item_name, o.expected_date, o.quantity,
               o.order_type, o.note, o.entered_by
          FROM Upcoming_Order o JOIN Dim_Product p ON p.product_id = o.product_id
         WHERE o.expected_date >= ?
         ORDER BY o.expected_date, o.order_id
    """, (since,)))


@app.post("/api/orders/upcoming")
def add_upcoming_order():
    """{product_id, expected_date, quantity, order_type: bulk | pre_order, note?}"""
    payload = request.get_json(silent=True) or {}
    c = con()
    errors = {}
    pid = _as_int(payload.get("product_id"))
    if pid is None or not dbmod.one(c, "SELECT 1 FROM Dim_Product WHERE product_id = ?", (pid,)):
        errors["product_id"] = "Select an item."
    expected = str(payload.get("expected_date") or "").strip()
    if not validation.is_iso_date(expected):
        errors["expected_date"] = "Pick the date it will be collected."
    elif expected < date.today().isoformat():
        errors["expected_date"] = "That date has passed - tally the sale instead."
    qty = _as_int(payload.get("quantity"))
    if qty is None or qty <= 0:
        errors["quantity"] = "Quantity must be a whole number > 0."
    order_type = order_types.normalize(payload.get("order_type"), default=None)
    if order_type not in ("bulk", "pre_order"):
        errors["order_type"] = "Choose bulk or pre-order."
    if errors:
        return jsonify({"ok": False, "errors": errors}), 400
    c.execute(dbmod.UPCOMING_ORDER_DDL)            # in case startup could not create it
    cur = c.execute("""
        INSERT INTO Upcoming_Order (product_id, expected_date, quantity, order_type, note, entered_by, date_logged)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (pid, expected, qty, order_type, (str(payload.get("note") or "").strip() or None),
          auth.current_user(), datetime.now().isoformat(timespec="seconds")))
    c.commit()
    return jsonify({"ok": True, "order_id": cur.lastrowid})


@app.delete("/api/orders/upcoming/<int:order_id>")
def delete_upcoming_order(order_id):
    gate = _needs_admin("Removing an upcoming order")
    if gate:
        return gate
    c = con()
    if "Upcoming_Order" not in _table_names(c) or not dbmod.one(
            c, "SELECT 1 FROM Upcoming_Order WHERE order_id = ?", (order_id,)):
        return jsonify({"ok": False, "error": "No such order."}), 404
    c.execute("DELETE FROM Upcoming_Order WHERE order_id = ?", (order_id,))
    c.commit()
    return jsonify({"ok": True})


@app.get("/api/tally/recent")
def get_recent_entries():
    limit = request.args.get("limit", default=50, type=int)
    rows = dbmod.rows(con(), """
        SELECT f.sale_id, f.product_id, d.calendar_date, f.quantity_sold,
               UPPER(f.transaction_type) AS transaction_type, f.imputation_flag, f.is_censored, f.entered_by,
               f.order_type, f.tally_date_flag, p.item_name, COALESCE(p.supplier_name, ?) AS supplier_name
        FROM Fact_Sales f
        JOIN Dim_Date d ON d.date_id = f.date_id
        JOIN Dim_Product p ON p.product_id = f.product_id
        WHERE f.quantity_sold > 0
        ORDER BY d.calendar_date DESC, f.sale_id DESC
        LIMIT ?
    """, (catalog.UNATTRIBUTED, limit))
    for r in rows:
        r["is_local"] = False
    return jsonify(rows)


@app.get("/api/tally")
def get_entries_by_date():
    iso_date = request.args.get("date")
    if not iso_date:
        return jsonify({"ok": False, "errors": {"calendar_date": "date=YYYY-MM-DD is required."}}), 400
    rows = dbmod.rows(con(), """
        SELECT f.sale_id, f.product_id, d.calendar_date, f.quantity_sold,
               UPPER(f.transaction_type) AS transaction_type, f.imputation_flag, f.is_censored, f.entered_by,
               f.order_type, f.tally_date_flag, p.item_name, COALESCE(p.supplier_name, ?) AS supplier_name
        FROM Fact_Sales f
        JOIN Dim_Date d ON d.date_id = f.date_id
        JOIN Dim_Product p ON p.product_id = f.product_id
        WHERE f.quantity_sold > 0 AND d.calendar_date = ?
        ORDER BY f.sale_id DESC
    """, (catalog.UNATTRIBUTED, iso_date))
    for r in rows:
        r["is_local"] = False
    return jsonify(rows)


@app.post("/api/tally")
def add_entry():
    payload = request.get_json(silent=True) or {}
    c = con()
    _extend_calendar(c)
    errors = validation.validate_entry(c, payload)
    transaction_type = str(payload.get("transaction_type") or "").upper()
    order_type = _order_type(transaction_type, payload.get("order_type"))
    if order_type is False:
        errors["order_type"] = "Choose walk-in, bulk or pre-order."
    if errors:
        return jsonify({"ok": False, "errors": errors}), 400

    product_id = int(payload["product_id"])
    quantity_sold = int(float(payload["quantity_sold"]))
    calendar_date = payload["calendar_date"]

    date_row = dbmod.one(c, "SELECT date_id FROM Dim_Date WHERE calendar_date = ?", (calendar_date,))
    # Stored in lower case like every historical row ('sale'); shown in upper case.
    cur = c.execute("""
        INSERT INTO Fact_Sales
            (product_id, date_id, quantity_sold, imputation_flag, tally_date_flag,
             transaction_type, entered_by, order_type)
        VALUES (?, ?, ?, 0, 0, ?, ?, ?)
    """, (product_id, date_row["date_id"], quantity_sold, transaction_type.lower(),
          auth.current_user(), order_type))
    c.commit()

    product = dbmod.one(c, "SELECT item_name, supplier_name FROM Dim_Product WHERE product_id = ?", (product_id,))
    entry = {
        "sale_id": cur.lastrowid,
        "product_id": product_id,
        "item_name": product["item_name"],
        "supplier_name": product["supplier_name"] or catalog.UNATTRIBUTED,
        "quantity_sold": quantity_sold,
        "calendar_date": calendar_date,
        "transaction_type": transaction_type,
        "order_type": order_type,
        "entered_by": auth.current_user(),
        "is_local": False,
    }
    return jsonify({"ok": True, "entry": entry})


@app.put("/api/tally/<int:sale_id>")
def edit_entry(sale_id):
    """Correct a tally entry: its date, quantity or transaction type.
    {calendar_date, quantity_sold, transaction_type, order_type?}. The item
    stays - a wrong item is removed and recorded again.

    An entry typed or imported here (tally_date_flag = 0) is simply changed.
    One loaded from the tally workbook (tally_date_flag = 1) is changed too,
    and the change is also kept in Tally_Correction: step2 deletes and reloads
    every workbook row on each pipeline run, and re-applies it from there
    (scripts/tally_corrections.py). The forecasts pick a change up at the
    next pipeline run."""
    gate = _needs_admin("Changing a tally entry")
    if gate:
        return gate
    payload = request.get_json(silent=True) or {}
    c = con()
    tally_corrections.ensure(c)
    row = _tally_row(c, sale_id)
    if row is None:
        return jsonify({"ok": False, "error": "That entry no longer exists."}), 404

    _extend_calendar(c)
    errors = validation.validate_entry(c, {**payload, "product_id": row["product_id"]})
    transaction_type = str(payload.get("transaction_type") or "").upper()
    order_type = _order_type(transaction_type, payload.get("order_type"))
    if order_type is False:
        errors["order_type"] = "Choose walk-in, bulk or pre-order."
    if errors:
        return jsonify({"ok": False, "errors": errors}), 400

    quantity_sold = int(float(payload["quantity_sold"]))
    calendar_date = payload["calendar_date"]
    date_row = dbmod.one(c, "SELECT date_id FROM Dim_Date WHERE calendar_date = ?", (calendar_date,))
    correction_id = _record_correction(c, row, calendar_date=calendar_date, quantity=quantity_sold,
                                       transaction_type=transaction_type.lower(), order_type=order_type)
    c.execute("""
        UPDATE Fact_Sales SET date_id = ?, quantity_sold = ?, transaction_type = ?, order_type = ?,
                              correction_id = ?
        WHERE sale_id = ?
    """, (date_row["date_id"], quantity_sold, transaction_type.lower(), order_type, correction_id, sale_id))
    c.commit()
    return jsonify({"ok": True, "entry": {
        "sale_id": sale_id, "calendar_date": calendar_date, "quantity_sold": quantity_sold,
        "transaction_type": transaction_type, "order_type": order_type,
    }})


@app.delete("/api/tally/<int:sale_id>")
def delete_entry(sale_id):
    """Remove a tally entry recorded by mistake. A workbook row's removal is
    kept in Tally_Correction, as edit_entry keeps a change."""
    gate = _needs_admin("Removing a tally entry")
    if gate:
        return gate
    c = con()
    tally_corrections.ensure(c)
    row = _tally_row(c, sale_id)
    if row is None:
        return jsonify({"ok": False, "error": "That entry no longer exists."}), 404
    _record_correction(c, row, is_deleted=1)
    c.execute("DELETE FROM Fact_Sales WHERE sale_id = ?", (sale_id,))
    c.commit()
    return jsonify({"ok": True, "deleted": {"sale_id": sale_id}})


def _tally_row(c, sale_id):
    return dbmod.one(c, """
        SELECT f.sale_id, f.product_id, f.quantity_sold, f.imputation_flag, f.tally_date_flag,
               f.correction_id, d.calendar_date, p.item_name
        FROM Fact_Sales f
        JOIN Dim_Date d ON d.date_id = f.date_id
        JOIN Dim_Product p ON p.product_id = f.product_id
        WHERE f.sale_id = ?
    """, (sale_id,))


def _record_correction(c, row, calendar_date=None, quantity=None, transaction_type=None,
                       order_type=None, is_deleted=0):
    """For a workbook row, the Tally_Correction step2 will re-apply: the one it
    already has, updated, or a new one naming the row as the workbook has it.
    None for an entry made here, which the pipeline never reloads."""
    if not row["tally_date_flag"]:
        return None
    values = (calendar_date, quantity, transaction_type, order_type, is_deleted)
    if row["correction_id"] and dbmod.one(c, "SELECT 1 AS ok FROM Tally_Correction WHERE correction_id = ?",
                                          (row["correction_id"],)):
        c.execute("""
            UPDATE Tally_Correction SET calendar_date = ?, quantity = ?, transaction_type = ?,
                   order_type = ?, is_deleted = ?, entered_by = ?, date_logged = ?
            WHERE correction_id = ?
        """, (*values, auth.current_user(), datetime.now().isoformat(timespec="seconds"), row["correction_id"]))
        return row["correction_id"]
    cur = c.execute("""
        INSERT INTO Tally_Correction (item_name, original_date, original_quantity, imputation_flag,
            calendar_date, quantity, transaction_type, order_type, is_deleted, entered_by, date_logged)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (row["item_name"], row["calendar_date"], row["quantity_sold"], row["imputation_flag"] or 0,
          *values, auth.current_user(), datetime.now().isoformat(timespec="seconds")))
    return cur.lastrowid


@app.post("/api/tally/import")
def import_tally():
    """Bulk sales-tally entries from a .csv/.xlsx, archived into rawdata/ as
    uploaded. Shaped like the source workbooks step0 reads:

        Date              required, ISO YYYY-MM-DD, not in the future
        Item              required, must match a Dim_Product item_name
        Total Quantity    required, whole number > 0  (aliases: Quantity, Qty)
        Transaction Type  optional, defaults to SALE
        Order Type        optional, for sales: walk-in (default), bulk, pre-order

    Dates must be ISO. DD/MM and MM/DD are refused rather than guessed - with
    both accepted, 05/11/2025 is valid under either reading and lands six
    months apart depending on which won, which is exactly the defect the
    pipeline was cleaned up to remove (README, "Key design decisions").

    Appends, it does not upsert: two tallies of the same item on the same day
    are two real movements. Rows that cannot be used are reported, not dropped.

    Item may be an item's name or any tally-sheet name the vocabulary maps
    onto it. A valid row whose name is neither is HELD (Pending_Import_Row),
    not rejected: because this appends, re-importing the file once the name
    was known would load every other row twice. It is applied when the name
    is settled under "Names to review".
    """
    storage = request.files.get("file")
    if storage is None or not storage.filename:
        return jsonify({"ok": False, "error": "Choose a .csv or .xlsx file to import."}), 400
    bad = _accepted_upload(storage)
    if bad:
        return jsonify({"ok": False, "error": bad}), 400

    saved = _archive_upload(storage)
    storage.stream.seek(0)
    try:
        _, rows = _read_table(storage)
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc), "saved_to": saved}), 400

    c = con()
    _extend_calendar(c)
    products = _product_index(c)
    dates = {r["calendar_date"]: r["date_id"]
             for r in dbmod.rows(c, "SELECT date_id, calendar_date FROM Dim_Date")}
    today_iso = date.today().isoformat()
    user = auth.current_user()
    imported = 0
    rejected = []
    held = {}

    for i, row in enumerate(rows, start=2):
        item = _pick(row, "Item", "Item Name", "Product", "Canonical Item Name")
        qty = _as_int(_pick(row, "Total Quantity", "Quantity", "Qty", "Units", "Quantity Sold"))
        iso = _as_iso_date(_pick(row, "Date", "Calendar Date", "Sale Date"))
        ttype = str(_pick(row, "Transaction Type", "Type") or "SALE").strip().upper()
        raw_order = _pick(row, "Order Type", "Order", "Sale Type")

        if not item:
            rejected.append({"row": i, "item": "", "reason": "No item name."}); continue
        item = str(item).strip()
        if iso is None:
            rejected.append({"row": i, "item": item,
                             "reason": "Date missing or not ISO YYYY-MM-DD."}); continue
        if iso > today_iso:
            rejected.append({"row": i, "item": item, "reason": "Date is in the future."}); continue
        if iso not in dates:
            rejected.append({"row": i, "item": item,
                             "reason": validation.calendar_gap_message(c, iso)}); continue
        if qty is None or qty <= 0:
            rejected.append({"row": i, "item": item,
                             "reason": "Quantity must be a whole number > 0."}); continue
        if ttype not in validation.TRANSACTION_TYPES:
            rejected.append({"row": i, "item": item,
                             "reason": f"Unknown transaction type '{ttype}'."}); continue
        order_type = _order_type(ttype, raw_order)
        if order_type is False:
            rejected.append({"row": i, "item": item,
                             "reason": f"Unknown order type '{raw_order}' (walk-in, bulk or pre-order)."}); continue
        pid = _lookup(products, item)
        if pid is None:
            names.hold_row(c, "tally", item, qty, storage.filename, calendar_date=iso,
                           transaction_type=ttype.lower(), entered_by=user, order_type=order_type)
            held[item] = held.get(item, 0) + 1
            continue

        c.execute("""
            INSERT INTO Fact_Sales
                (product_id, date_id, quantity_sold, imputation_flag, tally_date_flag,
                 transaction_type, entered_by, order_type)
            VALUES (?, ?, ?, 0, 0, ?, ?, ?)
        """, (pid, dates[iso], qty, ttype.lower(), user, order_type))
        imported += 1

    c.commit()
    return jsonify({
        "ok": True, "imported": imported, "updated": 0,
        "rejected": rejected[:50], "rejected_total": len(rejected),
        **_held_summary(held),
        "rows_read": len(rows), "saved_to": saved,
    })


# ---------------------------------------------------------------- events

@app.get("/api/events")
def get_events():
    rows = dbmod.rows(con(), "SELECT * FROM Event_Log ORDER BY event_date DESC, event_id DESC")
    for r in rows:
        r["is_local"] = False
        # The interface reads `calendar_date`, as POST's reply already names
        # it; without this the Flag an Event list showed labels with no date.
        r["calendar_date"] = r["event_date"]
    return jsonify(rows)


@app.get("/api/events/labels")
def get_event_labels():
    """Every event name already on record - the published calendar's event
    days (calendar_ranges.csv, read-only) and the Event_Log - so Flag an
    Event can offer them as the staff type. One entry per name, ignoring
    case, newest first; a reviewer's "(date uncertain - verify)" note on the
    calendar file is not part of the name."""
    dated = []
    try:
        with open(DATA_DIR / "calendar_ranges.csv", newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                if (row.get("flag") or "").strip() == "is_event_day":
                    dated.append(((row.get("start_date") or "").strip(), row.get("event_name") or ""))
    except (OSError, KeyError):
        pass
    dated += [(r["event_date"] or "", r["event_name"] or "")
              for r in dbmod.rows(con(), "SELECT event_date, event_name FROM Event_Log")]

    labels, seen = [], set()
    for _day, name in sorted(dated, key=lambda d: d[0], reverse=True):
        name = re.sub(r"\(date uncertain[^)]*\)", "", name).strip()
        if name and name.lower() not in seen:
            seen.add(name.lower())
            labels.append(name)
    return jsonify(labels)


def _calendar_gap(c, iso_date):
    """Why `iso_date` cannot hold a closure or event, or None when it can.
    Dim_Date is extended to the end of next year first (_extend_calendar), so
    this is a date before the calendar starts or improbably far ahead - and a
    flag on a day with no Dim_Date row would be silently lost."""
    if dbmod.one(c, "SELECT 1 AS ok FROM Dim_Date WHERE calendar_date = ?", (iso_date,)):
        return None
    return validation.calendar_gap_message(c, iso_date)


def _published_event_days():
    """Days calendar_ranges.csv itself marks is_event_day (Paskuhan and the
    like). Read-only - the file is protected (CLAUDE.md). Needed when an event
    moves off a day: that day keeps its flag if the published calendar has
    one there too, exactly as populate_dim_date.py would rebuild it."""
    days = set()
    try:
        with open(DATA_DIR / "calendar_ranges.csv", newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                if (row.get("flag") or "").strip() != "is_event_day":
                    continue
                start = date.fromisoformat(row["start_date"].strip())
                end = date.fromisoformat(row["end_date"].strip())
                while start <= end:
                    days.add(start.isoformat())
                    start += timedelta(days=1)
    except (OSError, ValueError, KeyError):
        pass
    return days


@app.put("/api/events/<int:event_id>")
def edit_event(event_id):
    """Change an event's date, label or description - an upcoming one when the
    calendar shifts, a past one to correct the record (it reaches the
    forecasts at the next pipeline run). {calendar_date, event_name,
    event_description?}.

    Dim_Date follows at once: the new day is flagged, and the old day loses
    its flag unless another logged event or the published calendar still
    puts one there (the same rule populate_dim_date.py rebuilds by)."""
    gate = _needs_admin("Changing an event")
    if gate:
        return gate
    payload = request.get_json(silent=True) or {}
    new_date = str(payload.get("calendar_date") or "").strip()
    name = (payload.get("event_name") or "").strip()
    description = (payload.get("event_description") or "").strip()

    c = con()
    row = dbmod.one(c, "SELECT * FROM Event_Log WHERE event_id = ?", (event_id,))
    if row is None:
        return jsonify({"ok": False, "error": "That event no longer exists."}), 404

    errors = {}
    if not validation.is_iso_date(new_date):
        errors["calendar_date"] = "Date must be YYYY-MM-DD."
    else:
        _extend_calendar(c)
        gap = _calendar_gap(c, new_date)
        if gap:
            errors["calendar_date"] = gap
    if not name:
        errors["event_name"] = "Give the event a label."
    if errors:
        return jsonify({"ok": False, "errors": errors}), 400

    old_date = row["event_date"]
    c.execute("UPDATE Event_Log SET event_date = ?, event_name = ?, event_description = ? "
              "WHERE event_id = ?", (new_date, name, description, event_id))
    c.execute("UPDATE Dim_Date SET is_event_day = 1 WHERE calendar_date = ?", (new_date,))
    if old_date != new_date:
        _refresh_event_flag(c, old_date)
    c.commit()
    return jsonify({"ok": True, "event": {
        "event_id": event_id, "event_date": new_date, "calendar_date": new_date,
        "event_name": name, "event_description": description, "previous_date": old_date,
    }})


def _refresh_event_flag(c, iso_date):
    """Dim_Date.is_event_day for one day, from what is still there: another
    logged event, or the published calendar (populate_dim_date.py's rule)."""
    still = c.execute("SELECT 1 FROM Event_Log WHERE event_date = ? LIMIT 1",
                      (iso_date,)).fetchone() or iso_date in _published_event_days()
    c.execute("UPDATE Dim_Date SET is_event_day = ? WHERE calendar_date = ?",
              (1 if still else 0, iso_date))


@app.delete("/api/events/<int:event_id>")
def delete_event(event_id):
    """Remove an event - an upcoming one that was cancelled, or a past one
    that was logged by mistake. A past removal reaches the forecasts at the
    next pipeline run."""
    gate = _needs_admin("Removing an event")
    if gate:
        return gate
    c = con()
    row = dbmod.one(c, "SELECT event_date, event_name FROM Event_Log WHERE event_id = ?", (event_id,))
    if row is None:
        return jsonify({"ok": False, "error": "That event no longer exists."}), 404
    c.execute("DELETE FROM Event_Log WHERE event_id = ?", (event_id,))
    _refresh_event_flag(c, row["event_date"])
    c.commit()
    return jsonify({"ok": True, "deleted": {"event_id": event_id, "calendar_date": row["event_date"],
                                            "event_name": row["event_name"]}})


@app.post("/api/events")
def add_event():
    payload = request.get_json(silent=True) or {}
    calendar_date = payload.get("calendar_date")
    event_name = (payload.get("event_name") or "").strip()
    event_description = (payload.get("event_description") or "").strip()

    errors = {}
    if not calendar_date:
        errors["calendar_date"] = "Pick a date."
    elif not validation.is_iso_date(calendar_date):
        # Stored as given, and read back by populate_dim_date.py on every
        # pipeline run - one malformed date used to fail every later run.
        errors["calendar_date"] = "Date must be YYYY-MM-DD."
    if not event_name:
        errors["event_name"] = "Give the event a label."
    if errors:
        return jsonify({"ok": False, "errors": errors}), 400

    c = con()
    _extend_calendar(c)     # see set_closure
    cur = c.execute("""
        INSERT INTO Event_Log (event_date, event_name, event_description, created_by, date_logged)
        VALUES (?, ?, ?, ?, ?)
    """, (calendar_date, event_name, event_description, auth.current_user(),
          datetime.now().isoformat(timespec="seconds")))
    c.execute("UPDATE Dim_Date SET is_event_day = 1 WHERE calendar_date = ?", (calendar_date,))
    c.commit()

    event = {
        "event_id": cur.lastrowid,
        "event_date": calendar_date,
        "calendar_date": calendar_date,
        "event_name": event_name,
        "event_description": event_description,
        "created_by": auth.current_user(),
        "is_local": False,
    }
    return jsonify({"ok": True, "event": event})


# -------------------------------------------------------------- forecast

def _forecastable_ids(c):
    """product_ids step4_forecast_model.py produced a forecast for.

    Empty when the pipeline has not run. Callers must read that as "no
    restriction is knowable yet", not as "nothing qualifies"."""
    if not _has_forecast_table(c):
        return set()
    return {r["product_id"] for r in dbmod.rows(
        c, "SELECT DISTINCT product_id FROM Result_Forecast "
           "WHERE model_type IS NOT 'policy_rate'")}


def _monthly_history(c, product_ids, until=None):
    """Observed units per month, summed over exactly `product_ids`, with the
    number of days the store tallied in that month.

    Sent alongside a forecast so the Demand Forecast chart can show the months
    leading up to it. Callers pass the same items the forecast covers - for
    the category model every product in its forecast_category, for the
    sum-of-items fallback only the Fast contributors - so the history and the
    forecast on one chart are totals of the same thing.

    `tally_days` is what lets the chart put a month on the same axis as the
    daily forecast (units / tally_days = units per day). It counts every day
    the store tallied ANY sales, from these items' first sale on: a tally day
    with no row for these items is a real zero and belongs in the average. A
    month nobody tallied at all is simply absent.

    `until` (an ISO date) stops the history where the forecast's own training
    data stops - the category model trims the zero-padded days after the last
    sale (its `history_end`), and averaging those days into the last month
    would show a dip that is not in the data the forecast was made from."""
    ids = list(product_ids)
    if not ids:
        return []
    marks = ",".join("?" * len(ids))
    return dbmod.rows(c, f"""
        WITH tally AS (
            SELECT DISTINCT d.calendar_date AS date, d.date_id
            FROM Fact_Sales f JOIN Dim_Date d ON d.date_id = f.date_id
            WHERE {dbmod.SALE_ONLY}
        ),
        mine AS (
            SELECT f.date_id, SUM(f.quantity_sold) AS units
            FROM Fact_Sales f WHERE f.product_id IN ({marks}) AND {dbmod.SALE_ONLY}
            GROUP BY f.date_id
        ),
        first_sale AS (
            SELECT MIN(t.date) AS date
            FROM tally t JOIN mine m ON m.date_id = t.date_id
            WHERE m.units > 0
        )
        SELECT substr(t.date, 1, 7) AS month,
               SUM(COALESCE(m.units, 0)) AS units,
               COUNT(*) AS tally_days
        FROM tally t LEFT JOIN mine m ON m.date_id = t.date_id
        WHERE t.date >= (SELECT date FROM first_sale)
          AND (? IS NULL OR t.date <= ?)
        GROUP BY 1 ORDER BY 1
    """, [*ids, until, until])


def _has_forecast_table(c):
    tables = {r[0] for r in c.execute(
        "SELECT name FROM sqlite_master WHERE type='table'")}
    return "Result_Forecast" in tables and c.execute(
        "SELECT COUNT(*) FROM Result_Forecast "
        "WHERE model_type IS NOT 'policy_rate'").fetchone()[0] > 0


# No column in Dim_Product says whether an item is still stocked - `is_active`
# is 1 for every one of the 519 products, and `payment_status` is "UNKNOWN" for
# the items this would matter for, so neither is a usable signal. This infers
# it instead: an item with no recorded sale in the last year is flagged as
# likely no longer stocked, so the screen can say so instead of showing a
# near-zero forecast with no explanation. It is a heuristic, not a confirmed
# status - a slow, genuinely seasonal item could trip it too, which is why the
# wording below says "likely" and gives the actual last-sale date to check.
DISCONTINUED_DAYS = 365


def _last_sale_date(c, product_id):
    """The most recent date this product recorded a real (>0) sale, or None
    if it never has."""
    row = dbmod.one(c, f"""
        SELECT MAX(d.calendar_date) AS d
        FROM Fact_Sales f JOIN Dim_Date d ON d.date_id = f.date_id
        WHERE f.product_id = ? AND f.quantity_sold > 0 AND {dbmod.SALE_ONLY}
    """, (product_id,))
    return row["d"] if row else None


def _discontinued_flags(last_sale_date, reference_date):
    """(days_since_last_sale, likely_discontinued) as of `reference_date`
    (a forecast's own snapshot_date/history_end, not wall-clock "now" - the
    calendar here runs past the real date). Never having sold at all counts
    as likely discontinued too, though that should not occur for a Fast item."""
    if not last_sale_date or not reference_date:
        return None, last_sale_date is None
    days = (date.fromisoformat(reference_date) - date.fromisoformat(last_sale_date)).days
    return days, days >= DISCONTINUED_DAYS


# Wording is aimed at the person using the dashboard, not at whoever wrote the
# ETL: from the frontend there is exactly one action that fixes this, and it is
# the Tally Interface's pipeline runner.
FORECAST_PENDING_REASON = (
    "The pipeline has not been run yet, please run it from the tally interface and this page will automatically display the results."
)

_FORECAST_PENDING = {
    "available": False,
    "reason": FORECAST_PENDING_REASON,
    "data": None,
}


# ------------------------------------------------------ forecast guidance

# Flagged dates no more than this many days apart are one window: a weekend or
# a holiday inside an exam week must not split it, and two exam periods two
# months apart must not merge into one "10/09 to 12/20" stretch.
WINDOW_GAP_DAYS = 3

GUIDANCE_DAYS = 30            # the horizon the forecast and the advice cover
REORDER_LIST_LIMIT = 8        # items named in a category's "reorder these" list


def _date_runs(dates, gap=WINDOW_GAP_DAYS):
    """ISO dates -> [{"start", "end", "days"}], one per unbroken window."""
    runs = []
    for d in sorted(set(dates)):
        if runs and (date.fromisoformat(d) - date.fromisoformat(runs[-1]["end"])).days <= gap:
            runs[-1]["end"] = d
            runs[-1]["days"] += 1
        else:
            runs.append({"start": d, "end": d, "days": 1})
    return runs


def _calendar_window(c, start, end):
    """What the school calendar and the event log hold between two dates:
    the calendar signals the forecast is fitted on, so the advice can name
    them rather than leave a spike on the chart unexplained."""
    days = dbmod.rows(c, """
        SELECT calendar_date, is_store_closed, is_enrollment_period, is_exam_week, is_sem_break
        FROM Dim_Date WHERE calendar_date BETWEEN ? AND ? ORDER BY calendar_date
    """, (start, end))
    events = dbmod.rows(c, """
        SELECT event_date AS date, event_name AS name FROM Event_Log
        WHERE event_date BETWEEN ? AND ? ORDER BY event_date
    """, (start, end)) if "Event_Log" in _table_names(c) else []

    def runs(col):
        return _date_runs([d["calendar_date"] for d in days if d[col]])

    return {
        "start": start, "end": end,
        "enrollment": runs("is_enrollment_period"),
        "exams": runs("is_exam_week"),
        "sem_break": runs("is_sem_break"),
        "closed_days": sum(1 for d in days if d["is_store_closed"]),
        "events": events,
    }


def _busiest_week(rows, total):
    """The 7 consecutive forecast days with the most demand."""
    if len(rows) < 7 or total <= 0:
        return None
    sums = [sum(r["yhat"] for r in rows[i:i + 7]) for i in range(len(rows) - 6)]
    i = max(range(len(sums)), key=sums.__getitem__)
    return {"start": rows[i]["forecast_date"][:10], "end": rows[i + 6]["forecast_date"][:10],
            "units": round(sums[i], 1), "share": round(sums[i] / total, 3)}


def _item_stock_advice(pid, stats, reorder, rate, total, high, today):
    """Stock against the forecast for one item, and the one thing to do.

    Dates are projected from TODAY at the forecast's average daily rate, not
    along the forecast's own dates: the forecast starts the day after the last
    tally, which can be weeks ago, and "order by" a date already gone is no
    advice at all. The reorder point and the suggested quantity are
    Result_Prescriptive's (shared with Reorder Alerts via _reorder_items), so
    this page and that one cannot disagree about an item."""
    st = stats.get(pid, {})
    r = reorder.get(pid)
    stock = st.get("current_stock")
    rop = r["reorder_point"] if r else None
    out = {
        "current_stock": stock, "stock_as_of": st.get("stock_as_of"), "reorder_point": rop,
        "days_cover": None, "runs_out_on": None, "order_by": None,
        "suggested_order_qty": r["suggested_order_qty"] if r else None,
        "to_cover_30d": None, "to_cover_30d_high": None,
    }
    if total < 0.5:
        return out, "no_demand"
    if stock is None:
        return out, "count_stock"

    out["to_cover_30d"] = max(0, math.ceil(total - stock))
    out["to_cover_30d_high"] = max(0, math.ceil(high - stock))
    if rate > 0:
        out["days_cover"] = round(stock / rate, 1)
        out["runs_out_on"] = (today + timedelta(days=int(stock / rate))).isoformat()
        if rop is not None and stock > rop:
            out["order_by"] = (today + timedelta(days=int((stock - rop) / rate))).isoformat()

    if r and r["needs_reorder"]:
        return out, "reorder_now"
    if out["order_by"] and out["order_by"] <= (today + timedelta(days=GUIDANCE_DAYS - 1)).isoformat():
        return out, "order_by"
    if out["to_cover_30d"] > 0:
        return out, "short"                    # no ROP to warn earlier, but the month outruns stock
    if out["to_cover_30d_high"] > 0:
        return out, "watch"                    # covers the expected month, not a busy one
    return out, "covered"


def _category_stock_advice(product_ids, stats, reorder, total):
    """A category's stock position is the items in it: which of them are at
    their reorder point now, from the same rule as Reorder Alerts. Category
    stock is not summed against the category forecast - counts exist for only
    some items, and a partial sum against a whole-category figure would make
    every category look short."""
    ids = set(product_ids)
    counted = sum(1 for pid in ids if stats.get(pid, {}).get("current_stock") is not None)
    due = sorted((reorder[pid] for pid in ids if pid in reorder and reorder[pid]["needs_reorder"]),
                 key=lambda i: -i["suggested_order_qty"])
    out = {
        "items_total": len(ids), "items_counted": counted,
        "reorder_now_total": len(due),
        "reorder_units_total": sum(i["suggested_order_qty"] for i in due),
        "reorder_now": [{k: i[k] for k in ("product_id", "item_name", "supplier_name", "current_stock",
                                           "reorder_point", "suggested_order_qty")}
                        for i in due[:REORDER_LIST_LIMIT]],
    }
    if total < 0.5:
        return out, "no_demand"
    if due:
        return out, "reorder_items"
    if counted == 0:
        return out, "count_stock"
    return out, "covered"


def _forecast_guidance(c, forecast, history, product_ids, single_item, rmse=None):
    """The forecast chart, read out in plain terms for whoever orders stock:
    how much to expect, against what has been selling, when the busy days
    are, what the calendar holds, and what to do about stock.

    Everything here is derived from numbers the pipeline already produced
    (Result_Forecast / Result_Category_Forecast, Result_Prescriptive, the
    stock counts) - it explains them, it does not add a second forecast.
    Returns None when there is nothing to read."""
    rows = [r for r in forecast if r.get("yhat") is not None]
    if not rows:
        return None
    n = len(rows)
    total = sum(r["yhat"] for r in rows)
    low = sum(r["yhat"] if r.get("yhat_lower") is None else r["yhat_lower"] for r in rows)
    high = sum(r["yhat"] if r.get("yhat_upper") is None else r["yhat_upper"] for r in rows)
    # An item's daily band is one day's spread repeated on every day (step4),
    # so its 30-day sum overstates the range several times over. With the
    # item's own 30-day RMSE the range is total +/- RMSE instead - the rule
    # the category band follows, and the range the Expected Demand table shows
    # for the same item (_range_30d), so the two pages cannot disagree.
    band = _range_30d(total, rmse)
    if band:
        low, high = band["low"], band["high"]
    rate = total / n
    today = date.today()
    start, end = rows[0]["forecast_date"][:10], rows[-1]["forecast_date"][:10]

    # The last month on record, as units per tally day - the same axis the
    # chart puts history on, so "up 20%" matches what the eye sees there.
    recent = next((h for h in reversed(history or []) if h.get("tally_days")), None)
    recent_rate = recent["units"] / recent["tally_days"] if recent else None

    # Calendar for the days the advice is about: the coming 30 from today, or
    # the forecast's own window when that has not started yet.
    cal_start = max(today.isoformat(), start)
    cal_end = (date.fromisoformat(cal_start) + timedelta(days=GUIDANCE_DAYS - 1)).isoformat()

    stats, _ = catalog.compute_stats(c)
    reorder = {i["product_id"]: i for i in (_reorder_items(c) or [])}
    if single_item:
        stock, action = _item_stock_advice(product_ids[0], stats, reorder, rate, total, high, today)
    else:
        stock, action = _category_stock_advice(product_ids, stats, reorder, total)

    return {
        "window": {"start": start, "end": end, "days": n, "passed": end < today.isoformat()},
        "expected": {"total": round(total, 1), "low": round(low, 1), "high": round(high, 1),
                     "per_day": round(rate, 2)},
        "recent": ({"month": recent["month"], "per_day": round(recent_rate, 2)} if recent else None),
        "change_pct": (round(100 * (rate - recent_rate) / recent_rate) if recent_rate else None),
        "busiest_week": _busiest_week(rows, total),
        "calendar": _calendar_window(c, cal_start, cal_end),
        "scope": "item" if single_item else "category",
        "stock": stock,
        "action": action,
        "today": today.isoformat(),
    }


@app.get("/api/forecast")
def get_forecast_general():
    c = con()
    if not _has_forecast_table(c):
        return jsonify(_FORECAST_PENDING)
    # Return a summary: list of product_ids that have forecasts
    products_with_forecast = dbmod.rows(c, """
        SELECT DISTINCT rf.product_id, p.item_name, p.fsn_class, p.is_hvl,
               rf.model_type, rf.is_heuristic, rf.snapshot_date
        FROM Result_Forecast rf
        JOIN Dim_Product p ON p.product_id = rf.product_id
        WHERE rf.model_type IS NOT 'policy_rate'
        GROUP BY rf.product_id
    """)
    return jsonify({
        "available": True,
        "reason": None,
        "data": {
            "product_count": len(products_with_forecast),
            "products": products_with_forecast,
        },
    })


@app.get("/api/forecast/<int:product_id>")
def get_forecast(product_id):
    c = con()
    if not _has_forecast_table(c):
        return jsonify(_FORECAST_PENDING)

    forecast_rows = dbmod.rows(c, """
        SELECT forecast_date, yhat, yhat_lower, yhat_upper,
               model_type, is_heuristic, snapshot_date
        FROM Result_Forecast
        WHERE product_id = ? AND model_type IS NOT 'policy_rate'
        ORDER BY forecast_date
    """, (product_id,))

    if not forecast_rows:
        return jsonify({
            "available": False,
            "reason": f"No forecast exists for product {product_id}.",
            "data": None,
        })

    # Metrics from Result_Forecast_Metrics
    tables = {r[0] for r in c.execute(
        "SELECT name FROM sqlite_master WHERE type='table'")}
    metrics = []
    if "Result_Forecast_Metrics" in tables:
        metrics = dbmod.rows(c, """
            SELECT tier, validation_method, period_scope, n_obs,
                   mae, rmse, mape, naive_mae, naive_rmse, naive_mape,
                   beats_naive_mae, meets_mape_threshold, snapshot_date
            FROM Result_Forecast_Metrics
            WHERE product_id = ?
        """, (product_id,))

    product = dbmod.one(c,
        "SELECT item_name, fsn_class, is_hvl FROM Dim_Product WHERE product_id = ?",
        (product_id,))

    snapshot_date = forecast_rows[0]["snapshot_date"]
    last_sale_date = _last_sale_date(c, product_id)
    days_since_last_sale, likely_discontinued = _discontinued_flags(last_sale_date, snapshot_date)
    history = _monthly_history(c, [product_id])
    overall = next((m for m in metrics if m["period_scope"] == "overall"), None)
    rmse = overall["rmse"] if overall else None
    total = sum(r["yhat"] or 0 for r in forecast_rows)

    return jsonify({
        "available": True,
        "reason": None,
        "data": {
            "product_id": product_id,
            "item_name": product["item_name"] if product else None,
            "fsn_class": product["fsn_class"] if product else None,
            "is_hvl": product["is_hvl"] if product else 0,
            "model_type": forecast_rows[0]["model_type"],
            "is_heuristic": bool(forecast_rows[0]["is_heuristic"]),
            "snapshot_date": snapshot_date,
            "last_sale_date": last_sale_date,
            "days_since_last_sale": days_since_last_sale,
            "likely_discontinued": likely_discontinued,
            "forecast": forecast_rows,
            "history": history,
            "metrics": metrics,
            "range_30d": _range_30d(total, rmse),
            "guidance": _forecast_guidance(c, forecast_rows, history, [product_id],
                                           single_item=True, rmse=rmse),
        },
    })


def _table_names(c):
    return {r[0] for r in c.execute(
        "SELECT name FROM sqlite_master WHERE type='table'")}


def _range_30d(total, rmse):
    """The likely 30-day range: the forecast plus or minus its usual miss
    (RMSE) over its own walk-forward 30-day checks, never below zero - the
    same rule the category model's band follows (step4c spreads its RMSE
    across the days, so the summed daily bounds are total +/- RMSE). None
    when the forecast has no checks to size it from."""
    if total is None or rmse is None:
        return None
    return {"low": round(max(total - rmse, 0.0), 1), "high": round(total + rmse, 1),
            "rmse": round(rmse, 1)}


def _item_rmse(c, product_ids):
    """product_id -> overall 30-day RMSE from Result_Forecast_Metrics.

    Not the item's daily band in Result_Forecast: step4 sets that to one
    day's standard deviation on EVERY day, so 30 of them summed is a range
    several times wider than the item's real 30-day miss."""
    if not product_ids or "Result_Forecast_Metrics" not in _table_names(c):
        return {}
    marks = ",".join("?" * len(product_ids))
    return {r["product_id"]: r["rmse"] for r in dbmod.rows(c, f"""
        SELECT product_id, rmse FROM Result_Forecast_Metrics
         WHERE period_scope = 'overall' AND rmse IS NOT NULL AND product_id IN ({marks})
    """, tuple(product_ids))}


def _attach_ranges(c, contributors):
    """Give each contributor its likely 30-day range (see _range_30d)."""
    rmse = _item_rmse(c, [r["product_id"] for r in contributors])
    for r in contributors:
        r["range_30d"] = _range_30d(r.get("yhat_30d"), rmse.get(r["product_id"]))
    return contributors


def _category_contributors(c, category):
    """The category's Fast items that have their own forecast, biggest first.

    `last_sale_date` rides along so the caller can flag an item that likely
    is not stocked any more (see DISCONTINUED_DAYS / _discontinued_flags) -
    computed here as one query rather than one round trip per contributor."""
    return dbmod.rows(c, """
        SELECT p.product_id, p.item_name, p.fsn_class, p.is_hvl,
               COALESCE(p.supplier_name, ?) AS supplier_name,
               SUM(f.yhat) AS yhat_30d,
               (SELECT MAX(d.calendar_date) FROM Fact_Sales f2
                JOIN Dim_Date d ON d.date_id = f2.date_id
                WHERE f2.product_id = p.product_id AND f2.quantity_sold > 0
                  AND LOWER(COALESCE(f2.transaction_type, 'sale')) = 'sale') AS last_sale_date
        FROM Result_Forecast f
        JOIN Dim_Product p ON p.product_id = f.product_id
        WHERE p.category = ? AND f.model_type IS NOT 'policy_rate'
        GROUP BY p.product_id
        ORDER BY yhat_30d DESC
    """, (catalog.UNATTRIBUTED, category))


def _flag_discontinued(contributors, reference_date):
    """Adds days_since_last_sale / likely_discontinued to each contributor
    row in place, as of the category's own reference date. Returns the same
    list for a fluent call at the point contributors are built."""
    for r in contributors:
        r["days_since_last_sale"], r["likely_discontinued"] = \
            _discontinued_flags(r.get("last_sale_date"), reference_date)
    return contributors


def _category_counts(c, category):
    return dbmod.one(c, """
        SELECT COUNT(*) AS n_products,
               SUM(CASE WHEN fsn_class='F' THEN 1 ELSE 0 END) AS n_fast
        FROM Dim_Product WHERE category = ?
    """, (category,))


@app.get("/api/forecast/categories")
def get_forecast_categories():
    """The categories that have a category-level forecast, biggest first.

    The Demand Forecast screen builds its category picker from this. It used
    to derive the picker from the Fast items that have a forecast, which left
    out any category with no Fast item at all (Home & Novelty, Apparel
    Accessories) even though those categories do sell."""
    c = con()
    rows = []
    if "Result_Category_Forecast" in _table_names(c):
        rows = dbmod.rows(c, """
            SELECT forecast_category AS category, SUM(yhat) AS total_30d,
                   MIN(model_type) AS model_type, MIN(snapshot_date) AS snapshot_date
            FROM Result_Category_Forecast
            GROUP BY forecast_category
            ORDER BY total_30d DESC
        """)
    if not rows:
        return jsonify({"available": False, "reason": FORECAST_PENDING_REASON, "data": None})
    return jsonify({"available": True, "reason": None, "data": {"categories": rows}})


@app.get("/api/forecast/category/<path:category>")
def get_forecast_category(category):
    """The 30-day forecast for a whole category.

    Served from Result_Category_Forecast (scripts/step4c_category_forecast.py):
    the category's sales are added up FIRST and that total is forecast, which
    covers every item in the category and measured more accurate than adding
    up item forecasts (see that script's docstring for the comparison). The
    response carries its own walk-forward `metrics` so the screen can say how
    far off it has typically been.

    `contributors` are the Fast items that have their own forecast, from
    Result_Forecast. They are a breakdown, NOT the parts of the total: each
    item is forecast on its own (step4), only Fast items are forecast at all,
    and the category model uses a different window, so the two do not add up.
    `contributors_total_30d` carries the items' own sum so the screen can show
    both numbers, and `source` tells it which shape it is looking at.

    When the category table is absent or has nothing for this category (the
    pipeline has not run step4c, or an older database), this falls back to the
    previous behaviour - the SUM of the Fast items' forecasts - and reports
    `source: "sum_of_items"` so the screen can word it accordingly.
    """
    c = con()
    tables = _table_names(c)
    cat_rows = []
    if "Result_Category_Forecast" in tables:
        cat_rows = dbmod.rows(c, """
            SELECT forecast_date, yhat, yhat_lower, yhat_upper
            FROM Result_Category_Forecast
            WHERE forecast_category = ?
            ORDER BY forecast_date
        """, (category,))
    if not cat_rows:
        return _category_forecast_from_items(c, category)

    head = dbmod.one(c, """
        SELECT model_type, is_heuristic, snapshot_date, history_end
        FROM Result_Category_Forecast WHERE forecast_category = ? LIMIT 1
    """, (category,))
    metrics = []
    if "Result_Category_Forecast_Metrics" in tables:
        metrics = dbmod.rows(c, """
            SELECT validation_method, period_scope, n_obs, mae, rmse, mape, mase,
                   naive_mae, naive_mase, beats_naive_mae, mean_actual_30d
            FROM Result_Category_Forecast_Metrics WHERE forecast_category = ?
        """, (category,))
    contributors = _category_contributors(c, category) if _has_forecast_table(c) else []
    _flag_discontinued(contributors, head["history_end"])
    _attach_ranges(c, contributors)
    meta = _category_counts(c, category)
    # History for the chart covers the same items the category model was
    # trained on: every product whose forecast_category this is, with a blank
    # one counted as "Uncategorised" - step4c_category_forecast.py's RESIDUE.
    series_ids = [r["product_id"] for r in dbmod.rows(c, """
        SELECT product_id FROM Dim_Product
        WHERE COALESCE(forecast_category, 'Uncategorised') = ?
    """, (category,))]
    history = _monthly_history(c, series_ids, until=head["history_end"])

    return jsonify({
        "available": True,
        "reason": None,
        "data": {
            "category": category,
            "source": "category_model",
            "model_type": head["model_type"],
            "snapshot_date": head["snapshot_date"],
            "history_end": head["history_end"],
            "is_heuristic": bool(head["is_heuristic"]),
            "n_products": meta["n_products"] if meta else 0,
            "n_fast": (meta["n_fast"] or 0) if meta else 0,
            "n_forecast": len(contributors),
            "total_30d": round(sum(r["yhat"] or 0 for r in cat_rows), 3),
            "contributors_total_30d": round(sum(r["yhat_30d"] or 0 for r in contributors), 3),
            "forecast": cat_rows,
            "history": history,
            "contributors": contributors,
            "metrics": metrics,
            "guidance": _forecast_guidance(c, cat_rows, history, series_ids, single_item=False),
        },
    })


def _category_forecast_from_items(c, category):
    """The category forecast built by SUMMING the per-SKU rows in
    Result_Forecast - what this endpoint did before step4c existed, kept as
    the fallback so a database without Result_Category_Forecast still works.

    Only Fast SKUs get a forecast, so this total covers the items in
    `contributors` and no others - the response says how many that is, and
    how many products the category holds in total, so the figure is never
    read as covering the whole category when it does not. It always equals the
    sum of the items shown underneath it, which the category model does not.
    """
    if not _has_forecast_table(c):
        return jsonify(_FORECAST_PENDING)

    rows = dbmod.rows(c, """
        SELECT f.forecast_date,
               SUM(f.yhat)       AS yhat,
               SUM(f.yhat_lower) AS yhat_lower,
               SUM(f.yhat_upper) AS yhat_upper
        FROM Result_Forecast f
        JOIN Dim_Product p ON p.product_id = f.product_id
        WHERE p.category = ? AND f.model_type IS NOT 'policy_rate'
        GROUP BY f.forecast_date
        ORDER BY f.forecast_date
    """, (category,))

    if not rows:
        total = dbmod.one(c,
            "SELECT COUNT(*) AS n FROM Dim_Product WHERE category = ?",
            (category,))
        return jsonify({
            "available": False,
            "reason": (f"No item in {category} has a forecast yet. "
                       "Only Fast-moving items are forecast, and this "
                       "category has none."
                       if total and total["n"] else
                       f"No products are in {category}."),
            "data": None,
        })

    contributors = _category_contributors(c, category)
    meta = _category_counts(c, category)

    head = dbmod.one(c, """
        SELECT f.model_type, f.snapshot_date,
               MAX(f.is_heuristic) AS any_heuristic
        FROM Result_Forecast f
        JOIN Dim_Product p ON p.product_id = f.product_id
        WHERE p.category = ? AND f.model_type IS NOT 'policy_rate'
    """, (category,))
    _flag_discontinued(contributors, head["snapshot_date"] if head else None)
    _attach_ranges(c, contributors)

    total = round(sum(r["yhat_30d"] or 0 for r in contributors), 3)
    ids = [r["product_id"] for r in contributors]
    history = _monthly_history(c, ids)
    return jsonify({
        "available": True,
        "reason": None,
        "data": {
            "category": category,
            "source": "sum_of_items",
            "model_type": head["model_type"] if head else None,
            "snapshot_date": head["snapshot_date"] if head else None,
            "history_end": None,
            "is_heuristic": bool(head["any_heuristic"]) if head else False,
            "n_products": meta["n_products"] if meta else 0,
            "n_fast": (meta["n_fast"] or 0) if meta else 0,
            "n_forecast": len(contributors),
            "total_30d": total,
            "contributors_total_30d": total,
            "forecast": rows,
            "history": history,
            "contributors": contributors,
            "metrics": [],
            "guidance": _forecast_guidance(c, rows, history, ids, single_item=False),
        },
    })


# ------------------------------------------------------------ advisories

@app.get("/api/advisories")
def get_advisories():
    """Calendar-contextual advisories: combine demand forecast (if any) with
    upcoming calendar events to produce actionable stocking recommendations.
    Returns an honest pending shape when the forecast hasn't been generated."""
    c = con()
    has_forecast = _has_forecast_table(c)
    # Read once: every advisory's item list and the decision-rule ones below
    # need the stock position, and each compute_catalog is a full scan.
    products = {p["product_id"]: p for p in catalog.compute_catalog(c)}

    def affected_items(fsn=("F",), limit=8, ids=None):
        """The SKUs an advisory is actually about, with the numbers that make
        it checkable: forecast, stock on hand, reorder point.

        An advisory that says "make sure Fast-moving items are stocked" is not
        actionable until it names them, so every advisory carries its own list
        and the interface shows it on click. LEFT JOINs throughout - a SKU with
        no forecast or no stock count still belongs in the list, reported as
        null rather than dropped.

        `ids`, when given, names the SKUs outright (already ranked by the
        caller) instead of taking a class."""
        if ids is not None:
            ids = list(ids)[:limit]
            if not ids:
                return []
            where, args = f"p.product_id IN ({','.join('?' * len(ids))})", tuple(ids)
        else:
            where, args = f"p.fsn_class IN ({','.join('?' * len(fsn))})", tuple(fsn)
        rows = dbmod.rows(c, f"""
            SELECT p.product_id, p.item_name, p.fsn_class, p.category,
                   COALESCE(p.supplier_name, ?) AS supplier_name,
                   {"(SELECT SUM(yhat) FROM Result_Forecast rf WHERE rf.product_id = p.product_id"
                    " AND rf.model_type IS NOT 'policy_rate')"
                    if has_forecast else "NULL"} AS forecast_30d,
                   (SELECT r.reorder_point FROM Result_Prescriptive r
                     WHERE r.product_id = p.product_id LIMIT 1) AS reorder_point
              FROM Dim_Product p
             WHERE {where}
             ORDER BY {"forecast_30d DESC," if has_forecast else ""} p.item_name
             LIMIT ?
        """, (catalog.UNATTRIBUTED, *args, limit))
        if ids is not None:
            rank = {pid: i for i, pid in enumerate(ids)}
            rows.sort(key=lambda r: rank[r["product_id"]])
        stock = {pid: p["current_stock"] for pid, p in products.items()}
        for r in rows:
            r["current_stock"] = stock.get(r["product_id"])
            r["below_rop"] = (
                r["current_stock"] is not None and r["reorder_point"] is not None
                and r["current_stock"] <= r["reorder_point"]
            )
        return rows

    # "Upcoming" is the store's own today. SQLite's date('now') is UTC, which
    # in Manila is still yesterday until 8 a.m.
    today = date.today().isoformat()

    # Upcoming events from Event_Log
    upcoming_events = dbmod.rows(c, """
        SELECT e.event_date, e.event_name, e.event_description,
               d.is_enrollment_period, d.is_exam_week, d.is_sem_break,
               d.semester_id, d.semester_week
        FROM Event_Log e
        LEFT JOIN Dim_Date d ON d.calendar_date = e.event_date
        WHERE e.event_date >= ?
        ORDER BY e.event_date
        LIMIT 10
    """, (today,))

    # Upcoming calendar periods (enrollment, exams) from Dim_Date. Every
    # upcoming flagged date, grouped into windows below: this used to take the
    # first 30 flagged dates and report the first and last of them as one
    # range, which turned midterms (Oct 9-14) and finals (Dec 14-20) into
    # "Exams scheduled 10/09 to 12/20" - ten weeks of quiet that are not there.
    upcoming_periods = dbmod.rows(c, """
        SELECT calendar_date, is_enrollment_period, is_exam_week
        FROM Dim_Date
        WHERE calendar_date >= ?
          AND (is_enrollment_period = 1 OR is_exam_week = 1)
        ORDER BY calendar_date
    """, (today,))

    advisories = []

    def span(run):
        return (_us_date(run["start"]) if run["start"] == run["end"]
                else f"{_us_date(run['start'])} to {_us_date(run['end'])}")

    # One advisory per window, so each carries its own dates.
    enrollment_runs = _date_runs([r["calendar_date"] for r in upcoming_periods if r["is_enrollment_period"]])
    for run in enrollment_runs:
        advisories.append({
            "type": "enrollment",
            "severity": "high",
            "title": "Enrollment Period Approaching",
            "description": f"Enrollment runs {span(run)}. "
                           "Historically the highest-volume sales window — "
                           "ensure Fast-moving items (uniforms, IDs, school supplies) are stocked.",
            "date_range": [run["start"], run["end"]],
            "has_forecast": has_forecast,
            "basis": f"{run['days']} upcoming date(s) in Dim_Date carry is_enrollment_period = 1. "
                     "The items below are the Fast-moving class (step3's 80th-percentile ADUS "
                     "cut), which is what an enrollment surge draws on.",
            "items": affected_items(("F",)),
        })

    for run in _date_runs([r["calendar_date"] for r in upcoming_periods if r["is_exam_week"]]):
        advisories.append({
            "type": "exam_week",
            "severity": "medium",
            "title": "Exam Week Upcoming",
            # Measured, not assumed: past exam weeks sold about the same per
            # day as regular days (tallied open days: ~160 vs ~158 units/day
            # over all history, ~185 vs ~184 over the last year), so the
            # advice is to restock as usual rather than to hold off.
            "description": f"Exams scheduled {span(run)}. "
                           "Past exam weeks sold about the same per day as regular days; "
                           "restock as usual.",
            "date_range": [run["start"], run["end"]],
            "has_forecast": has_forecast,
            "basis": f"{run['days']} upcoming date(s) in Dim_Date carry is_exam_week = 1. "
                     "Across past exam weeks, tallied sales averaged about the same per day "
                     "as regular days, though single weeks vary widely.",
            "items": affected_items(("F",)),
        })

    for event in upcoming_events:
        advisories.append({
            "type": "event",
            "severity": "medium",
            "title": event["event_name"],
            "description": event["event_description"] or
                           f"Event on {_us_date(event['event_date'])}. Check stock for high-demand items.",
            "date_range": [event["event_date"], event["event_date"]],
            "has_forecast": has_forecast,
            "basis": f"Logged in Event_Log for {_us_date(event['event_date'])}"
                     + (f", which falls in semester {event['semester_id']} week "
                        f"{event['semester_week']}." if event["semester_id"] else "."),
            "items": affected_items(("F",)),
        })

    # If forecast exists, add top fast-movers that may need restocking
    if has_forecast:
        # Get fast-moving items with low projected supply
        fast_items = dbmod.rows(c, """
            SELECT p.item_name, p.product_id,
                   SUM(rf.yhat) AS total_forecast_30d
            FROM Result_Forecast rf
            JOIN Dim_Product p ON p.product_id = rf.product_id
            WHERE p.fsn_class = 'F' AND rf.model_type IS NOT 'policy_rate'
            GROUP BY rf.product_id
            ORDER BY total_forecast_30d DESC
            LIMIT 5
        """)
        if fast_items:
            top_names = ", ".join(i["item_name"] for i in fast_items[:3])
            advisories.insert(0, {
                "type": "forecast_alert",
                "severity": "high",
                "title": "Top Forecasted Demand — Next 30 Days",
                "description": f"Highest projected demand: {top_names}. "
                               "Review stock levels against the forecast on the Demand Forecast page.",
                "date_range": None,
                "has_forecast": True,
                "basis": "Ranked by SUM(yhat) over the 30-day horizon in Result_Forecast, "
                         "Fast class only. Stock and reorder point come from the latest "
                         "count and Result_Prescriptive.",
                "items": affected_items(("F",), limit=5),
            })

    if not advisories:
        advisories.append({
            "type": "info",
            "severity": "low",
            "title": "No upcoming calendar signals",
            "description": "No enrollment, exam weeks, or flagged events are coming up. "
                           "Advisories appear here when calendar signals exist."
                           + ("" if has_forecast
                              else " The demand forecast has not been generated yet — "
                                   "run the pipeline from the Tally Interface to enable "
                                   "forecast-based advisories."),
            "date_range": None,
            "has_forecast": has_forecast,
            "basis": "No upcoming row in Dim_Date has is_enrollment_period, is_exam_week or "
                     "is_event_day set, and Event_Log has no future entry.",
            "items": [],
        })

    # After the calendar fallback on purpose: these are standing decisions
    # about the catalogue, not calendar signals, so they must not hide the
    # "no upcoming calendar signals" card when that is the truth.
    advisories.extend(_decision_rule_advisories(
        c, products, enrollment_runs, upcoming_events, affected_items, has_forecast))

    return jsonify({
        "available": True,
        "has_forecast": has_forecast,
        "advisories": advisories,
    })


def _next_high_traffic(enrollment_runs, upcoming_events):
    """The next enrollment window or logged event, whichever starts first, as
    (what, start, end) - or None when the calendar holds neither."""
    options = [("enrollment", r["start"], r["end"]) for r in enrollment_runs[:1]]
    options += [(e["event_name"], e["event_date"], e["event_date"]) for e in upcoming_events[:1]]
    return min(options, key=lambda o: o[1]) if options else None


def _decision_rule_advisories(c, products, enrollment_runs, upcoming_events,
                              affected_items, has_forecast):
    """The FSN decision rules (manuscript Sec. 3.3.1) as advisories: promote
    the Slow movers, review the ones slow for two semesters running, clear out
    the Non-moving items still holding stock. Each appears only when it names
    at least one item. The rule itself lives in decision_rules.py; this only
    words it and picks which items to list."""
    rules = decision_rules.evaluate(c)
    by_action = {}
    for item in rules["items"]:
        by_action.setdefault(item["action"], []).append(item["product_id"])

    def most_stock_first(pids):
        """Most units on hand first (uncounted last), then most units sold:
        the items tying up the most shelf stock lead the list."""
        def key(pid):
            p = products.get(pid, {})
            stock = p.get("current_stock")
            return (stock is None, -(stock or 0), -(p.get("total_units") or 0))
        return sorted(pids, key=key)

    def plural(n, word):
        return f"{n} {word}{'' if n == 1 else 's'}"

    out = []

    promote = by_action.get(decision_rules.PROMOTE, [])
    if promote:
        window = _next_high_traffic(enrollment_runs, upcoming_events)
        if window:
            what, start, end = window
            when = _us_date(start) if start == end else f"{_us_date(start)} to {_us_date(end)}"
            during = f"enrollment ({when})" if what == "enrollment" else f"{what} ({when})"
        else:
            during = "the next enrollment period or campus event"
        out.append({
            "type": "slow_moving",
            "severity": "medium",
            "title": "Promote Slow-Moving Items",
            "description": f"{plural(len(promote), 'slow-moving item')} "
                           f"{'sells' if len(promote) == 1 else 'sell'} only now and then. "
                           "Move them with bundled pricing, better shelf placement, or targeted "
                           f"marketing during {during}.",
            "date_range": [window[1], window[2]] if window else None,
            "has_forecast": has_forecast,
            "basis": "Slow-moving items sold in the records but below the Fast cut (step3's "
                     "80th-percentile ADUS), so they get no item forecast; the manuscript's rule "
                     "for them is promotion (Sec. 3.3.1). Listed by units on hand, then units "
                     "sold. Items slow for two semesters running are under Review for "
                     "Discontinuation instead.",
            "items": affected_items(ids=most_stock_first(promote)),
        })

    review = by_action.get(decision_rules.REVIEW, [])
    if review:
        sems = " and ".join(s["semester_id"] for s in rules["semesters"])
        out.append({
            "type": "discontinuation",
            "severity": "medium",
            "title": "Review for Discontinuation",
            "description": f"{plural(len(review), 'slow-moving item')} sold below "
                           f"{rules['threshold']:g} units per selling day in both {sems}. "
                           "Decide whether to keep carrying them.",
            "date_range": None,
            "has_forecast": has_forecast,
            "basis": "Velocity is ADUS per regular semester, measured as the FSN model measures "
                     "it. An item carried before a semester but absent from its tally sheets "
                     "counts as 0 for it. The threshold can be changed on the FSN Classification "
                     f"page; this card uses the default, {rules['default_threshold']:g}.",
            "items": affected_items(ids=most_stock_first(review)),
        })

    holding = [pid for pid in by_action.get(decision_rules.CLEAR_OUT, [])
               if (products.get(pid, {}).get("current_stock") or 0) > 0]
    if holding:
        units = sum(products[pid]["current_stock"] for pid in holding)
        out.append({
            "type": "non_moving",
            "severity": "medium",
            "title": "Non-Moving Items Holding Stock",
            "description": f"{plural(len(holding), 'item')} "
                           f"{'has' if len(holding) == 1 else 'have'} no recorded sales but still "
                           f"{'holds' if len(holding) == 1 else 'hold'} "
                           f"{units:,} units on hand. Evaluate returning them to the supplier, "
                           "marking them down for clearance, or discontinuing them.",
            "date_range": None,
            "has_forecast": has_forecast,
            "basis": "Non-moving items have no recorded sale (step3) and are left out of every "
                     "reorder calculation (Sec. 3.3.1). Only those with a stock count above zero "
                     "are listed: without stock there is nothing to clear out.",
            "items": affected_items(ids=most_stock_first(holding)),
        })

    return out


# ------------------------------------------------------ tally workbooks

@app.post("/api/tally/workbook")
def add_tally_workbook():
    """Add the store's own tally workbook (the monthly "<MONTH> - TBS" sheets)
    to rawdata/, where the next pipeline run's step0 reads it - so a new
    academic year's workbook needs no developer and no file copying.

    Checked before it is kept: it must have at least one tally sheet. A
    workbook of the same name already in rawdata/ is moved to
    rawdata/replaced/ first, never overwritten in place. The response says
    which months it covers and which of them another workbook also has -
    step0 reads those months from this one, the most recently added."""
    import openpyxl
    from step0_convert_sales_with_zeros import (
        APP_IMPORT_PREFIX, find_workbooks, sheet_month, tbs_sheet_names)

    storage = request.files.get("file")
    if storage is None or not storage.filename:
        return jsonify({"ok": False, "error": "Choose the tally workbook (.xlsx) to add."}), 400
    if not storage.filename.lower().endswith(".xlsx"):
        return jsonify({"ok": False, "error": "A tally workbook is an .xlsx file."}), 400
    payload = storage.read()
    # The store's own file name, kept as it is (spaces and all) so a newer copy
    # of the same workbook replaces the old one instead of sitting beside it.
    # Only a directory part and characters Windows forbids are removed.
    filename = re.sub(r'[<>:"|?*\x00-\x1f]', "_", re.split(r"[\\/]", storage.filename)[-1]).strip()
    if filename.startswith(APP_IMPORT_PREFIX):     # step0 skips import_* (app import archives)
        filename = "tally_" + filename
    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:                       # openpyxl raises a zoo
        return jsonify({"ok": False, "error": f"Could not read that workbook ({exc})."}), 400
    try:
        months = sorted({m for m in (sheet_month(wb[sn]) for sn in tbs_sheet_names(wb, filename)) if m})
    finally:
        wb.close()
    if not months:
        return jsonify({"ok": False, "error": "No tally sheets found. The pipeline reads sheets named like "
                                              "'AUGUST 2026 - TBS' with the dates across the first row."}), 400

    try:
        others = [p for p in find_workbooks(str(RAWDATA_DIR)) if Path(p).name != filename]
    except FileNotFoundError:
        others = []
    overlap = {}                                   # month -> other workbooks that have it
    for path in others:
        wbo = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sn in tbs_sheet_names(wbo, path):
                m = sheet_month(wbo[sn])
                if m in months:
                    overlap.setdefault(m, []).append(Path(path).name)
        finally:
            wbo.close()

    target = RAWDATA_DIR / filename
    replaced = None
    try:
        RAWDATA_DIR.mkdir(parents=True, exist_ok=True)
        if target.exists():
            old = RAWDATA_DIR / "replaced" / f"{datetime.now().strftime('%Y%m%d-%H%M%S')}_{filename}"
            old.parent.mkdir(exist_ok=True)
            target.replace(old)
            replaced = str(old.relative_to(RAWDATA_DIR.parent))
        target.write_bytes(payload)
    except OSError as exc:
        return jsonify({"ok": False, "error": f"Could not save the workbook to rawdata/ ({exc})."}), 500

    return jsonify({
        "ok": True, "saved_as": filename, "months": months,
        "replaced_file": replaced,
        "takes_over": [{"month": m, "from": sorted(set(f))} for m, f in sorted(overlap.items())],
    })


# --------------------------------------------------------------- pipeline

@app.post("/api/pipeline/run")
def run_pipeline():
    """Body (optional): {"skip": ["step4"]}, or the friendlier
    {"include_forecast": false} the Tally Interface sends.

    Only step4 can be opted out of (pipeline.SKIPPABLE) - it is the only step
    that costs hours rather than seconds and the only one whose output no
    later step reads. Everything else always runs, so a "quick" run still
    produces a fully rebuilt database, FSN classes and reorder points."""
    payload = request.get_json(silent=True) or {}
    skip = set(payload.get("skip") or [])
    if payload.get("include_forecast") is False:
        skip.add("step4")

    started = pipeline.start_pipeline(skip)
    if not started:
        return jsonify({"ok": False, "error": "The pipeline is already running."}), 409
    return jsonify({"ok": True, "skipped": sorted(skip & pipeline.SKIPPABLE)})


@app.post("/api/pipeline/stop")
def stop_pipeline():
    stopped = pipeline.stop_pipeline()
    if not stopped:
        return jsonify({"ok": False, "error": "No pipeline run is in progress."}), 409
    return jsonify({"ok": True})


@app.get("/api/pipeline/status")
def pipeline_status():
    return jsonify(pipeline.get_status())


@app.get("/api/pipeline/staleness")
def pipeline_staleness():
    """Whether the analytics on screen are older than what has been tallied.

    Everything the Digital Tallying Interface writes lands in the database
    immediately, but fsn_class / Result_Forecast / Result_Prescriptive are
    only recomputed when the pipeline runs - so a screen can be showing
    reorder points that predate a week of tally entries with nothing
    indicating it. See pipeline.get_staleness()."""
    state = pipeline.get_staleness(con())
    # Names settled under "Names to review" since the last run: the vocabulary
    # has them, but Dim_Product and everything after it does not yet.
    if state.get("supported"):
        n = names.awaiting_run_count(con(), VOCAB_CSV)
        state["pending"]["name_decisions"] = n
        state["total_pending"] += n
        state["stale"] = state["stale"] or n > 0
    return jsonify(state)


if __name__ == "__main__":
    app.run(debug=False, port=5000)
